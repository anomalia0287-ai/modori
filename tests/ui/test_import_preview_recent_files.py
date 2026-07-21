def test_import_preview_lists_variables_and_inferred_measures(tmp_path) -> None:
    from tests.ui.test_end_to_end_ui_flow import write_reference_csv
    from modori.ui.controller import UiController

    data_path = tmp_path / "survey.csv"
    write_reference_csv(data_path)
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True

    preview = controller.importPreviewText
    assert "미리 읽은 데이터: 20행 · 9개 변수" in preview
    assert "q1" in preview
    assert "group" in preview
    assert "순서형" in preview
    assert "no labels" not in preview
    assert "no missing" not in preview


def test_import_preview_failure_sets_visible_error(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "notes.txt"
    data_path.write_text("not,data\n", encoding="utf-8")
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is False

    assert controller.importPreviewText == "지원하지 않는 파일 형식입니다."
    assert controller.lastError == "지원하지 않는 파일 형식입니다."


def test_confirm_import_without_preview_sets_visible_error() -> None:
    from modori.ui.controller import UiController

    controller = UiController()

    assert controller.confirmPendingImport() is False

    assert controller.importPreviewText == "가져올 파일이 선택되지 않았습니다."
    assert controller.lastError == "가져올 파일이 선택되지 않았습니다."


def test_confirm_import_binds_preview_models_before_analysis_runs(tmp_path) -> None:
    from tests.ui.test_end_to_end_ui_flow import write_reference_csv
    from modori.ui.controller import UiController

    data_path = tmp_path / "survey.csv"
    write_reference_csv(data_path)
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert controller.confirmPendingImport() is True

    assert controller.dataModel is not None
    assert controller.dataModel.rowCount() == 20
    assert controller.dataModel.columnCount() == 9
    assert controller.variableModel is not None
    assert controller.variableModel.rowCount() == 9
    assert controller.dataViewNotice == "가져온 데이터 미리보기: 20행 · 9열"
    assert controller.resultSummary == ""


def test_confirm_import_keeps_public_data_warning_visible_in_data_notice(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "public-with-preamble.csv"
    data_path.write_text(
        "\n".join(
            [
                "서울시 인구 현황",
                "자료기준일: 2024-12-31",
                "단위: 명",
                "자치구,연도,인구",
                "종로구,2024,140000",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert controller.confirmPendingImport() is True

    assert "가져온 데이터 미리보기: 1행 · 3열" in controller.dataViewNotice
    assert "표 헤더 앞의 안내 행 3개" in controller.dataViewNotice


def test_confirm_import_can_exclude_warned_aggregate_rows(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "aggregate-row.csv"
    data_path.write_text(
        "\n".join(
            [
                "지역,인구",
                "합 계,300",
                "종로구,100",
                "중구,200",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert "집계/합계 행 1개를 감지했습니다." in controller.importPreviewText
    assert controller.confirmPendingImport(True) is True

    assert controller.dataModel is not None
    assert controller.dataModel.rowCount() == 2
    assert "집계/합계 행 1개를 제외했습니다." in controller.dataViewNotice


def test_adjust_pending_import_layout_repreviews_and_confirms_same_layout(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "manual-layout.csv"
    data_path.write_text(
        "\n".join(
            [
                "다운로드 조건,2026-07-06",
                "이 행은 표가 아닙니다,확인용",
                "city,value",
                "Seoul,10",
                "Busan,20",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert controller.previewPendingImportLayout(3, 1, 4, "") is True
    assert "사용자 지정 표 레이아웃을 적용했습니다." in controller.importPreviewText
    assert controller.confirmPendingImport() is True

    assert controller.dataModel is not None
    assert controller.dataModel.rowCount() == 2
    assert controller.dataModel.columnCount() == 2
    assert "가져온 데이터 미리보기: 2행 · 2열" in controller.dataViewNotice


def test_failed_layout_preview_keeps_pending_file_for_correction(tmp_path) -> None:
    from openpyxl import Workbook

    from modori.ui.controller import UiController

    data_path = tmp_path / "multi-sheet.xlsx"
    workbook = Workbook()
    notice = workbook.active
    notice.title = "안내"
    notice.append(["□ 안내문만 있는 시트입니다."])
    data = workbook.create_sheet("자료")
    data.append(["city", "value"])
    data.append(["Seoul", 10])
    workbook.active = workbook.sheetnames.index("자료")
    workbook.save(data_path)
    workbook.close()
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert controller.previewPendingImportLayout(1, 1, 2, "없는시트") is False
    assert "지정한 시트를 찾지 못했습니다." in controller.importPreviewText
    assert controller.previewPendingImportLayout(1, 1, 2, "자료") is True
    assert "사용자 지정 표 레이아웃을 적용했습니다." in controller.importPreviewText


def test_recent_files_persist_to_settings_file(tmp_path, monkeypatch) -> None:
    from tests.ui.test_end_to_end_ui_flow import write_reference_csv
    from modori.ui.controller import UiController

    settings_path = tmp_path / "settings.json"
    monkeypatch.setenv("MODORI_SETTINGS_PATH", str(settings_path))
    data_path = tmp_path / "survey.csv"
    write_reference_csv(data_path)

    first = UiController()
    assert first.openDataFilePath(str(data_path)) is True
    assert "survey.csv" in first.recentFilesText

    second = UiController()
    assert "survey.csv" in second.recentFilesText


def test_recent_file_entry_reopens_existing_file(tmp_path, monkeypatch) -> None:
    from tests.ui.test_end_to_end_ui_flow import write_reference_csv
    from modori.ui.controller import UiController

    settings_path = tmp_path / "settings.json"
    monkeypatch.setenv("MODORI_SETTINGS_PATH", str(settings_path))
    data_path = tmp_path / "survey.csv"
    write_reference_csv(data_path)

    first = UiController()
    assert first.openDataFilePath(str(data_path)) is True

    second = UiController()
    assert second.openRecentFileAt(0) is True

    assert second.pipeline is not None
    assert second.lastError == ""
    assert "survey.csv" in second.recentFilesText


def test_recent_file_entry_reports_missing_file(tmp_path, monkeypatch) -> None:
    from tests.ui.test_end_to_end_ui_flow import write_reference_csv
    from modori.ui.controller import UiController

    settings_path = tmp_path / "settings.json"
    monkeypatch.setenv("MODORI_SETTINGS_PATH", str(settings_path))
    data_path = tmp_path / "survey.csv"
    write_reference_csv(data_path)

    first = UiController()
    assert first.openDataFilePath(str(data_path)) is True
    data_path.unlink()

    second = UiController()
    assert second.openRecentFileAt(0) is False

    assert second.lastError == "최근 파일을 찾을 수 없습니다."
    assert second.pipeline is None


def test_recent_files_can_be_disabled_and_cleared(tmp_path, monkeypatch) -> None:
    from tests.ui.test_end_to_end_ui_flow import write_reference_csv
    from modori.ui.controller import UiController

    settings_path = tmp_path / "settings.json"
    monkeypatch.setenv("MODORI_SETTINGS_PATH", str(settings_path))
    data_path = tmp_path / "survey.csv"
    write_reference_csv(data_path)

    controller = UiController()
    assert controller.openDataFilePath(str(data_path)) is True
    assert controller.setRecentFilesEnabled(False) is True
    assert controller.recentFilesText == ""

    reloaded = UiController()
    assert reloaded.recentFilesText == ""
    assert reloaded.recentFilesEnabled is False


def test_import_preview_exposes_review_rows_with_roles(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "preamble.csv"
    lines = ["자료기준: 2026-07-06", "지역,인구", "종로구,100", "중구,200"]
    data_path.write_text("\n".join(lines), encoding="utf-8")
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True

    rows = controller.importReviewRows
    assert [entry["role"] for entry in rows] == ["skipped", "header", "data", "data"]
    assert rows[1]["cells"] == "지역 | 인구"


def test_import_review_rows_follow_manual_layout_override(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "preamble.csv"
    lines = ["자료기준: 2026-07-06", "지역,인구", "종로구,100", "중구,200"]
    data_path.write_text("\n".join(lines), encoding="utf-8")
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert controller.previewPendingImportLayout(3, 1, 4, "") is True

    rows = controller.importReviewRows
    assert [entry["role"] for entry in rows] == ["skipped", "skipped", "header", "data"]


def test_import_review_rows_are_empty_without_preview() -> None:
    from modori.ui.controller import UiController

    assert UiController().importReviewRows == []


def test_import_flow_supports_duplicate_row_exclusion(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "dupes.csv"
    data_path.write_text("지역,인구\n종로구,100\n종로구,100\n중구,200\n", encoding="utf-8")
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert "완전히 동일한 중복 행 1개를 감지했습니다" in controller.importPreviewText

    assert controller.previewPendingImportLayout(1, 1, 2, "", False, True) is True
    assert "중복 행 1개를 제외했습니다." in controller.importPreviewText

    assert controller.confirmPendingImport(False, True) is True
    import_step = controller.pipeline.steps[0]
    assert import_step.params["drop_duplicate_rows"] is True
    assert len(controller.pipeline.current_dataset.df) == 2


def test_confirm_import_persists_selected_columns_from_preview(tmp_path) -> None:
    from modori.ui.controller import UiController

    data_path = tmp_path / "curated.csv"
    data_path.write_text("지역,인구,비고\n종로구,100,메모\n", encoding="utf-8")
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert controller.importColumnRows == [
        {"name": "지역", "included": True},
        {"name": "인구", "included": True},
        {"name": "비고", "included": True},
    ]
    assert controller.previewPendingImportLayout(
        1,
        1,
        2,
        "",
        False,
        False,
        ["지역", "인구"],
    ) is True
    assert controller.confirmPendingImport(False, False, ["지역", "인구"]) is True

    import_step = controller.pipeline.steps[0]
    assert import_step.params["import_selection"]["included_columns"] == ["지역", "인구"]
    assert controller.dataModel is not None
    assert controller.dataModel.columnCount() == 2
    assert "2열" in controller.dataViewNotice


def test_confirm_import_rebinds_wide_preview_to_full_selected_columns(tmp_path) -> None:
    from modori.ui.controller import UiController

    columns = [f"v{index}" for index in range(55)]
    data_path = tmp_path / "wide.csv"
    data_path.write_text(
        ",".join(columns) + "\n" + ",".join(str(index) for index in range(55)) + "\n",
        encoding="utf-8",
    )
    controller = UiController()

    assert controller.previewDataFilePath(str(data_path)) is True
    assert len(controller.importColumnRows) == 55
    assert controller.confirmPendingImport() is True

    assert controller.pipeline.current_dataset.df.shape[1] == 55
    assert controller.dataModel is not None
    assert controller.dataModel.columnCount() == 55
