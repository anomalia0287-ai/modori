from __future__ import annotations

from contextlib import closing
import json
from pathlib import Path
import re
import sqlite3

from docx import Document
import pandas as pd
from PySide6.QtCore import Q_ARG, QMetaObject, QObject
import pytest
from scipy.stats import spearmanr

from modori.research_flow import DurableDecision, ResearchFlowState
from modori.research_os import PrimaryAction
from modori.ui.contracts import ReportExportOptions
from modori.ui.controller import UiController
from tests.ui.test_research_flow_controller import ControllableWorker
from tests.ui.test_research_flow_qml import _app, _load_panel


ROOT = Path(__file__).resolve().parents[2]
DEMO_DATA = ROOT / "examples" / "build-week-demo" / "student-study-and-grades.csv"


def _finish(worker: ControllableWorker) -> None:
    worker.finish(worker.execute())
    _app().processEvents()


def _invoke(panel: QObject, method: str, argument: str | None = None) -> bool:
    if argument is None:
        return bool(QMetaObject.invokeMethod(panel, method))
    return bool(
        QMetaObject.invokeMethod(
            panel,
            method,
            Q_ARG("QVariant", argument),
        )
    )


def _word_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(values)


def test_english_guided_qml_real_data_flow_is_numerically_and_ledger_bound(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv(
        "LOCALAPPDATA",
        str((tmp_path / "local-app-data").resolve()),
    )
    data_path = tmp_path / DEMO_DATA.name
    data_path.write_bytes(DEMO_DATA.read_bytes())
    source_frame = pd.read_csv(data_path)
    reference = spearmanr(
        source_frame["weekly_study_time_band"],
        source_frame["final_grade"],
    )
    worker = ControllableWorker()
    host = UiController(worker=worker)

    assert host.setUiLanguage("en") is True
    assert host.chooseMode("guided") is True
    assert host.previewDataFilePath(str(data_path)) is True
    assert "Preview sample: 30 rows · 6 variables" in host.importPreviewText
    assert (
        "Only the first 30 rows are previewed; the full dataset is loaded after confirmation."
        in host.importPreviewText
    )
    assert "Evidence: The first row was recognized as the header." in (
        host.importPreviewText
    )
    assert len(host.importColumnRows) == 6
    assert host.confirmPendingImport() is True
    imported_dataset = host._services.pipeline_ops.current_dataset()
    assert imported_dataset is not None
    assert len(imported_dataset.df) == 649

    assert host.changeVariableMeasure("weekly_study_time_band", "ordinal") is True
    assert host.updateVariableMetadataFieldsFromText(
        "weekly_study_time_band",
        "Weekly study time",
        (
            "1=Under 2 hours; 2=2 to 5 hours; "
            "3=5 to 10 hours; 4=Over 10 hours"
        ),
        "",
    ) is True
    assert host.changeVariableMeasure("final_grade", "scale") is True
    assert host.updateVariableMetadataFieldsFromText(
        "final_grade",
        "Final grade",
        "",
        "",
    ) is True

    flow = host.researchFlow
    engine, panel = _load_panel(flow)
    assert engine._bootstrap.setLanguage("en") is True
    _app().processEvents()
    try:
        assert _invoke(panel, "invokeCommand", "start")
        _finish(worker)
        assert flow.current_view.state is ResearchFlowState.INTAKE_CAUSAL
        assert flow.stateModel["language"] == "en"
        assert _invoke(panel, "invokeCommand", "causal_no")
        assert _invoke(panel, "chooseProfile", "rank_co_movement")

        outcome = panel.findChild(QObject, "researchRoleOutcome")
        predictor = panel.findChild(QObject, "researchRolePredictor")
        assert outcome is not None
        assert predictor is not None
        outcome.setProperty("text", "final_grade")
        predictor.setProperty("text", "weekly_study_time_band")
        _app().processEvents()
        assert _invoke(panel, "submitRoleDraft")
        assert flow.current_view.state is ResearchFlowState.MEANING_REVIEWING
        assert flow._runtime._handle is None
        assert flow._runtime._record is None

        _finish(worker)
        assert flow.current_view.state is ResearchFlowState.VARIABLE_MEANING_REVIEW
        meaning_rows = flow.stateModel["meaningReview"]["rows"]
        assert [row["role"] for row in meaning_rows] == [
            "outcome",
            "focal_predictor",
        ]
        assert [row["label"] for row in meaning_rows] == [
            "Final grade",
            "Weekly study time",
        ]
        assert [row["measure"] for row in meaning_rows] == ["scale", "ordinal"]
        assert meaning_rows[1]["valueLabels"] == [
            {"value": "1", "label": "Under 2 hours"},
            {"value": "2", "label": "2 to 5 hours"},
            {"value": "3", "label": "5 to 10 hours"},
            {"value": "4", "label": "Over 10 hours"},
        ]
        meaning_surface = panel.findChild(QObject, "researchVariableMeaningReview")
        assert meaning_surface is not None
        assert meaning_surface.property("visible") is True
        meaning_digest = flow._views.meaning_review.review_digest
        assert _invoke(panel, "confirmMeaningReview")
        _finish(worker)

        question = panel.findChild(QObject, "researchQuestionCard")
        assert question is not None
        expected_answers = {
            "confirm_dependence": "independent",
            "confirm_weight_use": "variables",
            "confirm_cluster_use": "variables",
        }
        observed_questions: list[str] = []
        while flow.current_view.state is ResearchFlowState.CLARIFY_READY:
            record = flow._runtime._record
            assert record is not None
            assert record.passport.clarify is not None
            question_id = record.passport.clarify.clarification_ref.question_id
            assert question_id in expected_answers
            question_text = str(flow.stateModel["question"]["questionText"])
            assert re.search(r"[가-힣]", question_text) is None
            observed_questions.append(question_id)
            option_id = expected_answers[question_id]
            if option_id == "independent":
                question.setProperty("selectedOptionId", option_id)
            assert QMetaObject.invokeMethod(question, "submitAnswer")
            _finish(worker)

        assert len(observed_questions) == 3
        assert set(observed_questions) == set(expected_answers)
        assert flow.current_view.state is ResearchFlowState.CANDIDATE_READY
        candidate = flow.stateModel["candidate"]
        assert candidate["methodLabel"] == "Spearman rank correlation"
        assert "association" in candidate["claimBoundary"].lower()

        record = flow._runtime._record
        assert isinstance(record, DurableDecision)
        assert record.action is PrimaryAction.RECOMMEND_LOCAL
        assert record.passport_digest == flow.current_view.decision_identity_digest
        role_facts = record.request.estimand.target_roles
        assert all(
            f"variable-meaning-review:v1:{meaning_digest}"
            in role.variable_ids.provenance_refs
            for role in role_facts
        )
        handle = flow._runtime._handle
        assert handle is not None
        ledger_uri = f"{Path(handle.ledger_path).as_uri()}?mode=ro"
        with closing(sqlite3.connect(ledger_uri, uri=True)) as connection:
            event_count = connection.execute(
                "SELECT COUNT(*) FROM ledger_events"
            ).fetchone()[0]
            passport_count = connection.execute(
                "SELECT COUNT(*) FROM ledger_artifacts "
                "WHERE artifact_kind='analysis_passport'"
            ).fetchone()[0]
        assert (event_count, passport_count) == (8, 4)

        submissions_before_prepare = len(worker.submissions)
        assert _invoke(panel, "invokeCommand", "prepare")
        _finish(worker)
        assert flow.current_view.state is ResearchFlowState.PREPARE_REVIEW
        assert len(worker.submissions) == submissions_before_prepare + 1
        preparation = flow._views.preparation
        assert preparation is not None
        assert preparation.step_type == "stats.correlation"
        assert json.loads(preparation.canonical_step_params) == {
            "method": "spearman",
            "missing_policy": "pairwise",
            "p_adjust": "none",
            "pairs": [["final_grade", "weekly_study_time_band"]],
            "schema_version": 1,
        }
        assert preparation.experimental is True
        assert preparation.requires_explicit_configure_confirm_run is True

        submissions_before_confirm = len(worker.submissions)
        assert _invoke(panel, "confirmReview")
        assert flow.current_view.state is ResearchFlowState.CONFIRMED
        assert len(worker.submissions) == submissions_before_confirm
        assert host.resultsModel == []

        early_export = host.exportReport(
            ReportExportOptions(language="en", include_figures=False)
        )
        assert early_export.ok is False
        submissions_before_run = len(worker.submissions)
        assert host.rerun().ok is True
        _finish(worker)
        assert len(worker.submissions) == submissions_before_run + 1
        assert len(host.resultsModel) == 1
        display = host.resultsModel[0]
        assert display.result_id == "correlation"
        assert display.kind == "correlation"
        assert display.prose_en.startswith(
            "This result summarizes 1 correlation pair "
        )

        raw = host._services.pipeline_ops.analysis_objects()["correlation"]
        pair = raw.pairs[0]
        assert pair.method == "spearman"
        assert pair.x_label == "Final grade"
        assert pair.y_label == "Weekly study time"
        assert pair.coefficient == pytest.approx(float(reference.statistic), abs=1e-12)
        assert pair.p_value == pytest.approx(float(reference.pvalue), rel=1e-12)
        assert pair.n == 649
        assert pair.excluded_n == 0

        table = display.tables[0]
        assert table.rows_en is not None
        table_row = {
            column.label: value
            for column, value in zip(table.columns, table.rows_en[0], strict=True)
        }
        assert table_row == {
            "x": "Final grade",
            "y": "Weekly study time",
            "method": "spearman",
            "statistic": "rho",
            "coefficient": "0.275",
            "p_value": "0.000",
            "p_adjusted": "",
            "n": "649",
            "excluded_n": "0",
            "warnings": (
                "Spearman correlation handled tied ranks using SciPy's spearmanr "
                "policy. P-values for data with many ties can differ across "
                "statistical software."
            ),
        }
        assert re.search(r"[가-힣]", host.resultTableText) is None

        exported = host.exportReport(
            ReportExportOptions(
                language="en",
                include_descriptives=False,
                include_reliability=False,
                include_comparison=False,
                include_association=True,
                include_group_models=False,
                include_dimension_reduction=False,
                include_regression=False,
                include_figures=False,
            )
        )
        assert exported.ok is True
        report_path = Path(exported.result_ids[0])
        assert report_path.is_file()
        report_text = _word_text(report_path)
        assert "Research OS" in report_text
        assert "Spearman" in report_text
        assert "Final grade" in report_text
        assert "Weekly study time" in report_text
        assert "0.275" in report_text
        assert "recommendation validity" in report_text
        assert re.search(r"[가-힣]", report_text) is None

        assert host.updateVariableMetadata(
            "weekly_study_time_band",
            {"measure": "scale"},
        ).ok
        _app().processEvents()
        assert flow.current_view.state is ResearchFlowState.REPLAN_REQUIRED
        blocked = host.rerun()
        assert blocked.ok is False
        assert blocked.error_code == "experimental_confirmation_required"
        assert _invoke(panel, "invokeCommand", "replan")
        _finish(worker)
        assert flow.current_view.state is ResearchFlowState.INTAKE_CAUSAL
    finally:
        panel.deleteLater()
        _app().processEvents()
        del engine
