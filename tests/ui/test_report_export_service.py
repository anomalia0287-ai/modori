from __future__ import annotations

import hashlib
import os
from pathlib import Path

from docx import Document

from modori.steps.reporting import RESEARCH_OS_SELECTION_DISCLOSURE
import modori.ui.report_export as report_export_module
from modori.ui.contracts import ReportExportOptions
from modori.ui.report_export import (
    EXPERIMENTAL_DISCLOSURE_EN,
    EXPERIMENTAL_DISCLOSURE_KO,
    RESEARCH_OS_DISCLOSURE_EN,
    RESEARCH_OS_DISCLOSURE_KO,
    ReportExportService,
)


def test_report_options_preserve_distinct_research_os_origin() -> None:
    options = ReportExportOptions(selection_provenance="research_os_assisted")

    assert options.selection_provenance == "research_os_assisted"
    assert options.selection_origin == "research_os_assisted"


def test_report_export_service_rejects_missing_pipeline() -> None:
    result = ReportExportService().export(
        pipeline=None,
        options=ReportExportOptions(),
        exporter=lambda pipeline, options: "unused.docx",
        pipeline_version=7,
    )

    assert result.ok is False
    assert result.error_code == "no_pipeline"
    assert result.pipeline_version == 7


def test_report_export_service_requires_created_output_file(tmp_path) -> None:
    missing_path = tmp_path / "missing.docx"

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(),
        exporter=lambda pipeline, options: missing_path,
        pipeline_version=3,
    )

    assert result.ok is False
    assert result.error_code == "engine_error"
    assert result.message_ko == "보고서 파일이 생성되지 않았습니다."


def test_report_export_service_returns_success_with_existing_file(tmp_path) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"docx")

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(language="ko"),
        exporter=lambda pipeline, options: output_path,
        pipeline_version=11,
    )

    assert result.ok is True
    assert result.result_ids == [str(output_path)]
    assert result.pipeline_version == 11


def test_report_export_adds_localized_experimental_selection_disclosure(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"

    def exporter(pipeline, options):
        document = Document()
        document.add_paragraph("본문" if options.language == "ko" else "Body")
        document.save(output_path)
        return output_path

    service = ReportExportService()

    korean = service.export(
        pipeline=object(),
        options=ReportExportOptions(
            language="ko",
            selection_provenance="experimental_candidate_assisted",
        ),
        exporter=exporter,
        pipeline_version=1,
    )

    assert korean.ok is True
    assert EXPERIMENTAL_DISCLOSURE_KO in [
        paragraph.text for paragraph in Document(output_path).paragraphs
    ]

    english = service.export(
        pipeline=object(),
        options=ReportExportOptions(
            language="en",
            selection_provenance="experimental_candidate_assisted",
        ),
        exporter=exporter,
        pipeline_version=1,
    )

    paragraphs = [paragraph.text for paragraph in Document(output_path).paragraphs]
    assert english.ok is True
    assert EXPERIMENTAL_DISCLOSURE_EN in paragraphs
    assert EXPERIMENTAL_DISCLOSURE_KO not in paragraphs


def test_manual_report_export_has_no_experimental_selection_disclosure(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"

    def exporter(pipeline, options):
        document = Document()
        document.add_paragraph("본문")
        document.save(output_path)
        return output_path

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(language="ko", selection_provenance="manual"),
        exporter=exporter,
        pipeline_version=2,
    )

    paragraphs = [paragraph.text for paragraph in Document(output_path).paragraphs]
    assert result.ok is True
    assert EXPERIMENTAL_DISCLOSURE_KO not in paragraphs
    assert EXPERIMENTAL_DISCLOSURE_EN not in paragraphs


def test_report_export_uses_distinct_research_os_disclosure(tmp_path) -> None:
    output_path = tmp_path / "report.docx"

    def exporter(pipeline, options):
        document = Document()
        document.add_paragraph("Body")
        document.save(output_path)
        return output_path

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(
            language="en",
            selection_provenance="research_os_assisted",
        ),
        exporter=exporter,
        pipeline_version=5,
    )

    paragraphs = [paragraph.text for paragraph in Document(output_path).paragraphs]
    assert result.ok is True
    assert RESEARCH_OS_DISCLOSURE_EN in paragraphs
    assert RESEARCH_OS_DISCLOSURE_KO not in paragraphs
    assert EXPERIMENTAL_DISCLOSURE_EN not in paragraphs


def test_report_export_replaces_calculation_disclosure_instead_of_duplicating_it(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"

    def exporter(pipeline, options):
        document = Document()
        document.add_paragraph(RESEARCH_OS_SELECTION_DISCLOSURE["en"])
        document.add_paragraph("Body")
        document.save(output_path)
        return output_path

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(
            language="en",
            selection_provenance="research_os_assisted",
        ),
        exporter=exporter,
        pipeline_version=5,
    )

    paragraphs = [paragraph.text for paragraph in Document(output_path).paragraphs]
    assert result.ok is True
    assert RESEARCH_OS_SELECTION_DISCLOSURE["en"] not in paragraphs
    assert paragraphs.count(RESEARCH_OS_DISCLOSURE_EN) == 1
    assert paragraphs.count("Body") == 1


def test_experimental_report_export_fails_closed_when_docx_cannot_hold_disclosure(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"not-a-docx")

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(
            selection_provenance="experimental_candidate_assisted"
        ),
        exporter=lambda pipeline, options: output_path,
        pipeline_version=4,
    )

    assert result.ok is False
    assert result.error_code == "engine_error"
    assert output_path.exists() is False
    assert list(tmp_path.iterdir()) == []


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_report_export_rejects_existing_destination_before_exporter_runs(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"original-report")
    before = _sha256(output_path)
    calls: list[bool] = []

    def exporter(pipeline, options):
        calls.append(True)
        output_path.write_bytes(b"replacement")
        return output_path

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(),
        exporter=exporter,
        pipeline_version=8,
        expected_output_path=output_path,
    )

    assert result.ok is False
    assert result.error_code == "report_destination_exists"
    assert result.result_ids == [str(output_path.resolve())]
    assert calls == []
    assert _sha256(output_path) == before


def test_approved_replacement_restores_original_when_exporter_fails(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"original-report")
    before = _sha256(output_path)

    def failing_exporter(pipeline, options):
        output_path.write_bytes(b"partial-new-report")
        raise RuntimeError("failed after mutation")

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(replace_existing=True),
        exporter=failing_exporter,
        pipeline_version=9,
        expected_output_path=output_path,
    )

    assert result.ok is False
    assert result.error_code == "engine_error"
    assert _sha256(output_path) == before
    assert not list(tmp_path.glob(".*.backup"))


def test_approved_replacement_retries_transient_windows_restore_lock(
    tmp_path,
    monkeypatch,
) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"original-report")
    before = _sha256(output_path)
    real_replace = os.replace
    restore_attempts = 0

    def transient_replace(source, destination):
        nonlocal restore_attempts
        if Path(source).suffix == ".backup" and Path(destination) == output_path:
            restore_attempts += 1
            if restore_attempts == 1:
                raise PermissionError("transient Windows file lock")
        real_replace(source, destination)

    def failing_exporter(pipeline, options):
        output_path.write_bytes(b"partial-new-report")
        raise RuntimeError("failed after mutation")

    monkeypatch.setattr(report_export_module.os, "replace", transient_replace)
    monkeypatch.setattr(report_export_module, "sleep", lambda _delay: None)

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(replace_existing=True),
        exporter=failing_exporter,
        pipeline_version=9,
        expected_output_path=output_path,
    )

    assert result.ok is False
    assert result.error_code == "engine_error"
    assert restore_attempts == 2
    assert _sha256(output_path) == before
    assert not list(tmp_path.glob(".*.backup"))


def test_persistent_restore_lock_retains_backup_and_returns_recovery_path(
    tmp_path,
    monkeypatch,
) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"original-report")
    before = _sha256(output_path)

    def locked_replace(source, destination):
        if Path(source).suffix == ".backup" and Path(destination) == output_path:
            raise PermissionError("persistent Windows file lock")
        os.replace(source, destination)

    def failing_exporter(pipeline, options):
        output_path.write_bytes(b"partial-new-report")
        raise RuntimeError("failed after mutation")

    monkeypatch.setattr(report_export_module.os, "replace", locked_replace)
    monkeypatch.setattr(report_export_module, "sleep", lambda _delay: None)

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(replace_existing=True),
        exporter=failing_exporter,
        pipeline_version=9,
        expected_output_path=output_path,
    )

    assert result.ok is False
    assert result.error_code == "report_restore_failed"
    assert len(result.result_ids) == 1
    backup_path = Path(result.result_ids[0])
    assert backup_path.suffix == ".backup"
    assert _sha256(backup_path) == before
    assert output_path.read_bytes() == b"partial-new-report"


def test_approved_replacement_commits_valid_docx_and_discards_backup(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"original-report")

    def exporter(pipeline, options):
        document = Document()
        document.add_paragraph("replacement")
        document.save(output_path)
        return output_path

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(replace_existing=True),
        exporter=exporter,
        pipeline_version=10,
        expected_output_path=output_path,
    )

    assert result.ok is True
    assert [paragraph.text for paragraph in Document(output_path).paragraphs] == [
        "replacement"
    ]
    assert not list(tmp_path.glob(".*.backup"))


def test_approved_replacement_restores_original_when_disclosure_fails(
    tmp_path,
) -> None:
    output_path = tmp_path / "report.docx"
    output_path.write_bytes(b"original-report")
    before = _sha256(output_path)

    def invalid_exporter(pipeline, options):
        output_path.write_bytes(b"not-a-docx")
        return output_path

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(
            selection_provenance="experimental_candidate_assisted",
            replace_existing=True,
        ),
        exporter=invalid_exporter,
        pipeline_version=11,
        expected_output_path=output_path,
    )

    assert result.ok is False
    assert result.error_code == "engine_error"
    assert _sha256(output_path) == before
    assert not list(tmp_path.glob(".*.backup"))


def test_approved_replacement_rejects_unexpected_export_path(tmp_path) -> None:
    expected_path = tmp_path / "report.docx"
    unexpected_path = tmp_path / "other.docx"
    expected_path.write_bytes(b"original-report")
    before = _sha256(expected_path)

    def exporter(pipeline, options):
        Document().save(unexpected_path)
        return unexpected_path

    result = ReportExportService().export(
        pipeline=object(),
        options=ReportExportOptions(replace_existing=True),
        exporter=exporter,
        pipeline_version=12,
        expected_output_path=expected_path,
    )

    assert result.ok is False
    assert result.error_code == "engine_error"
    assert _sha256(expected_path) == before
    assert unexpected_path.exists() is True
    assert not list(tmp_path.glob(".*.backup"))
