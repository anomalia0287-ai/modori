from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from modori.recommendation_benchmark import BenchmarkContractError, PredictionRecord
from modori.recommendation_benchmark_io import read_jsonl
from scripts.build_recommendation_pilot import (
    build_pilot_pack,
    pilot_case_definitions,
    validate_pilot_pack,
)


def _mapping(value: object) -> dict[str, object]:
    assert isinstance(value, dict)
    return value


def _file_hashes(root: Path) -> dict[str, str]:
    return {
        path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def test_pilot_definitions_cover_20_unique_unlabeled_cases() -> None:
    cases = pilot_case_definitions()

    assert len(cases) == 20
    assert len({case.case_id for case in cases}) == 20
    assert {case.language for case in cases} == {"ko", "mixed"}
    assert all(case.sensitivity == "synthetic_no_real_pii" for case in cases)
    assert all(case.generation_seed > 0 for case in cases)
    assert sum(case.data_file == "data/pilot-019-020-shared.csv" for case in cases) == 2
    serialized = json.dumps(
        [case.to_mapping() for case in cases],
        ensure_ascii=False,
        sort_keys=True,
    )
    for forbidden in (
        "gold",
        "action_class",
        "acceptable_recommendations",
        "required_clarification_facts",
    ):
        assert forbidden not in serialized


def test_generated_pilot_pack_has_valid_manifest_checksums_and_blank_workbooks(
    tmp_path: Path,
) -> None:
    root = tmp_path / "recommendation_benchmark"

    build_pilot_pack(root)
    report = validate_pilot_pack(root)

    assert report.case_count == 20
    assert report.unique_data_file_count == 19
    assert report.manifest_count == 20
    manifest = read_jsonl(root / "manifest.jsonl", _mapping)
    cases = read_jsonl(root / "public" / "pilot" / "cases.jsonl", _mapping)
    predictions = read_jsonl(
        root / "public" / "pilot" / "baseline-a-predictions.jsonl",
        PredictionRecord.from_mapping,
    )
    metadata = json.loads(
        (root / "public" / "pilot" / "baseline-a-metadata.json").read_text(
            encoding="utf-8"
        )
    )
    assert len(manifest) == len(cases) == 20
    assert len(predictions) == 20
    assert metadata["variant"] == "A"
    assert metadata["adapter_version"] == "current-recommendation-service-v1"
    assert metadata["prediction_count"] == 20
    assert metadata["prediction_checksum"].startswith("sha256:")
    assert metadata["case_set_checksum"].startswith("sha256:")
    assert metadata["source_fingerprint"].startswith("sha256:")
    assert "score" not in metadata
    assert {(prediction.case_id, prediction.evidence_stage) for prediction in predictions} == {
        (str(case["case_id"]), str(case["evidence_stage"])) for case in cases
    }
    assert all(record["split_role"] == "economics_pilot" for record in manifest)
    assert all(record["license"] == "LicenseRef-Modori-Synthetic-Benchmark-1.0" for record in manifest)
    assert all(record["sensitivity"] == "synthetic_no_real_pii" for record in manifest)
    for record in manifest:
        data_path = root / str(record["data_file"])
        digest = f"sha256:{hashlib.sha256(data_path.read_bytes()).hexdigest()}"
        assert record["checksum"] == digest

    for workbook_name in ("reviewer-a.xlsx", "reviewer-b.xlsx", "adjudication.xlsx"):
        workbook = load_workbook(root / "public" / "pilot" / workbook_name, data_only=False)
        assert workbook["Instructions"]["B2"].value is None
        reviews = workbook["Case Reviews"]
        assert all(
            reviews.cell(row=row, column=column).value is None
            for row in range(2, 22)
            for column in range(3, 7)
        )


def test_pilot_pack_generation_is_deterministic_for_text_and_data_files(
    tmp_path: Path,
) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"

    build_pilot_pack(first)
    build_pilot_pack(second)

    assert _file_hashes(first) == _file_hashes(second)


def test_pilot_pack_validation_detects_data_tampering(tmp_path: Path) -> None:
    root = tmp_path / "recommendation_benchmark"
    build_pilot_pack(root)
    data_path = root / "public" / "pilot" / "data" / "pilot-001-descriptives.csv"
    data_path.write_text(
        data_path.read_text(encoding="utf-8") + "tampered,row\n",
        encoding="utf-8",
    )

    with pytest.raises(BenchmarkContractError, match="checksum mismatch"):
        validate_pilot_pack(root)


def test_pilot_pack_validation_detects_incomplete_baseline_predictions(
    tmp_path: Path,
) -> None:
    root = tmp_path / "recommendation_benchmark"
    build_pilot_pack(root)
    prediction_path = root / "public" / "pilot" / "baseline-a-predictions.jsonl"
    lines = prediction_path.read_text(encoding="utf-8").splitlines()
    prediction_path.write_text("\n".join(lines[:-1]) + "\n", encoding="utf-8")

    with pytest.raises(BenchmarkContractError, match="baseline prediction checksum mismatch"):
        validate_pilot_pack(root)


def test_pilot_pack_validation_detects_valid_shape_baseline_tampering(
    tmp_path: Path,
) -> None:
    root = tmp_path / "recommendation_benchmark"
    build_pilot_pack(root)
    prediction_path = root / "public" / "pilot" / "baseline-a-predictions.jsonl"
    rows = [json.loads(line) for line in prediction_path.read_text(encoding="utf-8").splitlines()]
    rows[0]["level"] = "candidate"
    prediction_path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in rows) + "\n",
        encoding="utf-8",
    )

    with pytest.raises(BenchmarkContractError, match="baseline prediction checksum mismatch"):
        validate_pilot_pack(root)


def test_checked_in_pilot_pack_is_complete_and_valid() -> None:
    root = Path("tests/fixtures/recommendation_benchmark")

    report = validate_pilot_pack(root)

    assert report.case_count == 20
    assert report.unique_data_file_count == 19
