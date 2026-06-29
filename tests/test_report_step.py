from pathlib import Path
import struct
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pandas as pd
import pingouin as pg
import pytest
from docx import Document
from matplotlib import image as mpimg
from scipy import stats

from modori.core import Dataset, Measure, Pipeline, PipelineContext, Variable
from modori.results import ChartSpec, ReliabilityResult, ReportResult
from modori.steps import (
    CompareGroupsStep,
    ComposeScaleStep,
    ImportStep,
    RecodeReverseStep,
    ReliabilityStep,
    ReportStep,
)


def report_dataset() -> Dataset:
    comparison_group1 = [10, 11, 9, 10, 12, 11, 10, 9, 11, 10]
    comparison_group2 = [12, 13, 11, 12, 14, 13, 12, 11, 13, 12]
    frame = pd.DataFrame(
        {
            "q1": [1, 2, 3, 4, 5, 5, 4, 3, 1, 2, 3, 4, 5, 5, 4, 3, 1, 2, 3, 4],
            "q2": [1, 2, 3, 4, 4, 5, 4, 3, 1, 2, 3, 4, 4, 5, 4, 3, 1, 2, 3, 4],
            "q3": [2, 2, 3, 3, 5, 5, 4, 4, 2, 2, 3, 3, 5, 5, 4, 4, 2, 2, 3, 3],
            "q4": [1, 3, 3, 4, 5, 4, 4, 3, 1, 3, 3, 4, 5, 4, 4, 3, 1, 3, 3, 4],
            "job_sat": comparison_group1 + comparison_group2,
            "group": [1.0] * 10 + [2.0] * 10,
        }
    )
    variables = {
        column: Variable(
            name=column,
            label=column,
            measure=Measure.ORDINAL,
            value_labels={},
            missing_values=[],
            dtype="float",
            origin_step_id=None,
        )
        for column in ["q1", "q2", "q3", "q4"]
    }
    variables["job_sat"] = Variable(
        name="job_sat",
        label="직무만족",
        measure=Measure.SCALE,
        value_labels={},
        missing_values=[],
        dtype="float",
        origin_step_id="compose",
    )
    variables["group"] = Variable(
        name="group",
        label="집단",
        measure=Measure.NOMINAL,
        value_labels={1.0: "통제집단", 2.0: "처치집단"},
        missing_values=[],
        dtype="float",
        origin_step_id=None,
    )
    return Dataset(df=frame, variables=variables)


def report_step(output_dir: Path) -> ReportStep:
    return ReportStep(
        id="report",
        title="APA report",
        params={
            "include": ["reliability:job_sat", "comparison:job_sat:group"],
            "output_dir": str(output_dir),
            "filename": "report.docx",
            "language": "ko",
        },
    )


def reliability_result_with_chart(chart_type: str) -> ReliabilityResult:
    return ReliabilityResult(
        scale_name="job_sat",
        n_items=4,
        n_cases=20,
        cronbach_alpha=0.90,
        alpha_ci=(0.80, 0.95),
        mcdonald_omega=0.91,
        item_total_corr={"q1": 0.7, "q2": 0.8},
        alpha_if_deleted={"q1": 0.85, "q2": 0.86},
        apa_template_id="reliability.v1",
        chart_spec=ChartSpec(
            type=chart_type,
            title="Reliability chart",
            data={"values": {"q1": 0.7, "q2": 0.8}},
            x_label="x",
            y_label="y",
        ),
    )


def _png_pixels_per_meter(path: Path) -> tuple[int, int, int]:
    data = path.read_bytes()
    assert data.startswith(b"\x89PNG\r\n\x1a\n")
    index = 8
    while index < len(data):
        length = struct.unpack(">I", data[index:index + 4])[0]
        chunk_type = data[index + 4:index + 8]
        chunk_data = data[index + 8:index + 8 + length]
        if chunk_type == b"pHYs":
            x_ppm, y_ppm, unit = struct.unpack(">IIB", chunk_data)
            return x_ppm, y_ppm, unit
        index += 12 + length
    raise AssertionError(f"PNG pHYs chunk not found: {path}")


def test_report_step_generates_korean_apa_prose_figures_and_docx(tmp_path) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        CompareGroupsStep(
            id="compare",
            title="Compare groups",
            params={
                "dv": "job_sat",
                "group": "group",
                "routing_policy": {"preset": "modern"},
            },
        )
    )
    pipeline.add(report_step(tmp_path))

    pipeline.recompute(dirty_from=None)

    report = pipeline.analysis_objects["report"]
    assert isinstance(report, ReportResult)
    assert Path(report.docx_path).exists()
    assert any("Cronbach's \u03b1 = .97" in sentence for sentence in report.prose)
    assert any("독립표본 t검정 결과" in sentence for sentence in report.prose)
    assert any("직무만족 점수 차이" in sentence for sentence in report.prose)
    assert any("p < .001" in sentence for sentence in report.prose)
    assert any("Cohen's d = -2.11" in sentence for sentence in report.prose)
    assert "reliability:job_sat" in report.tables
    assert report.tables["reliability:job_sat"][0]["item"] == "q1"
    assert "comparison:job_sat:group" in report.tables
    assert report.tables["comparison:job_sat:group"][0]["test"] == "student_t"
    assert sorted(Path(path).suffix for paths in report.figure_paths.values() for path in paths) == [
        ".eps",
        ".eps",
        ".png",
        ".png",
        ".svg",
        ".svg",
    ]
    assert all(Path(path).exists() for paths in report.figure_paths.values() for path in paths)
    assert all(Path(path).stat().st_size > 0 for paths in report.figure_paths.values() for path in paths)
    png_paths = [
        Path(path)
        for paths in report.figure_paths.values()
        for path in paths
        if path.endswith(".png")
    ]
    svg_paths = [
        Path(path)
        for paths in report.figure_paths.values()
        for path in paths
        if path.endswith(".svg")
    ]
    eps_paths = [
        Path(path)
        for paths in report.figure_paths.values()
        for path in paths
        if path.endswith(".eps")
    ]
    png_images = [mpimg.imread(path) for path in png_paths]
    assert all(image.size > 0 for image in png_images)
    assert all(image.shape[0] >= 100 and image.shape[1] >= 100 for image in png_images)
    assert all(_png_pixels_per_meter(path) == (11811, 11811, 1) for path in png_paths)
    assert all(ET.parse(path).getroot().tag.endswith("svg") for path in svg_paths)
    assert all(path.read_text(encoding="utf-8", errors="ignore").startswith("%!PS-Adobe") for path in eps_paths)
    document = Document(report.docx_path)
    assert len(document.inline_shapes) == len(png_paths)
    assert all(shape.width > 0 and shape.height > 0 for shape in document.inline_shapes)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert all(sentence in paragraphs for sentence in report.prose)
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    assert "student_t" in table_cells
    assert "cohen_d" in table_cells
    with ZipFile(report.docx_path) as docx_archive:
        embedded_media = [
            name
            for name in docx_archive.namelist()
            if name.startswith("word/media/")
        ]
    assert len(embedded_media) == len(png_paths)


def test_report_step_recomputes_after_import_data_changes(tmp_path) -> None:
    csv_path = tmp_path / "survey.csv"
    csv_path.write_text(
        "job_sat,group\n"
        "10,1\n11,1\n9,1\n10,1\n12,1\n11,1\n10,1\n9,1\n11,1\n10,1\n"
        "12,2\n13,2\n11,2\n12,2\n14,2\n13,2\n12,2\n11,2\n13,2\n12,2\n",
        encoding="utf-8",
    )
    pipeline = Pipeline(Dataset.empty())
    pipeline.add(
        ImportStep(
            id="import",
            title="Import CSV",
            params={"path": str(csv_path), "file_type": "csv"},
        )
    )
    pipeline.add(
        CompareGroupsStep(
            id="compare",
            title="Compare groups",
            params={
                "dv": "job_sat",
                "group": "group",
                "routing_policy": {"preset": "modern"},
            },
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["comparison:job_sat:group"],
                "output_dir": str(tmp_path),
                "filename": "rerun.docx",
                "language": "ko",
            },
        )
    )
    pipeline.recompute(dirty_from=None)
    first_prose = pipeline.analysis_objects["report"].prose

    csv_path.write_text(
        "job_sat,group\n"
        "10,1\n11,1\n9,1\n10,1\n12,1\n11,1\n10,1\n9,1\n11,1\n10,1\n"
        "10,2\n11,2\n9,2\n10,2\n12,2\n11,2\n10,2\n9,2\n11,2\n10,2\n",
        encoding="utf-8",
    )
    pipeline.recompute(dirty_from="import")
    second_prose = pipeline.analysis_objects["report"].prose

    assert first_prose != second_prose
    assert any("통계적으로 유의하였다" in sentence for sentence in first_prose)
    assert any("통계적으로 유의하지 않았다" in sentence for sentence in second_prose)


def test_public_reference_data_runs_from_csv_import_to_report(tmp_path) -> None:
    cronbach_wide = (
        pg.read_dataset("cronbach_alpha")
        .pivot(index="Subj", columns="Items", values="Scores")
        .reset_index(drop=True)
        .rename(columns=lambda column: str(column))
    )
    august_scores = (
        pg.read_dataset("mixed_anova")
        .loc[lambda frame: frame["Time"] == "August", ["Scores", "Group"]]
        .reset_index(drop=True)
    )
    repeated_items = pd.concat([cronbach_wide] * 4, ignore_index=True).iloc[
        : len(august_scores)
    ]
    frame = pd.concat([repeated_items, august_scores], axis=1)
    csv_path = tmp_path / "public-reference.csv"
    frame.to_csv(csv_path, index=False)

    pipeline = Pipeline(Dataset.empty())
    pipeline.add(
        ImportStep(
            id="import",
            title="Import public reference CSV",
            params={"path": str(csv_path), "file_type": "csv"},
        )
    )
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={
                "items": list(cronbach_wide.columns),
                "scale_name": "cronbach_alpha_reference",
            },
        )
    )
    pipeline.add(
        CompareGroupsStep(
            id="compare",
            title="Compare groups",
            params={
                "dv": "Scores",
                "group": "Group",
                "routing_policy": {"preset": "modern"},
            },
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": [
                    "reliability:cronbach_alpha_reference",
                    "comparison:Scores:Group",
                ],
                "output_dir": str(tmp_path),
                "filename": "public-reference-report.docx",
                "language": "ko",
            },
        )
    )

    pipeline.recompute(dirty_from=None)

    reliability = pipeline.analysis_objects["reliability:cronbach_alpha_reference"]
    comparison = pipeline.analysis_objects["comparison:Scores:Group"]
    report = pipeline.analysis_objects["report"]
    assert reliability.cronbach_alpha == pytest.approx(
        0.5917188485995826,
        abs=1e-12,
    )
    assert reliability.alpha_ci == pytest.approx((0.418, 0.730), abs=0.001)
    assert comparison.test_name == "student_t"
    assert comparison.statistic == pytest.approx(0.31602196533393784, abs=1e-12)
    assert comparison.p_value == pytest.approx(0.7531203054939072, abs=1e-12)
    assert Path(report.docx_path).exists()
    assert any("Cronbach's \u03b1 = .59" in sentence for sentence in report.prose)
    assert any("통계적으로 유의하지 않았다" in sentence for sentence in report.prose)


def test_psych_bfi_public_survey_runs_from_csv_import_to_report(tmp_path) -> None:
    source_path = Path(__file__).parent / "fixtures" / "psych_bfi.csv"
    source_frame = pd.read_csv(source_path)
    keyed_frame = source_frame.copy()
    keyed_frame["A1_R"] = 7 - keyed_frame["A1"]
    agree_items = ["A1_R", "A2", "A3", "A4", "A5"]
    valid_agree = keyed_frame[agree_items].notna().sum(axis=1) >= 4
    keyed_frame["agree"] = keyed_frame[agree_items].mean(axis=1, skipna=True).where(
        valid_agree
    )

    complete_reliability = keyed_frame[agree_items].dropna()
    covariance = complete_reliability.cov()
    alpha_manual = (len(agree_items) / (len(agree_items) - 1)) * (
        1 - covariance.to_numpy().diagonal().sum() / covariance.to_numpy().sum()
    )
    comparison_frame = keyed_frame[["agree", "gender"]].dropna()
    gender_1 = comparison_frame.loc[comparison_frame["gender"] == 1, "agree"]
    gender_2 = comparison_frame.loc[comparison_frame["gender"] == 2, "agree"]
    welch = stats.ttest_ind(gender_1, gender_2, equal_var=False)
    first_variance = gender_1.var(ddof=1)
    second_variance = gender_2.var(ddof=1)
    first_term = first_variance / len(gender_1)
    second_term = second_variance / len(gender_2)
    welch_df = (first_term + second_term) ** 2 / (
        first_term**2 / (len(gender_1) - 1)
        + second_term**2 / (len(gender_2) - 1)
    )
    pooled_sd = (
        (
            (len(gender_1) - 1) * first_variance
            + (len(gender_2) - 1) * second_variance
        )
        / (len(gender_1) + len(gender_2) - 2)
    ) ** 0.5
    signed_cohen_d = (gender_1.mean() - gender_2.mean()) / pooled_sd

    pipeline = Pipeline(Dataset.empty())
    pipeline.add(
        ImportStep(
            id="import",
            title="Import psych bfi CSV",
            params={"path": str(source_path), "file_type": "csv"},
        )
    )
    pipeline.add(
        RecodeReverseStep(
            id="reverse",
            title="Reverse keyed BFI items",
            params={"columns": ["A1"], "scale_min": 1, "scale_max": 6},
        )
    )
    pipeline.add(
        ComposeScaleStep(
            id="compose",
            title="Compose agreeableness",
            params={
                "items": agree_items,
                "method": "mean",
                "name": "agree",
                "missing_policy": {"preset": "survey", "min_valid": 0.8},
            },
        )
    )
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": agree_items, "scale_name": "agree"},
        )
    )
    pipeline.add(
        CompareGroupsStep(
            id="compare",
            title="Compare gender groups",
            params={
                "dv": "agree",
                "group": "gender",
                "routing_policy": {"preset": "modern"},
            },
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:agree", "comparison:agree:gender"],
                "output_dir": str(tmp_path),
                "filename": "psych-bfi-report.docx",
                "language": "ko",
            },
        )
    )

    pipeline.recompute(dirty_from=None)

    reliability = pipeline.analysis_objects["reliability:agree"]
    comparison = pipeline.analysis_objects["comparison:agree:gender"]
    report = pipeline.analysis_objects["report"]
    assert pipeline.current_dataset.df.loc[0, "A1_R"] == pytest.approx(5.0)
    assert pipeline.current_dataset.df["agree"].notna().sum() == 2790
    assert reliability.n_cases == 2709
    assert reliability.cronbach_alpha == pytest.approx(alpha_manual, abs=1e-12)
    assert reliability.cronbach_alpha == pytest.approx(0.7037558943748362, abs=1e-12)
    assert comparison.test_name == "welch_t"
    assert comparison.route_reason == "unequal variance -> Welch correction"
    assert comparison.statistic == pytest.approx(welch.statistic, abs=1e-12)
    assert comparison.df == pytest.approx(welch_df, abs=1e-9)
    assert comparison.effect_value == pytest.approx(signed_cohen_d, abs=1e-12)
    assert comparison.p_value < 1e-20
    assert Path(report.docx_path).exists()
    assert any("Welch 보정" in sentence for sentence in report.prose)
    assert any("Cronbach's \u03b1 = .70" in sentence for sentence in report.prose)


def test_report_step_supports_english_secondary_prose(tmp_path) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(tmp_path),
                "filename": "report-en.docx",
                "language": "en",
            },
        )
    )

    pipeline.recompute(dirty_from=None)

    assert pipeline.analysis_objects["report"].prose == [
        "The job_sat scale showed excellent internal consistency "
        "(Cronbach's \u03b1 = .97, McDonald's \u03c9 = .97)."
    ]


def test_report_step_rejects_output_filename_path_traversal(tmp_path) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(tmp_path),
                "filename": "..\\escape.docx",
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="Report filename must not contain path separators"):
        pipeline.recompute(dirty_from=None)


@pytest.mark.parametrize("filename", ["../escape.docx", "/escape.docx", "C:/escape.docx"])
def test_report_step_rejects_absolute_or_forward_slash_filenames(
    tmp_path,
    filename,
) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(tmp_path),
                "filename": filename,
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="Report filename must not contain path separators"):
        pipeline.recompute(dirty_from=None)


def test_report_step_rejects_output_dir_that_is_an_existing_file(tmp_path) -> None:
    output_file = tmp_path / "not-a-directory"
    output_file.write_text("occupied", encoding="utf-8")
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(output_file),
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="Report output_dir exists and is not a directory"):
        pipeline.recompute(dirty_from=None)


def test_report_step_rejects_existing_docx_symlink_destination(tmp_path) -> None:
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    outside = tmp_path / "outside.docx"
    outside.write_text("outside", encoding="utf-8")
    docx_link = output_dir / "report.docx"
    try:
        docx_link.symlink_to(outside)
    except (NotImplementedError, OSError) as exc:
        pytest.skip(f"symlink creation unavailable: {exc}")
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(output_dir),
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="Report output path must not be a symbolic link"):
        pipeline.recompute(dirty_from=None)

    assert outside.read_text(encoding="utf-8") == "outside"


def test_report_step_rejects_symlink_output_dir(tmp_path) -> None:
    output_dir = tmp_path / "out"
    outside_dir = tmp_path / "outside"
    outside_dir.mkdir()
    try:
        output_dir.symlink_to(outside_dir, target_is_directory=True)
    except (NotImplementedError, OSError) as exc:
        pytest.skip(f"directory symlink creation unavailable: {exc}")
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(output_dir),
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="Report output_dir must not be a symbolic link"):
        pipeline.recompute(dirty_from=None)

    assert not (outside_dir / "report.docx").exists()


def test_report_step_rejects_unsupported_language(tmp_path) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(tmp_path),
                "filename": "report.docx",
                "language": "jp",
            },
        )
    )

    with pytest.raises(ValueError, match="Unsupported report language"):
        pipeline.recompute(dirty_from=None)


def test_report_step_reports_missing_upstream_analysis_key_clearly(tmp_path) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:missing"],
                "output_dir": str(tmp_path),
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="Missing analysis result: reliability:missing"):
        pipeline.recompute(dirty_from=None)


def test_report_step_does_not_write_partial_figures_when_include_key_is_missing(
    tmp_path,
) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat", "comparison:missing"],
                "output_dir": str(tmp_path),
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="Missing analysis result: comparison:missing"):
        pipeline.recompute(dirty_from=None)

    assert not list(tmp_path.glob("*.png"))
    assert not list(tmp_path.glob("*.svg"))
    assert not list(tmp_path.glob("*.eps"))
    assert not (tmp_path / "report.docx").exists()


def test_report_step_cleans_generated_artifacts_when_chart_render_fails(
    tmp_path,
) -> None:
    step = ReportStep(
        id="report",
        title="APA report",
        params={
            "include": ["ok", "bad"],
            "output_dir": str(tmp_path),
            "filename": "report.docx",
            "language": "ko",
        },
    )

    with pytest.raises(ValueError, match="Unsupported chart type"):
        step.compute(
            PipelineContext(
                dataset=Dataset.empty(),
                analyses={
                    "ok": reliability_result_with_chart("horizontal_bar"),
                    "bad": reliability_result_with_chart("unsupported"),
                },
            )
        )

    assert not list(tmp_path.glob("*.png"))
    assert not list(tmp_path.glob("*.svg"))
    assert not list(tmp_path.glob("*.eps"))
    assert not (tmp_path / "report.docx").exists()


def test_report_step_failure_preserves_preexisting_empty_output_dir(tmp_path) -> None:
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    step = ReportStep(
        id="report",
        title="APA report",
        params={
            "include": ["ok", "bad"],
            "output_dir": str(output_dir),
            "filename": "report.docx",
            "language": "ko",
        },
    )

    with pytest.raises(ValueError, match="Unsupported chart type"):
        step.compute(
            PipelineContext(
                dataset=Dataset.empty(),
                analyses={
                    "ok": reliability_result_with_chart("horizontal_bar"),
                    "bad": reliability_result_with_chart("unsupported"),
                },
            )
        )

    assert output_dir.exists()
    assert output_dir.is_dir()


def test_report_step_cleanup_does_not_delete_renderer_path_outside_output_dir(
    tmp_path,
    monkeypatch,
) -> None:
    output_dir = tmp_path / "out"
    outside = tmp_path / "outside.png"

    def fake_render_chart(spec, output_path, key=None):
        if spec.type == "unsupported":
            raise ValueError("Unsupported chart type")
        outside.write_text("user image", encoding="utf-8")
        return str(outside)

    monkeypatch.setattr("modori.steps.reporting.render_chart", fake_render_chart)
    step = ReportStep(
        id="report",
        title="APA report",
        params={
            "include": ["ok", "bad"],
            "output_dir": str(output_dir),
            "filename": "report.docx",
            "language": "ko",
        },
    )

    with pytest.raises(ValueError, match="Unsupported chart type"):
        step.compute(
            PipelineContext(
                dataset=Dataset.empty(),
                analyses={
                    "ok": reliability_result_with_chart("horizontal_bar"),
                    "bad": reliability_result_with_chart("unsupported"),
                },
            )
        )

    assert outside.read_text(encoding="utf-8") == "user image"


def test_report_step_cleans_generated_artifacts_when_docx_write_fails(
    tmp_path,
    monkeypatch,
) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(tmp_path),
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )

    def failing_write_docx(prose, tables, figure_paths, path):
        path.write_text("partial", encoding="utf-8")
        raise RuntimeError("docx failure")

    monkeypatch.setattr("modori.steps.reporting.write_docx", failing_write_docx)

    with pytest.raises(RuntimeError, match="docx failure"):
        pipeline.recompute(dirty_from=None)

    assert not list(tmp_path.glob("*.png"))
    assert not list(tmp_path.glob("*.svg"))
    assert not list(tmp_path.glob("*.eps"))
    assert not (tmp_path / "report.docx").exists()


def test_report_step_uses_distinct_figure_paths_across_recompute(tmp_path) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(tmp_path),
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )
    pipeline.recompute(dirty_from=None)
    first_paths = set(pipeline.analysis_objects["report"].figure_paths["reliability:job_sat"])

    pipeline.recompute(dirty_from="report")
    second_paths = set(pipeline.analysis_objects["report"].figure_paths["reliability:job_sat"])

    assert first_paths.isdisjoint(second_paths)

@pytest.mark.parametrize("param_name", ["output_docx", "docx_path", "output_path"])
def test_report_step_rejects_direct_output_path_aliases(param_name, tmp_path) -> None:
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                param_name: str(tmp_path / "direct.docx"),
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="direct output paths are not supported"):
        pipeline.recompute(dirty_from=None)


@pytest.mark.parametrize("chart_dir", ["../outside", "/outside"])
def test_report_step_rejects_chart_dir_outside_output_dir(chart_dir, tmp_path) -> None:
    output_dir = tmp_path / "out"
    pipeline = Pipeline(report_dataset())
    pipeline.add(
        ReliabilityStep(
            id="reliability",
            title="Reliability",
            params={"items": ["q1", "q2", "q3", "q4"], "scale_name": "job_sat"},
        )
    )
    pipeline.add(
        ReportStep(
            id="report",
            title="APA report",
            params={
                "include": ["reliability:job_sat"],
                "output_dir": str(output_dir),
                "chart_dir": chart_dir,
                "filename": "report.docx",
                "language": "ko",
            },
        )
    )

    with pytest.raises(ValueError, match="chart_dir must stay within output_dir"):
        pipeline.recompute(dirty_from=None)
