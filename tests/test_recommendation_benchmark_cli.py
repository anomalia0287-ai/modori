from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from modori.recommendation_benchmark import ScorerConfig, scorer_fingerprint
from scripts.build_recommendation_pilot import build_pilot_pack
from scripts.recommendation_benchmark import main


@pytest.fixture
def completed_pack(tmp_path: Path) -> Path:
    root = tmp_path / "completed-recommendation-pack"
    build_pilot_pack(root)
    pilot = root / "public" / "pilot"
    for name, reviewer_id, minutes in (
        ("reviewer-a.xlsx", "reviewer-a", 10.0),
        ("reviewer-b.xlsx", "reviewer-b", 12.0),
    ):
        path = pilot / name
        workbook = load_workbook(path)
        workbook["Instructions"]["B2"] = reviewer_id
        reviews = workbook["Case Reviews"]
        abstentions = workbook["Abstentions"]
        for row in range(2, 22):
            reviews.cell(row=row, column=3).value = "abstention_required"
            reviews.cell(row=row, column=4).value = "E4"
            reviews.cell(row=row, column=5).value = minutes
            abstentions.append(
                [
                    reviews.cell(row=row, column=1).value,
                    reviews.cell(row=row, column=2).value,
                    "unsupported_design",
                ]
            )
        workbook.save(path)
    path = pilot / "adjudication.xlsx"
    workbook = load_workbook(path)
    workbook["Instructions"]["B2"] = "adjudicator"
    reviews = workbook["Case Reviews"]
    abstentions = workbook["Abstentions"]
    resolution = workbook["Resolution Minutes"]
    for row in range(2, 22):
        reviews.cell(row=row, column=3).value = "abstention_required"
        reviews.cell(row=row, column=4).value = "E4"
        abstentions.append(
            [
                reviews.cell(row=row, column=1).value,
                reviews.cell(row=row, column=2).value,
                "unsupported_design",
            ]
        )
        resolution.cell(row=row, column=3).value = 4.0
    workbook.save(path)
    return root


def test_cli_runs_as_documented_script_from_repository_root(tmp_path: Path) -> None:
    output = tmp_path / "validation.json"
    env = dict(os.environ)
    env["PYTHONPATH"] = str(Path("src").resolve())

    completed = subprocess.run(
        [
            sys.executable,
            "scripts/recommendation_benchmark.py",
            "validate-pack",
            "--pack-root",
            "tests/fixtures/recommendation_benchmark",
            "--output",
            str(output),
        ],
        capture_output=True,
        text=True,
        check=False,
        env=env,
    )

    assert completed.returncode == 0, completed.stderr
    assert json.loads(output.read_text(encoding="utf-8"))["case_count"] == 20


def test_validate_pack_and_predict_a_write_explicit_outputs(tmp_path: Path) -> None:
    pack = Path("tests/fixtures/recommendation_benchmark")
    validation_output = tmp_path / "validation.json"
    prediction_output = tmp_path / "predictions.jsonl"

    assert main(
        [
            "validate-pack",
            "--pack-root",
            str(pack),
            "--output",
            str(validation_output),
        ]
    ) == 0
    assert json.loads(validation_output.read_text(encoding="utf-8"))["case_count"] == 20

    assert main(
        [
            "predict-a",
            "--pack-root",
            str(pack),
            "--output",
            str(prediction_output),
        ]
    ) == 0
    assert len(prediction_output.read_text(encoding="utf-8").splitlines()) == 20
    assert main(
        [
            "predict-a",
            "--pack-root",
            str(pack),
            "--output",
            str(prediction_output),
        ]
    ) == 1


def test_validate_submissions_agreement_and_cost_commands(
    completed_pack: Path,
    tmp_path: Path,
) -> None:
    pilot = completed_pack / "public" / "pilot"
    common = [
        "--pack-root",
        str(completed_pack),
        "--reviewer-a",
        str(pilot / "reviewer-a.xlsx"),
        "--reviewer-b",
        str(pilot / "reviewer-b.xlsx"),
    ]
    validation_output = tmp_path / "submissions.json"
    agreement_output = tmp_path / "agreement.json"
    cost_output = tmp_path / "cost.json"

    assert main(
        [
            "validate-submissions",
            *common,
            "--adjudication",
            str(pilot / "adjudication.xlsx"),
            "--output",
            str(validation_output),
        ]
    ) == 0
    assert json.loads(validation_output.read_text(encoding="utf-8"))["case_count"] == 20

    assert main(["agreement", *common, "--output", str(agreement_output)]) == 0
    agreement = json.loads(agreement_output.read_text(encoding="utf-8"))
    assert agreement["primary_action_status"] == "not_estimable"

    assert main(
        [
            "cost",
            *common,
            "--adjudication",
            str(pilot / "adjudication.xlsx"),
            "--stage-cases",
            "150",
            "--reviewer-rate",
            "100",
            "--adjudicator-rate",
            "150",
            "--output",
            str(cost_output),
        ]
    ) == 0
    cost = json.loads(cost_output.read_text(encoding="utf-8"))
    assert cost["stage_case_count"] == 150
    assert cost["median_reviewer_minutes"] == 11.0


def test_score_refuses_wrong_fingerprint_and_scores_completed_gold(
    completed_pack: Path,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    pilot = completed_pack / "public" / "pilot"
    output = tmp_path / "score.json"
    base_args = [
        "score",
        "--pack-root",
        str(completed_pack),
        "--adjudication",
        str(pilot / "adjudication.xlsx"),
        "--predictions",
        str(pilot / "baseline-a-predictions.jsonl"),
        "--output",
        str(output),
    ]

    assert main([*base_args, "--scorer-fingerprint", "sha256:wrong"]) == 1
    assert "fingerprint mismatch" in capsys.readouterr().err
    assert not output.exists()

    assert main(
        [
            *base_args,
            "--scorer-fingerprint",
            scorer_fingerprint(ScorerConfig()),
        ]
    ) == 0
    score = json.loads(output.read_text(encoding="utf-8"))
    assert score["abstention_accuracy"]["successes"] == 0
    assert score["errors_by_severity"] == [["E1", 0], ["E2", 0], ["E3", 0], ["E4", 20], ["E5", 0]]


def test_blank_submission_is_rejected_with_nonzero_exit(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    pack = Path("tests/fixtures/recommendation_benchmark")
    pilot = pack / "public" / "pilot"
    output = tmp_path / "invalid.json"

    assert main(
        [
            "validate-submissions",
            "--pack-root",
            str(pack),
            "--reviewer-a",
            str(pilot / "reviewer-a.xlsx"),
            "--reviewer-b",
            str(pilot / "reviewer-b.xlsx"),
            "--adjudication",
            str(pilot / "adjudication.xlsx"),
            "--output",
            str(output),
        ]
    ) == 1
    assert "reviewer ID" in capsys.readouterr().err
    assert not output.exists()
