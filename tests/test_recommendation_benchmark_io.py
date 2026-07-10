from __future__ import annotations

from pathlib import Path

import pytest
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook

from modori.recommendation_benchmark import BenchmarkContractError
from modori.recommendation_benchmark_io import (
    PilotCaseSummary,
    _normalize_core_properties,
    build_blank_pilot_workbooks,
    load_adjudication_workbook,
    load_reviewer_workbook,
    read_jsonl,
    workbook_schema_fingerprint,
    write_jsonl,
)


def test_core_property_normalization_rejects_xml_entities() -> None:
    malicious = b"""<?xml version='1.0' encoding='UTF-8'?>
<!DOCTYPE coreProperties [<!ENTITY injected 'unsafe'>]>
<coreProperties xmlns:dcterms='http://purl.org/dc/terms/'>
  <dcterms:created>&injected;</dcterms:created>
</coreProperties>
"""

    with pytest.raises(DefusedXmlException):
        _normalize_core_properties(malicious)


def test_jsonl_round_trip_is_canonical_and_accepts_utf8_bom(tmp_path: Path) -> None:
    path = tmp_path / "records.jsonl"
    path.write_bytes(b"\xef\xbb\xbf" + '{"b":2,"a":1}\n'.encode())

    records = read_jsonl(path, lambda mapping: mapping)

    assert records == ({"a": 1, "b": 2},)

    output = tmp_path / "output.jsonl"
    write_jsonl(output, records)
    assert output.read_text(encoding="utf-8") == '{"a":1,"b":2}\n'


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ('{"a":1,"a":2}\n', "duplicate JSON key"),
        ('{"a":1}\n\n', "blank line"),
        ('[1,2,3]\n', "must be an object"),
        ('{"a":NaN}\n', "non-standard JSON constant"),
    ],
)
def test_read_jsonl_rejects_ambiguous_or_noncanonical_records(
    tmp_path: Path,
    content: str,
    message: str,
) -> None:
    path = tmp_path / "bad.jsonl"
    path.write_text(content, encoding="utf-8")

    with pytest.raises(BenchmarkContractError, match=message):
        read_jsonl(path, lambda mapping: mapping)


def test_write_jsonl_refuses_implicit_overwrite(tmp_path: Path) -> None:
    path = tmp_path / "records.jsonl"
    path.write_text("existing\n", encoding="utf-8")

    with pytest.raises(BenchmarkContractError, match="already exists"):
        write_jsonl(path, ({"case_id": "case-1"},))

    write_jsonl(path, ({"case_id": "case-1"},), overwrite=True)
    assert path.read_text(encoding="utf-8") == '{"case_id":"case-1"}\n'


def pilot_cases() -> tuple[PilotCaseSummary, ...]:
    return (
        PilotCaseSummary(
            case_id="case-1",
            evidence_stage="cold_start",
            title_ko="두 집단 점수",
            research_question_ko="두 집단의 점수 차이를 확인한다.",
            data_file="data/case-1.csv",
            unit_of_observation="한 행은 참여자 한 명이다.",
            sampling="두 독립 집단의 합성 표본이다.",
            grouping="arm이 집단을 구분한다.",
            time_structure="한 시점이다.",
            weights_clusters="없다.",
            variable_meanings="arm: 집단\nscore: 결과 점수",
            known_missing_codes="없음",
            facts_visible="각 참여자는 한 집단에만 속한다.",
            facts_clarification_only="없음",
        ),
        PilotCaseSummary(
            case_id="case-2",
            evidence_stage="cold_start",
            title_ko="지원하지 않는 군집 설계",
            research_question_ko="군집 표본의 차이를 확인한다.",
            data_file="data/case-2.csv",
            unit_of_observation="한 행은 참여자 한 명이다.",
            sampling="군집 표본이다.",
            grouping="arm이 집단을 구분한다.",
            time_structure="한 시점이다.",
            weights_clusters="cluster_id와 weight를 반영해야 한다.",
            variable_meanings="cluster_id: 군집\nweight: 가중치",
            known_missing_codes="없음",
            facts_visible="가중치와 군집 열이 있다.",
            facts_clarification_only="없음",
        ),
    )


def test_blank_workbooks_have_korean_instructions_stable_schema_and_no_answers(
    tmp_path: Path,
) -> None:
    reviewer_a, reviewer_b, adjudication = build_blank_pilot_workbooks(
        pilot_cases(),
        tmp_path,
    )

    workbook = load_workbook(reviewer_a, data_only=False)
    assert workbook.sheetnames == [
        "Instructions",
        "Study Cards",
        "Case Reviews",
        "Recommendations",
        "Clarifications",
        "Abstentions",
    ]
    assert "독립적으로" in str(workbook["Instructions"]["A1"].value)
    assert "검토자" in str(workbook["Instructions"]["A2"].value)
    assert "A1:B1" in {str(cell_range) for cell_range in workbook["Instructions"].merged_cells.ranges}
    assert workbook["Instructions"].row_dimensions[1].height >= 60
    assert workbook["Instructions"]["B2"].value is None
    assert workbook["Study Cards"]["E2"].value == "한 행은 참여자 한 명이다."
    assert "score: 결과 점수" in workbook["Study Cards"]["J2"].value
    assert workbook["Study Cards"].freeze_panes == "C2"
    reviews = workbook["Case Reviews"]
    assert [reviews.cell(row=row, column=1).value for row in (2, 3)] == [
        "case-1",
        "case-2",
    ]
    assert [reviews.cell(row=2, column=column).value for column in range(3, 7)] == [
        None,
        None,
        None,
        None,
    ]
    assert len(reviews.data_validations.dataValidation) == 2
    assert workbook_schema_fingerprint(reviewer_a) == workbook_schema_fingerprint(
        reviewer_b
    )
    assert workbook_schema_fingerprint(reviewer_a) != workbook_schema_fingerprint(
        adjudication
    )


def test_completed_reviewer_workbook_loads_strict_annotation_contract(
    tmp_path: Path,
) -> None:
    reviewer_a, _reviewer_b, _adjudication = build_blank_pilot_workbooks(
        pilot_cases(),
        tmp_path,
    )
    workbook = load_workbook(reviewer_a)
    workbook["Instructions"]["B2"] = "reviewer-a"
    reviews = workbook["Case Reviews"]
    reviews["C2"] = "recommendation_eligible"
    reviews["D2"] = "E3"
    reviews["E2"] = 8.5
    reviews["C3"] = "abstention_required"
    reviews["D3"] = "E4"
    reviews["E3"] = 10.0
    recommendations = workbook["Recommendations"]
    recommendations.append(
        [
            "case-1",
            "cold_start",
            1,
            "compare_groups",
            "independent",
            "score",
            "arm",
            None,
            None,
            None,
            None,
            None,
        ]
    )
    workbook["Abstentions"].append(
        ["case-2", "cold_start", "unsupported_design"]
    )
    workbook.save(reviewer_a)

    submission = load_reviewer_workbook(reviewer_a, pilot_cases())

    assert submission.reviewer_id == "reviewer-a"
    assert len(submission.annotations) == 2
    assert submission.annotations[0].action_class == "recommendation_eligible"
    recommendation = submission.annotations[0].acceptable_recommendations[0]
    assert recommendation.family == "compare_groups"
    assert dict(recommendation.roles) == {
        "group": ("arm",),
        "outcome": ("score",),
    }
    assert submission.annotations[1].acceptable_abstention_reasons == (
        "unsupported_design",
    )


def test_workbook_loader_rejects_formula_cells_and_schema_drift(tmp_path: Path) -> None:
    reviewer_a, reviewer_b, _adjudication = build_blank_pilot_workbooks(
        pilot_cases(),
        tmp_path,
    )
    formula_workbook = load_workbook(reviewer_a)
    formula_workbook["Instructions"]["B2"] = "reviewer-a"
    formula_workbook["Case Reviews"]["C2"] = '=IF(1=1,"recommendation_eligible","")'
    formula_workbook.save(reviewer_a)

    with pytest.raises(BenchmarkContractError, match="formula cells are forbidden"):
        load_reviewer_workbook(reviewer_a, pilot_cases())

    drifted_workbook = load_workbook(reviewer_b)
    drifted_workbook.remove(drifted_workbook["Clarifications"])
    drifted_workbook.save(reviewer_b)

    with pytest.raises(BenchmarkContractError, match="workbook sheets do not match"):
        load_reviewer_workbook(reviewer_b, pilot_cases())


def test_workbook_builder_writes_formula_like_case_metadata_as_literal_text(
    tmp_path: Path,
) -> None:
    cases = (
        PilotCaseSummary(
            case_id="formula-case",
            evidence_stage="cold_start",
            title_ko='=HYPERLINK("https://invalid.example","열기")',
            research_question_ko="+1은 연구 질문의 일부다.",
            data_file="data/formula-case.csv",
        ),
    )

    reviewer_a, _reviewer_b, _adjudication = build_blank_pilot_workbooks(
        cases,
        tmp_path,
    )
    workbook = load_workbook(reviewer_a, data_only=False)

    assert workbook["Case Reviews"]["G2"].data_type == "s"
    assert workbook["Case Reviews"]["G2"].value == cases[0].title_ko


def test_reviewer_workbook_rejects_missing_time_and_incomplete_case_coverage(
    tmp_path: Path,
) -> None:
    missing_dir = tmp_path / "missing"
    reviewer_a, _reviewer_b, _adjudication = build_blank_pilot_workbooks(
        pilot_cases(),
        missing_dir,
    )
    workbook = load_workbook(reviewer_a)
    workbook["Instructions"]["B2"] = "reviewer-a"
    reviews = workbook["Case Reviews"]
    for row in (2, 3):
        reviews.cell(row=row, column=3).value = "abstention_required"
        reviews.cell(row=row, column=4).value = "E4"
        workbook["Abstentions"].append(
            [reviews.cell(row=row, column=1).value, "cold_start", "unsupported_design"]
        )
    reviews["E2"] = 5.0
    workbook.save(reviewer_a)

    with pytest.raises(BenchmarkContractError, match="active_minutes must be numeric"):
        load_reviewer_workbook(reviewer_a, pilot_cases())

    incomplete_dir = tmp_path / "incomplete"
    reviewer_a, _reviewer_b, _adjudication = build_blank_pilot_workbooks(
        pilot_cases(),
        incomplete_dir,
    )
    workbook = load_workbook(reviewer_a)
    workbook["Instructions"]["B2"] = "reviewer-a"
    workbook["Case Reviews"].delete_rows(3)
    workbook["Case Reviews"]["C2"] = "abstention_required"
    workbook["Case Reviews"]["D2"] = "E4"
    workbook["Case Reviews"]["E2"] = 5.0
    workbook["Abstentions"].append(
        ["case-1", "cold_start", "unsupported_design"]
    )
    workbook.save(reviewer_a)

    with pytest.raises(BenchmarkContractError, match="cover every expected case-stage"):
        load_reviewer_workbook(reviewer_a, pilot_cases())


def test_completed_adjudication_workbook_yields_gold_and_resolution_minutes(
    tmp_path: Path,
) -> None:
    _reviewer_a, _reviewer_b, adjudication = build_blank_pilot_workbooks(
        pilot_cases(),
        tmp_path,
    )
    workbook = load_workbook(adjudication)
    workbook["Instructions"]["B2"] = "adjudicator"
    reviews = workbook["Case Reviews"]
    reviews["C2"] = "recommendation_eligible"
    reviews["D2"] = "E3"
    reviews["C3"] = "abstention_required"
    reviews["D3"] = "E4"
    workbook["Recommendations"].append(
        [
            "case-1",
            "cold_start",
            1,
            "compare_groups",
            "independent",
            "score",
            "arm",
            None,
            None,
            None,
            None,
            None,
        ]
    )
    workbook["Abstentions"].append(
        ["case-2", "cold_start", "unsupported_design"]
    )
    resolution = workbook["Resolution Minutes"]
    resolution["C2"] = 3.0
    resolution["C3"] = 1.0
    workbook.save(adjudication)

    result = load_adjudication_workbook(adjudication, pilot_cases())

    assert [record.action_class for record in result.gold_records] == [
        "recommendation_eligible",
        "abstention_required",
    ]
    assert [record.resolution_minutes for record in result.adjudications] == [3.0, 1.0]
