import pandas as pd
import pyreadstat
import pytest
from openpyxl import load_workbook

from modori.table_io import (
    FullReadLimits,
    PreviewReadLimits,
    TableHeaderResult,
    TableLayoutOverride,
    TablePreviewResult,
    TableReadError,
    TableReadResult,
    read_full,
    read_header,
    read_header_result,
    read_preview,
)


def test_read_full_returns_structured_result_that_stays_tuple_unpackable(tmp_path) -> None:
    path = tmp_path / "survey.csv"
    pd.DataFrame({"score": [1, 2]}).to_csv(path, index=False)

    result = read_full(path, "csv")
    frame, metadata = result

    assert isinstance(result, TableReadResult)
    assert frame.equals(result.frame)
    assert metadata is None
    assert result.source.path == path
    assert result.source.file_type == "csv"
    assert result.columns == ("score",)
    assert result.row_count == 2


def test_read_preview_returns_xlsx_source_context_and_sample_rows(tmp_path) -> None:
    path = tmp_path / "survey.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({"respondent_id": [101, 102, 103], "score": [5, 4, 3]}).to_excel(
            writer,
            sheet_name="Survey Responses",
            index=False,
        )
        pd.DataFrame({"code": [1, 2], "label": ["control", "treatment"]}).to_excel(
            writer,
            sheet_name="Codebook",
            index=False,
        )

    result = read_preview(path, "xlsx", limits=PreviewReadLimits(max_rows=2))
    frame, metadata = result

    assert isinstance(result, TablePreviewResult)
    assert metadata is None
    assert frame.shape == (2, 2)
    assert result.source.file_type == "xlsx"
    assert result.source.path == path
    assert result.source.sheet_name == "Survey Responses"
    assert result.source.sheet_names == ("Survey Responses", "Codebook")
    assert result.preview_limit == 2
    assert result.previewed_rows == 2
    assert result.columns == ("respondent_id", "score")
    assert result.sample_rows == (
        {"respondent_id": 101, "score": 5},
        {"respondent_id": 102, "score": 4},
    )


def test_read_header_result_exposes_xlsx_source_context_and_compat_header(tmp_path) -> None:
    path = tmp_path / "survey.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({"q1": [1], "q2": [2]}).to_excel(
            writer,
            sheet_name="Responses",
            index=False,
        )
        pd.DataFrame({"code": [1]}).to_excel(writer, sheet_name="Codebook", index=False)

    result = read_header_result(path, "xlsx")

    assert isinstance(result, TableHeaderResult)
    assert result.columns == ("q1", "q2")
    assert tuple(result) == ("q1", "q2")
    assert result.source.file_type == "xlsx"
    assert result.source.sheet_name == "Responses"
    assert result.source.sheet_names == ("Responses", "Codebook")
    assert read_header(path, "xlsx") == ["q1", "q2"]


def test_read_header_applies_user_csv_layout_override(tmp_path) -> None:
    path = tmp_path / "manual-layout.csv"
    path.write_text(
        "\n".join(
            [
                "다운로드 조건,2026-07-06",
                "이 행은 표가 아닙니다,확인용",
                "city,value",
                "Seoul,10",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    columns = read_header(
        path,
        "csv",
        layout=TableLayoutOverride(header_row_index=2, header_row_count=1),
    )

    assert columns == ["city", "value"]


def test_xlsx_full_header_and_preview_use_same_active_sheet(tmp_path) -> None:
    path = tmp_path / "survey.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({"code": [1], "label": ["treatment"]}).to_excel(
            writer,
            sheet_name="Codebook",
            index=False,
        )
        pd.DataFrame({"respondent_id": [101, 102], "score": [5, 4]}).to_excel(
            writer,
            sheet_name="Responses",
            index=False,
        )

    workbook = load_workbook(path)
    workbook.active = workbook.sheetnames.index("Responses")
    workbook.save(path)
    workbook.close()

    preview = read_preview(path, "xlsx", limits=PreviewReadLimits(max_rows=1))
    header = read_header_result(path, "xlsx")
    full = read_full(path, "xlsx")

    assert preview.source.sheet_name == "Responses"
    assert header.source.sheet_name == "Responses"
    assert full.source.sheet_name == "Responses"
    assert preview.columns == ("respondent_id", "score")
    assert header.columns == ("respondent_id", "score")
    assert full.columns == ("respondent_id", "score")
    assert preview.frame.to_dict(orient="list") == {
        "respondent_id": [101],
        "score": [5],
    }
    assert full.frame.to_dict(orient="list") == {
        "respondent_id": [101, 102],
        "score": [5, 4],
    }


def test_read_preview_uses_bounded_csv_columns(tmp_path, monkeypatch) -> None:
    path = tmp_path / "survey.csv"
    path.write_text("a,b,c\n1,2,3\n", encoding="utf-8")
    calls = []

    def read_csv_spy(path_arg, *args, **kwargs):
        calls.append(kwargs)
        if kwargs.get("nrows") == 0:
            return pd.DataFrame(columns=["a", "b", "c"])
        return pd.DataFrame({"a": [1], "b": [2]})

    monkeypatch.setattr("modori.table_io.pd.read_csv", read_csv_spy)

    result = read_preview(path, "csv", limits=PreviewReadLimits(max_rows=2, max_columns=2))

    assert result.columns == ("a", "b")
    assert result.warnings == ("미리보기 열 제한: 3개 중 2개 열만 표시합니다.",)
    assert calls == [
        {"nrows": 0},
        {"nrows": 2, "usecols": ["a", "b"]},
    ]


def test_default_preview_limits_do_not_block_by_file_size() -> None:
    assert PreviewReadLimits().max_file_bytes is None


def test_read_preview_limits_xlsx_columns(tmp_path) -> None:
    path = tmp_path / "survey.xlsx"
    pd.DataFrame(
        {
            "a": [1, 4],
            "b": [2, 5],
            "c": [3, 6],
        }
    ).to_excel(path, index=False)

    result = read_preview(path, "xlsx", limits=PreviewReadLimits(max_rows=2, max_columns=2))

    assert result.columns == ("a", "b")
    assert result.frame.to_dict(orient="list") == {"a": [1, 4], "b": [2, 5]}
    assert result.warnings == ("미리보기 열 제한: 3개 중 2개 열만 표시합니다.",)


def test_read_full_decodes_cp949_public_csv(tmp_path) -> None:
    path = tmp_path / "public.csv"
    text = "자치구,연도,인구\n종로구,2024,140000\n중구,2024,120000\n"
    path.write_bytes(text.encode("cp949"))

    result = read_full(path, "csv")

    assert result.columns == ("자치구", "연도", "인구")
    assert result.frame.to_dict(orient="list") == {
        "자치구": ["종로구", "중구"],
        "연도": [2024, 2024],
        "인구": [140000, 120000],
    }
    assert "CSV 인코딩: cp949" in result.warnings


def test_read_preview_skips_public_csv_metadata_rows(tmp_path) -> None:
    path = tmp_path / "public-with-preamble.csv"
    path.write_text(
        "\n".join(
            [
                "서울시 인구 현황",
                "자료기준일: 2024-12-31",
                "단위: 명",
                "자치구,연도,인구",
                "종로구,2024,\"140,000\"",
                "중구,2024,\"120,000\"",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    result = read_preview(path, "csv", limits=PreviewReadLimits(max_rows=2))

    assert result.columns == ("자치구", "연도", "인구")
    assert result.sample_rows == (
        {"자치구": "종로구", "연도": 2024, "인구": "140,000"},
        {"자치구": "중구", "연도": 2024, "인구": "120,000"},
    )
    assert "표 헤더 앞의 안내 행 3개를 건너뛰었습니다." in result.warnings


def test_read_full_skips_xlsx_metadata_rows(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "public-with-preamble.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "공개자료"
    worksheet.append(["서울시 시설 현황"])
    worksheet.append(["자료기준일", "2024-12-31"])
    worksheet.append(["단위", "개소"])
    worksheet.append(["자치구", "연도", "시설수"])
    worksheet.append(["종로구", 2024, 12])
    worksheet.append(["중구", 2024, 10])
    workbook.save(path)
    workbook.close()

    result = read_full(path, "xlsx")

    assert result.columns == ("자치구", "연도", "시설수")
    assert result.frame.to_dict(orient="list") == {
        "자치구": ["종로구", "중구"],
        "연도": [2024, 2024],
        "시설수": [12, 10],
    }
    assert "표 헤더 앞의 안내 행 3개를 건너뛰었습니다." in result.warnings


def test_read_full_sanitizes_blank_and_duplicate_xlsx_headers(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "public-duplicate-headers.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["자치구", None, "인구", "인구"])
    worksheet.append(["종로구", None, 140000, 150000])
    workbook.save(path)
    workbook.close()

    result = read_full(path, "xlsx")

    assert result.columns == ("자치구", "인구", "인구_2")
    assert result.frame.to_dict(orient="list") == {
        "자치구": ["종로구"],
        "인구": [140000],
        "인구_2": [150000],
    }
    assert "빈 열 1개를 제외했습니다." in result.warnings
    assert "중복 열 이름 1개를 고유한 이름으로 바꿨습니다." in result.warnings


def test_read_full_flattens_two_row_public_csv_headers(tmp_path) -> None:
    path = tmp_path / "kosis-two-row.csv"
    text = "\n".join(
        [
            "행정구역별(1),특성별(1),특성별(2),2025,2025,2025",
            "행정구역별(1),특성별(1),특성별(2),계 (%),매우 만족,약간 만족",
            "전국,전체,계,100.0,11.5,27.9",
            "서울특별시,전체,계,100.0,11.0,30.3",
        ]
    )
    path.write_bytes((text + "\n").encode("cp949"))

    result = read_full(path, "csv")

    assert result.columns == (
        "행정구역별(1)",
        "특성별(1)",
        "특성별(2)",
        "2025 계 (%)",
        "2025 매우 만족",
        "2025 약간 만족",
    )
    assert result.frame.to_dict(orient="records")[0] == {
        "행정구역별(1)": "전국",
        "특성별(1)": "전체",
        "특성별(2)": "계",
        "2025 계 (%)": 100.0,
        "2025 매우 만족": 11.5,
        "2025 약간 만족": 27.9,
    }
    assert "다중 헤더 2행을 하나의 열 이름으로 합쳤습니다." in result.warnings


def test_read_preview_reports_two_row_public_csv_inference(tmp_path) -> None:
    path = tmp_path / "kosis-two-row.csv"
    text = "\n".join(
        [
            "행정구역별(1),특성별(1),특성별(2),2025,2025,2025",
            "행정구역별(1),특성별(1),특성별(2),계 (%),매우 만족,약간 만족",
            "전국,전체,계,100.0,11.5,27.9",
            "서울특별시,전체,계,100.0,11.0,30.3",
        ]
    )
    path.write_bytes((text + "\n").encode("cp949"))

    result = read_preview(path, "csv")

    report = result.inference_report
    assert report is not None
    assert report.file_type == "csv"
    assert report.header_row_index == 0
    assert report.header_row_count == 2
    assert report.data_start_row_index == 2
    assert report.column_count == 6
    assert report.confidence == "high"
    assert report.requires_user_confirmation is False
    assert "다중 헤더 2행을 하나의 열 이름으로 합쳤습니다." in report.reasons


def test_read_full_keeps_categorical_survey_rows_as_data(tmp_path) -> None:
    path = tmp_path / "survey.csv"
    text = "\n".join(
        [
            "group,region,score",
            "A,Seoul,1",
            "B,Busan,2",
        ]
    )
    path.write_text(text + "\n", encoding="utf-8")

    result = read_full(path, "csv")

    assert result.columns == ("group", "region", "score")
    assert result.frame.to_dict(orient="records") == [
        {"group": "A", "region": "Seoul", "score": 1},
        {"group": "B", "region": "Busan", "score": 2},
    ]
    assert "다중 헤더 2행을 하나의 열 이름으로 합쳤습니다." not in result.warnings


def test_read_preview_reports_plain_csv_inference(tmp_path) -> None:
    path = tmp_path / "survey.csv"
    text = "\n".join(
        [
            "group,region,score",
            "A,Seoul,1",
            "B,Busan,2",
        ]
    )
    path.write_text(text + "\n", encoding="utf-8")

    result = read_preview(path, "csv")

    report = result.inference_report
    assert report is not None
    assert report.file_type == "csv"
    assert report.header_row_index == 0
    assert report.header_row_count == 1
    assert report.data_start_row_index == 1
    assert report.column_count == 3
    assert report.confidence == "high"
    assert report.requires_user_confirmation is False
    assert report.reasons == ("첫 번째 행을 헤더로 인식했습니다.",)


def test_read_preview_applies_user_csv_header_override(tmp_path) -> None:
    path = tmp_path / "manual-layout.csv"
    text = "\n".join(
        [
            "다운로드 조건,2026-07-06",
            "이 행은 표가 아닙니다,확인용",
            "city,value",
            "Seoul,10",
            "Busan,20",
        ]
    )
    path.write_text(text + "\n", encoding="utf-8")

    result = read_preview(
        path,
        "csv",
        layout=TableLayoutOverride(header_row_index=2, header_row_count=1),
    )

    assert result.columns == ("city", "value")
    assert result.sample_rows == (
        {"city": "Seoul", "value": 10},
        {"city": "Busan", "value": 20},
    )
    report = result.inference_report
    assert report is not None
    assert report.header_row_index == 2
    assert report.header_row_count == 1
    assert report.data_start_row_index == 3
    assert report.confidence == "high"
    assert "사용자 지정 표 레이아웃을 적용했습니다." in report.reasons


def test_read_preview_rejects_out_of_range_csv_layout_override(tmp_path) -> None:
    path = tmp_path / "manual-layout.csv"
    path.write_text("city,value\nSeoul,10\n", encoding="utf-8")

    with pytest.raises(TableReadError) as exc_info:
        read_preview(
            path,
            "csv",
            layout=TableLayoutOverride(header_row_index=10, header_row_count=1),
        )

    assert exc_info.value.message_ko == "지정한 표 레이아웃을 적용할 수 없습니다. 헤더 행과 데이터 시작 행을 확인해 주세요."


def test_read_full_applies_user_csv_data_start_override(tmp_path) -> None:
    path = tmp_path / "manual-data-start.csv"
    text = "\n".join(
        [
            "city,value",
            "주석,아래부터 데이터",
            "Seoul,10",
            "Busan,20",
        ]
    )
    path.write_text(text + "\n", encoding="utf-8")

    result = read_full(
        path,
        "csv",
        layout=TableLayoutOverride(
            header_row_index=0,
            header_row_count=1,
            data_start_row_index=2,
        ),
    )

    assert result.columns == ("city", "value")
    assert result.frame.to_dict(orient="records") == [
        {"city": "Seoul", "value": 10},
        {"city": "Busan", "value": 20},
    ]
    report = result.inference_report
    assert report is not None
    assert report.data_start_row_index == 2
    assert "사용자 지정 표 레이아웃을 적용했습니다." in report.reasons


def test_read_preview_reports_medium_confidence_for_deep_preamble(tmp_path) -> None:
    path = tmp_path / "real-estate.csv"
    preamble = [f"검색조건 {index}: 값" for index in range(15)]
    text = "\n".join(
        [
            *preamble,
            "시군구,번지,거래금액",
            "서울특별시 종로구,1-1,100000",
        ]
    )
    path.write_text(text + "\n", encoding="utf-8")

    result = read_preview(path, "csv")

    report = result.inference_report
    assert report is not None
    assert report.header_row_index == 15
    assert report.header_row_count == 1
    assert report.data_start_row_index == 16
    assert report.column_count == 3
    assert report.confidence == "medium"
    assert report.requires_user_confirmation is False
    assert "표 헤더 앞의 안내 행 15개를 건너뛰었습니다." in report.reasons


def test_read_full_keeps_all_text_survey_rows_as_data(tmp_path) -> None:
    path = tmp_path / "all-text-survey.csv"
    text = "\n".join(
        [
            "group,region,status",
            "A,Seoul,ready",
            "B,Busan,done",
        ]
    )
    path.write_text(text + "\n", encoding="utf-8")

    result = read_full(path, "csv")

    assert result.columns == ("group", "region", "status")
    assert result.frame.to_dict(orient="records") == [
        {"group": "A", "region": "Seoul", "status": "ready"},
        {"group": "B", "region": "Busan", "status": "done"},
    ]
    assert "다중 헤더 2행을 하나의 열 이름으로 합쳤습니다." not in result.warnings


def test_read_full_flattens_merged_xlsx_header_rows(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "merged-public-header.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["2020년_ [재적 학생 현황 (대학원)]"])
    worksheet.append(["작성자 : 교육혁신과"])
    worksheet.append([])
    worksheet.append(["순번", "기준연도", "학과", "재학생(A)", None, None, None])
    worksheet.append([None, None, None, "계", None, "남", None])
    worksheet.append([None, None, None, "정원내", "정원외", "정원내", "정원외"])
    worksheet.append(["합 계", None, None, 3240, 743, 1667, 417])
    worksheet.append([2, "2020", "ICT융합학과", 0, 37, 0, 27])
    workbook.save(path)
    workbook.close()

    result = read_full(path, "xlsx")

    assert result.columns == (
        "순번",
        "기준연도",
        "학과",
        "재학생(A) 계 정원내",
        "재학생(A) 계 정원외",
        "재학생(A) 남 정원내",
        "재학생(A) 남 정원외",
    )
    first_row = result.frame.iloc[0]
    assert first_row["순번"] == "합 계"
    assert pd.isna(first_row["기준연도"])
    assert pd.isna(first_row["학과"])
    assert first_row["재학생(A) 계 정원내"] == 3240
    assert first_row["재학생(A) 계 정원외"] == 743
    assert first_row["재학생(A) 남 정원내"] == 1667
    assert first_row["재학생(A) 남 정원외"] == 417
    assert "표 헤더 앞의 안내 행 3개를 건너뛰었습니다." in result.warnings
    assert "다중 헤더 3행을 하나의 열 이름으로 합쳤습니다." in result.warnings


def test_read_preview_warns_about_aggregate_rows_and_can_drop_them(tmp_path) -> None:
    path = tmp_path / "aggregate-row.csv"
    path.write_text(
        "\n".join(
            [
                "지역,인구",
                "합 계,300",
                "종로구,100",
                "중구,200",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    kept = read_preview(path, "csv")

    assert kept.sample_rows[0] == {"지역": "합 계", "인구": 300}
    assert "집계/합계 행 1개를 감지했습니다. 필요한 경우 가져오기 창에서 제외할 수 있습니다." in kept.warnings

    dropped = read_preview(path, "csv", drop_aggregate_rows=True)

    assert dropped.sample_rows == (
        {"지역": "종로구", "인구": 100},
        {"지역": "중구", "인구": 200},
    )
    assert "집계/합계 행 1개를 제외했습니다." in dropped.warnings


def test_read_preview_warns_and_drops_kosis_nationwide_total_row(tmp_path) -> None:
    path = tmp_path / "kosis-two-row.csv"
    text = "\n".join(
        [
            "행정구역별(1),특성별(1),특성별(2),2025,2025,2025",
            "행정구역별(1),특성별(1),특성별(2),계 (%),매우 만족,약간 만족",
            "전국,전체,계,100.0,11.5,27.9",
            "서울특별시,전체,계,100.0,11.0,30.3",
        ]
    )
    path.write_text(text + "\n", encoding="utf-8")

    kept = read_preview(path, "csv")

    assert kept.sample_rows[0]["행정구역별(1)"] == "전국"
    assert "집계/합계 행 1개를 감지했습니다. 필요한 경우 가져오기 창에서 제외할 수 있습니다." in kept.warnings

    dropped = read_preview(path, "csv", drop_aggregate_rows=True)

    assert dropped.sample_rows == (
        {
            "행정구역별(1)": "서울특별시",
            "특성별(1)": "전체",
            "특성별(2)": "계",
            "2025 계 (%)": 100.0,
            "2025 매우 만족": 11.0,
            "2025 약간 만족": 30.3,
        },
    )
    assert "집계/합계 행 1개를 제외했습니다." in dropped.warnings


def test_read_full_detects_aggregate_labels_after_row_number_markers(tmp_path) -> None:
    path = tmp_path / "numbered-aggregate-row.csv"
    path.write_text(
        "\n".join(
            [
                "순번,지역,인구",
                "1,소계,300",
                "2,종로구,100",
                "-,총합계,300",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    result = read_full(path, "csv", drop_aggregate_rows=True)

    assert result.frame.to_dict(orient="records") == [
        {"순번": "2", "지역": "종로구", "인구": 100},
    ]
    assert "집계/합계 행 2개를 제외했습니다." in result.warnings


def test_read_preview_rejects_long_same_width_preamble_instead_of_high_confidence_misread(
    tmp_path,
) -> None:
    path = tmp_path / "long-preamble.csv"
    lines = [f"검색조건 {index},값,메모" for index in range(21)]
    lines.extend(["시군구,번지,거래금액", "서울특별시 종로구,1-1,100000"])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    with pytest.raises(TableReadError) as exc_info:
        read_preview(path, "csv")

    assert exc_info.value.message_ko == (
        "표 헤더를 자동으로 찾지 못했습니다. 가져오기 창에서 헤더 행을 지정해 주세요."
    )


def test_read_preview_reports_merged_xlsx_header_inference(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "merged-public-header.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "재적학생"
    worksheet.append(["2020년_ [재적 학생 현황 (대학원)]"])
    worksheet.append(["작성자 : 교육혁신과"])
    worksheet.append([])
    worksheet.append(["순번", "기준연도", "학과", "재학생(A)", None, None, None])
    worksheet.append([None, None, None, "계", None, "남", None])
    worksheet.append([None, None, None, "정원내", "정원외", "정원내", "정원외"])
    worksheet.append(["합 계", None, None, 3240, 743, 1667, 417])
    worksheet.append([2, "2020", "ICT융합학과", 0, 37, 0, 27])
    workbook.save(path)
    workbook.close()

    result = read_preview(path, "xlsx")

    report = result.inference_report
    assert report is not None
    assert report.file_type == "xlsx"
    assert report.sheet_name == "재적학생"
    assert report.sheet_names == ("재적학생",)
    assert report.header_row_index == 3
    assert report.header_row_count == 3
    assert report.data_start_row_index == 6
    assert report.column_count == 7
    assert report.confidence == "high"
    assert report.requires_user_confirmation is False
    assert "표 헤더 앞의 안내 행 3개를 건너뛰었습니다." in report.reasons
    assert "다중 헤더 3행을 하나의 열 이름으로 합쳤습니다." in report.reasons


def test_read_preview_applies_user_xlsx_sheet_override(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "multi-sheet-public.xlsx"
    workbook = Workbook()
    notice = workbook.active
    notice.title = "안내"
    notice.append(["□ 안내문만 있는 시트입니다."])
    data = workbook.create_sheet("자료")
    data.append(["city", "value"])
    data.append(["Seoul", 10])
    data.append(["Busan", 20])
    workbook.save(path)
    workbook.close()

    result = read_preview(
        path,
        "xlsx",
        layout=TableLayoutOverride(sheet_name="자료"),
    )

    assert result.source.sheet_name == "자료"
    assert result.columns == ("city", "value")
    assert result.sample_rows == (
        {"city": "Seoul", "value": 10},
        {"city": "Busan", "value": 20},
    )
    report = result.inference_report
    assert report is not None
    assert report.sheet_name == "자료"
    assert "사용자 지정 표 레이아웃을 적용했습니다." in report.reasons


def test_read_preview_rejects_missing_xlsx_sheet_override(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "multi-sheet-public.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "자료"
    worksheet.append(["city", "value"])
    worksheet.append(["Seoul", 10])
    workbook.save(path)
    workbook.close()

    with pytest.raises(TableReadError) as exc_info:
        read_preview(
            path,
            "xlsx",
            layout=TableLayoutOverride(sheet_name="없는시트"),
        )

    assert exc_info.value.message_ko == "지정한 시트를 찾지 못했습니다. 시트 이름을 확인해 주세요."


def test_read_preview_reads_tab_delimited_text_with_xls_extension(tmp_path) -> None:
    path = tmp_path / "weather.xls"
    text = "지점\t지점명\t일시\t기온(°C)\n108\t서울\t2026-07-06 01:00\t25.1\n"
    path.write_bytes(text.encode("cp949"))

    result = read_preview(path, "xls", limits=PreviewReadLimits(max_rows=1))

    assert result.columns == ("지점", "지점명", "일시", "기온(°C)")
    assert result.sample_rows == (
        {"지점": 108, "지점명": "서울", "일시": "2026-07-06 01:00", "기온(°C)": 25.1},
    )
    assert "XLS 확장자이지만 텍스트 표로 읽었습니다." in result.warnings
    report = result.inference_report
    assert report is not None
    assert report.file_type == "xls"
    assert report.header_row_index == 0
    assert report.header_row_count == 1
    assert report.data_start_row_index == 1
    assert report.column_count == 4
    assert report.confidence == "high"
    assert "XLS 확장자이지만 텍스트 표로 읽었습니다." in report.reasons


def test_read_preview_uses_excel_reader_for_binary_xls(tmp_path, monkeypatch) -> None:
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    def read_excel_spy(path_arg, *args, **kwargs):
        if kwargs.get("header") is None and kwargs.get("nrows") == 30:
            return pd.DataFrame([["시도", "값"], ["서울", 1]])
        if kwargs.get("header") == 0 and kwargs.get("nrows") == 0:
            return pd.DataFrame(columns=["시도", "값"])
        if kwargs.get("header") == 0 and kwargs.get("nrows") == 1:
            return pd.DataFrame({"시도": ["서울"], "값": [1]})
        return pd.DataFrame({"시도": ["서울"], "값": [1]})

    monkeypatch.setattr("modori.table_io.pd.read_excel", read_excel_spy)

    result = read_preview(path, "xls", limits=PreviewReadLimits(max_rows=1))

    assert result.columns == ("시도", "값")
    assert result.sample_rows == ({"시도": "서울", "값": 1},)
    report = result.inference_report
    assert report is not None
    assert report.file_type == "xls"
    assert report.header_row_index == 0
    assert report.header_row_count == 1
    assert report.data_start_row_index == 1
    assert report.column_count == 2
    assert report.confidence == "high"
    assert report.reasons == ("첫 번째 행을 헤더로 인식했습니다.",)


def test_read_preview_flattens_two_row_binary_xls_headers(tmp_path, monkeypatch) -> None:
    path = tmp_path / "legacy-kosis.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    def read_excel_spy(path_arg, *args, **kwargs):
        if kwargs.get("header") is None and kwargs.get("nrows") == 30:
            return pd.DataFrame(
                [
                    ["행정구역별(1)", "특성별(1)", "2025", "2025"],
                    ["행정구역별(1)", "특성별(1)", "계 (%)", "매우 만족"],
                    ["전국", "전체", 100.0, 11.5],
                    ["서울특별시", "전체", 100.0, 11.0],
                ]
            )
        if kwargs.get("header") is None and kwargs.get("names"):
            return pd.DataFrame(
                [["전국", "전체", 100.0, 11.5]],
                columns=kwargs["names"],
            )
        raise AssertionError(f"unexpected read_excel kwargs: {kwargs}")

    monkeypatch.setattr("modori.table_io.pd.read_excel", read_excel_spy)

    result = read_preview(path, "xls", limits=PreviewReadLimits(max_rows=1))

    assert result.columns == (
        "행정구역별(1)",
        "특성별(1)",
        "2025 계 (%)",
        "2025 매우 만족",
    )
    assert result.sample_rows == (
        {
            "행정구역별(1)": "전국",
            "특성별(1)": "전체",
            "2025 계 (%)": 100.0,
            "2025 매우 만족": 11.5,
        },
    )
    assert "다중 헤더 2행을 하나의 열 이름으로 합쳤습니다." in result.warnings
    report = result.inference_report
    assert report is not None
    assert report.file_type == "xls"
    assert report.header_row_count == 2
    assert report.data_start_row_index == 2
    assert report.column_count == 4
    assert "다중 헤더 2행을 하나의 열 이름으로 합쳤습니다." in report.reasons


def test_read_full_flattens_two_row_binary_xls_headers(tmp_path, monkeypatch) -> None:
    path = tmp_path / "legacy-kosis.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    def read_excel_spy(path_arg, *args, **kwargs):
        if kwargs.get("header") is None and kwargs.get("nrows") == 30:
            return pd.DataFrame(
                [
                    ["행정구역별(1)", "특성별(1)", "2025", "2025"],
                    ["행정구역별(1)", "특성별(1)", "계 (%)", "매우 만족"],
                    ["전국", "전체", 100.0, 11.5],
                ]
            )
        if kwargs.get("header") is None and kwargs.get("names"):
            return pd.DataFrame(
                [["전국", "전체", 100.0, 11.5]],
                columns=kwargs["names"],
            )
        raise AssertionError(f"unexpected read_excel kwargs: {kwargs}")

    monkeypatch.setattr("modori.table_io.pd.read_excel", read_excel_spy)

    result = read_full(path, "xls")

    assert result.columns == (
        "행정구역별(1)",
        "특성별(1)",
        "2025 계 (%)",
        "2025 매우 만족",
    )
    assert result.frame.to_dict(orient="records") == [
        {
            "행정구역별(1)": "전국",
            "특성별(1)": "전체",
            "2025 계 (%)": 100.0,
            "2025 매우 만족": 11.5,
        }
    ]
    assert "다중 헤더 2행을 하나의 열 이름으로 합쳤습니다." in result.warnings


def test_read_preview_rejects_single_notice_cell_xlsx(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "notice-only.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["□ 본 서비스에서 제공하는 정보는 참고용으로만 활용하시기 바랍니다."])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="표 데이터"):
        read_preview(path, "xlsx")


def test_read_preview_rejects_single_header_only_xlsx(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "header-only.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["행정구역별(1)"])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="표 데이터"):
        read_preview(path, "xlsx")


def test_read_full_rejects_csv_when_explicit_row_limit_is_exceeded(tmp_path) -> None:
    path = tmp_path / "survey.csv"
    pd.DataFrame({"a": [1, 2, 3], "b": [4, 5, 6]}).to_csv(path, index=False)

    with pytest.raises(ValueError, match="row limit"):
        read_full(path, "csv", limits=FullReadLimits(max_rows=2))


def test_read_full_rejects_csv_when_explicit_column_limit_is_exceeded(tmp_path) -> None:
    path = tmp_path / "survey.csv"
    pd.DataFrame({"a": [1], "b": [2]}).to_csv(path, index=False)

    with pytest.raises(ValueError, match="column limit"):
        read_full(path, "csv", limits=FullReadLimits(max_columns=1))


def test_read_full_uses_bounded_csv_read_when_row_limit_is_set(
    tmp_path,
    monkeypatch,
) -> None:
    path = tmp_path / "survey.csv"
    path.write_text("a\n1\n", encoding="utf-8")
    calls = []

    def read_csv_spy(path_arg, *args, **kwargs):
        calls.append(kwargs)
        return pd.DataFrame({"a": [1]})

    monkeypatch.setattr("modori.table_io.pd.read_csv", read_csv_spy)

    read_full(path, "csv", limits=FullReadLimits(max_rows=2))

    assert calls == [{"nrows": 3}]


def test_read_full_uses_streaming_xlsx_read_when_row_limit_is_set(
    tmp_path,
    monkeypatch,
) -> None:
    path = tmp_path / "survey.xlsx"
    path.write_bytes(b"placeholder")
    calls = []

    def read_excel_spy(*args, **kwargs):
        raise AssertionError("pandas.read_excel should not be used for bounded xlsx")

    class FakeWorksheet:
        def iter_rows(self, *, values_only):
            assert values_only is True
            return iter(
                [
                    ("a", "b"),
                    (1, 2),
                    (3, 4),
                    (5, 6),
                ]
            )

    class FakeWorkbook:
        active = FakeWorksheet()

        def close(self):
            calls.append("closed")

    def load_workbook_spy(path_arg, *args, **kwargs):
        calls.append(
            {
                "path": path_arg,
                "read_only": kwargs.get("read_only"),
                "data_only": kwargs.get("data_only"),
            }
        )
        return FakeWorkbook()

    monkeypatch.setattr("modori.table_io.pd.read_excel", read_excel_spy)
    monkeypatch.setattr("openpyxl.load_workbook", load_workbook_spy)

    with pytest.raises(ValueError, match="row limit"):
        read_full(path, "xlsx", limits=FullReadLimits(max_rows=2))

    assert calls == [
        {"path": path, "read_only": True, "data_only": True},
        "closed",
    ]


def test_read_full_passes_sav_row_limit_when_row_limit_is_set(
    tmp_path,
    monkeypatch,
) -> None:
    path = tmp_path / "survey.sav"
    path.write_bytes(b"placeholder")
    calls = []

    def read_sav_spy(path_arg, *args, **kwargs):
        calls.append({"path": path_arg, **kwargs})
        return pd.DataFrame({"a": [1, 2]}), object()

    monkeypatch.setattr(pyreadstat, "read_sav", read_sav_spy)

    read_full(path, "sav", limits=FullReadLimits(max_rows=2))

    assert calls == [{"path": path, "row_limit": 3, "user_missing": True}]


def test_preview_inference_report_exposes_leading_rows_for_preamble_csv(tmp_path) -> None:
    path = tmp_path / "preamble.csv"
    lines = [f"검색조건 {index}: 값" for index in range(3)]
    lines.extend(["지역,인구", "종로구,100", "중구,200"])
    path.write_text("\n".join(lines), encoding="utf-8")

    report = read_preview(path, "csv").inference_report

    assert report is not None
    assert report.header_row_index == 3
    assert len(report.leading_rows) == 6
    assert report.leading_rows[0] == ("검색조건 0: 값",)
    assert report.leading_rows[3] == ("지역", "인구")
    assert report.leading_rows[4] == ("종로구", "100")


def test_preview_inference_leading_rows_are_capped(tmp_path) -> None:
    path = tmp_path / "wide.csv"
    header = ",".join(f"c{index}" for index in range(12))
    long_cell = "x" * 200
    row = ",".join([long_cell] + [str(index) for index in range(11)])
    path.write_text("\n".join([header] + [row] * 30), encoding="utf-8")

    report = read_preview(path, "csv").inference_report

    assert report is not None
    assert len(report.leading_rows) == 4
    assert all(len(cells) <= 8 for cells in report.leading_rows)
    assert len(report.leading_rows[1][0]) == 60


def test_preview_inference_report_exposes_leading_rows_for_merged_xlsx(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["대학 현황"])
    sheet.append(["기준일: 2026-04-01"])
    sheet.append([])
    sheet.append(["순번", "학과", "재학생(A)", "재학생(A)"])
    sheet.append([None, None, "정원내", "정원외"])
    sheet.append([1, "사회학과", 40, 5])
    sheet.append([2, "심리학과", 35, 4])
    workbook.save(path)

    report = read_preview(path, "xlsx").inference_report

    assert report is not None
    assert report.header_row_index == 3
    assert report.header_row_count == 2
    assert report.data_start_row_index == 5
    assert len(report.leading_rows) == 7
    assert report.leading_rows[0] == ("대학 현황",)
    assert report.leading_rows[3] == ("순번", "학과", "재학생(A)", "재학생(A)")
    assert report.leading_rows[5] == ("1", "사회학과", "40", "5")
