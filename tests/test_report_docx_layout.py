from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt

from modori.steps.reporting import write_docx


def test_write_docx_makes_wide_result_tables_readable(tmp_path: Path) -> None:
    output_path = tmp_path / "wide-report.docx"
    write_docx(
        prose=["One correlation pair was reviewed."],
        tables={
            "correlation": [
                {
                    "x": "Final grade",
                    "y": "Weekly study time",
                    "method": "spearman",
                    "statistic": "rho",
                    "coefficient": "0.275",
                    "p_value": "0.000",
                    "p_adjusted": "",
                    "n": "649",
                    "excluded_n": "0",
                    "warnings": "Tied ranks were handled by the calculation engine.",
                }
            ]
        },
        figure_paths={},
        path=output_path,
    )

    document = Document(output_path)
    section = document.sections[0]
    table = document.tables[0]

    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width > section.page_height
    assert section.left_margin <= Inches(0.5)
    assert section.right_margin <= Inches(0.5)
    assert table.autofit is False

    body_sizes = {
        run.font.size
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.text
    }
    assert body_sizes == {Pt(8)}

    grid_widths = [column.w for column in table._tbl.tblGrid.gridCol_lst]
    warnings_index = len(grid_widths) - 1
    n_index = 7
    assert grid_widths[warnings_index] > grid_widths[n_index]
