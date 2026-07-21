from __future__ import annotations

import math
import os
import re
import uuid
from collections.abc import Callable, Iterable, Mapping
from dataclasses import dataclass
from pathlib import Path

from modori.cache import matplotlib_cache_dir

_CACHE_DIR = matplotlib_cache_dir()
os.environ["MPLCONFIGDIR"] = str(_CACHE_DIR)

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import pandas as pd
from docx import Document
from scipy import stats

from modori.core import PipelineContext, Step, StepResult
from modori.descriptives_table1_reporting import (
    prose_for_descriptives,
    table_for_descriptives,
)
from modori.descriptives_table1_results import DescriptivesTableResult
from modori.frequency_crosstab_reporting import (
    prose_for_frequency_crosstab,
    table_for_frequency_crosstab,
)
from modori.frequency_crosstab_results import FrequencyCrosstabResult
from modori.correlation_reporting import prose_for_correlation, table_for_correlation
from modori.correlation_results import CorrelationResult
from modori.anova_oneway_reporting import (
    prose_for_anova_oneway,
    table_for_anova_oneway,
)
from modori.anova_oneway_results import OneWayAnovaResult
from modori.friedman_reporting import prose_for_friedman, table_for_friedman
from modori.friedman_results import FriedmanResult
from modori.kruskal_wallis_reporting import (
    prose_for_kruskal_wallis,
    table_for_kruskal_wallis,
)
from modori.kruskal_wallis_results import KruskalWallisResult
from modori.mediation_reporting import prose_for_mediation, table_for_mediation
from modori.mediation_results import MediationResult
from modori.moderated_mediation_reporting import (
    prose_for_moderated_mediation,
    table_for_moderated_mediation,
)
from modori.moderated_mediation_results import ModeratedMediationResult
from modori.repeated_measures_anova_reporting import (
    prose_for_repeated_measures_anova,
    table_for_repeated_measures_anova,
)
from modori.repeated_measures_anova_results import RepeatedMeasuresAnovaResult
from modori.ancova_reporting import prose_for_ancova, table_for_ancova
from modori.ancova_results import AncovaResult
from modori.factor_pca_reporting import (
    component_table_for_factor_pca,
    loading_table_for_factor_pca,
    prose_for_factor_pca,
)
from modori.factor_pca_results import FactorPcaResult
from modori.factorial_anova_reporting import (
    prose_for_factorial_anova,
    table_for_factorial_anova,
)
from modori.factorial_anova_results import FactorialAnovaResult
from modori.logistic_regression_reporting import (
    prose_for_logistic,
    table_for_logistic,
)
from modori.logistic_regression_results import LogisticRegressionResult
from modori.results import (
    ChartSpec,
    CoefficientRow,
    ComparisonResult,
    RegressionResult,
    ReliabilityResult,
    ReportResult,
)


_FONT_CONFIGURED = False


def _configure_fonts() -> None:
    global _FONT_CONFIGURED
    if _FONT_CONFIGURED:
        return

    candidates = [
        "Malgun Gothic",
        "AppleGothic",
        "Noto Sans CJK KR",
        "NanumGothic",
        "DejaVu Sans",
    ]
    available = {font.name for font in font_manager.fontManager.ttflist}
    for name in candidates:
        if name in available:
            plt.rcParams["font.family"] = [name]
            break
    plt.rcParams["axes.unicode_minus"] = False
    _FONT_CONFIGURED = True


def _apa_number(
    value: float, digits: int = 2, omit_leading_zero: bool = True, **kwargs: object
) -> str:
    """Format a number with APA-style optional leading-zero omission."""

    if "decimals" in kwargs:
        digits = int(kwargs.pop("decimals"))
    if kwargs:
        unexpected = ", ".join(sorted(kwargs))
        raise TypeError(f"Unexpected keyword argument(s): {unexpected}")
    text = f"{float(value):.{digits}f}"
    if omit_leading_zero:
        text = re.sub(r"^(-?)0\.", r"\1.", text)
    return text


def _apa_p(value: float) -> str:
    p_value = float(value)
    if p_value < 0.001:
        return "p < .001"
    return f"p = {_apa_number(p_value, 3)}"


def _safe_chart_key(key: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", key).strip("._-")
    return safe or "chart"


def _alpha_qualifier_en(alpha: float) -> str:
    if alpha >= 0.90:
        return "excellent"
    if alpha >= 0.80:
        return "good"
    if alpha >= 0.70:
        return "acceptable"
    if alpha >= 0.60:
        return "questionable"
    return "low"


def _alpha_qualifier_ko(alpha: float) -> str:
    if alpha >= 0.90:
        return "매우 높은"
    if alpha >= 0.80:
        return "높은"
    if alpha >= 0.70:
        return "수용 가능한"
    if alpha >= 0.60:
        return "다소 낮은"
    return "낮은"


def _reliability_prose(result: ReliabilityResult, language: str = "ko") -> str:
    alpha = _apa_number(result.cronbach_alpha, omit_leading_zero=True)
    omega = _apa_number(result.mcdonald_omega, omit_leading_zero=True)
    if language == "en":
        return (
            f"The {result.scale_name} scale showed {_alpha_qualifier_en(result.cronbach_alpha)} internal consistency "
            f"(Cronbach's α = {alpha}, McDonald's ω = {omega})."
        )
    scale_name = (
        "선택한 문항" if result.scale_name == "selected_scale" else result.scale_name
    )
    return (
        f"{scale_name} 척도는 {_alpha_qualifier_ko(result.cronbach_alpha)} 내적 일관성을 보였다"
        f"(Cronbach's α = {alpha}, McDonald's ω = {omega})."
    )


def _comparison_groups(result: ComparisonResult) -> list[tuple[str, object]]:
    return list(result.groups.items())


def _comparison_mean_difference(result: ComparisonResult) -> float:
    groups = _comparison_groups(result)
    if len(groups) < 2:
        return float("nan")
    return float(groups[0][1].mean - groups[1][1].mean)


def _comparison_effect_label(result: ComparisonResult) -> str:
    if result.effect_name == "cohen_d":
        return "Cohen's d"
    if result.effect_name == "cohen_dz":
        return "Cohen's dz"
    if result.effect_name == "hedges_g":
        return "Hedges' g"
    if result.effect_name == "eta_squared":
        return "eta squared"
    if result.effect_name == "rank_biserial":
        return "rank-biserial r"
    return result.effect_name


def _comparison_test_label_ko(result: ComparisonResult) -> str:
    if result.test_name == "welch_t":
        return "독립표본 t검정(Welch 보정)"
    if result.test_name == "student_t":
        return "독립표본 t검정"
    if result.test_name == "paired_t":
        return "대응표본 t검정"
    if result.test_name == "wilcoxon":
        return "Wilcoxon 부호순위검정"
    if result.test_name == "mann_whitney":
        return "Mann-Whitney U 검정"
    if result.test_name == "anova_oneway":
        return "일원분산분석"
    return result.test_name


def _comparison_test_label_en(result: ComparisonResult) -> str:
    if result.test_name == "welch_t":
        return "Welch independent-samples t test"
    if result.test_name == "student_t":
        return "independent-samples t test"
    if result.test_name == "paired_t":
        return "paired-samples t test"
    if result.test_name == "wilcoxon":
        return "Wilcoxon signed-rank test"
    if result.test_name == "mann_whitney":
        return "Mann-Whitney U test"
    if result.test_name == "anova_oneway":
        return "one-way ANOVA"
    return result.test_name


def _comparison_significance_ko(result: ComparisonResult) -> str:
    if result.p_value < 0.05:
        return "통계적으로 유의하였다"
    return "통계적으로 유의하지 않았다"


def _comparison_significance_en(result: ComparisonResult) -> str:
    if result.p_value < 0.05:
        return "a statistically significant"
    return "no statistically significant"


def _comparison_case_count_sentence(result: ComparisonResult, language: str) -> str:
    if language == "en":
        excluded_case = "case" if result.n_dropped == 1 else "cases"
        excluded_verb = "was" if result.n_dropped == 1 else "were"
        return (
            f" The analysis used {result.n_obs} cases; "
            f"{result.n_dropped} {excluded_case} {excluded_verb} excluded for missing values."
        )
    return f" 분석에는 {result.n_obs}명이 사용되었고 {result.n_dropped}명은 결측으로 제외되었다."


def _comparison_paired_labels(result: ComparisonResult) -> tuple[str, str]:
    groups = _comparison_groups(result)
    before_fallback = str(groups[0][0]) if len(groups) > 0 else "before"
    after_fallback = str(groups[1][0]) if len(groups) > 1 else "after"
    return (
        result.before_label or before_fallback,
        result.after_label or after_fallback,
    )


def _comparison_prose(result: ComparisonResult, language: str = "ko") -> str:
    p_text = _apa_p(result.p_value)
    effect = f"{_comparison_effect_label(result)} = {_apa_number(result.effect_value, omit_leading_zero=True)}"
    case_count_sentence = _comparison_case_count_sentence(result, language)
    ci_text = ""
    if result.mean_diff_ci is not None:
        ci_text = f", 95% CI [{_apa_number(result.mean_diff_ci[0])}, {_apa_number(result.mean_diff_ci[1])}]"
    groups = _comparison_groups(result)
    first_label, first_group = groups[0] if len(groups) > 0 else ("group 1", None)
    second_label, second_group = groups[1] if len(groups) > 1 else ("group 2", None)
    first_desc = ""
    second_desc = ""
    if first_group is not None:
        first_desc = (
            f"{first_label}(M = {_apa_number(first_group.mean)}, "
            f"SD = {_apa_number(first_group.sd)}, n = {first_group.n})"
        )
    if second_group is not None:
        second_desc = (
            f"{second_label}(M = {_apa_number(second_group.mean)}, "
            f"SD = {_apa_number(second_group.sd)}, n = {second_group.n})"
        )

    if result.test_name == "paired_t":
        before_label, after_label = _comparison_paired_labels(result)
        df_text = _apa_number(result.df, 2) if result.df is not None else ""
        stat_text = (
            f"t({df_text}) = {_apa_number(result.statistic)}"
            if df_text
            else f"t = {_apa_number(result.statistic)}"
        )
        if language == "en":
            return (
                f"A {_comparison_test_label_en(result)} showed {_comparison_significance_en(result)} "
                f"mean difference between {before_label} and {after_label}, "
                f"{stat_text}, {p_text}{ci_text}, {effect}.{case_count_sentence}"
            )
        return (
            f"{_comparison_test_label_ko(result)} 결과, {before_label}와 {after_label}의 평균 차이는 "
            f"{_comparison_significance_ko(result)}, {stat_text}, {p_text}{ci_text}, "
            f"{effect}.{case_count_sentence}"
        )

    if result.test_name == "wilcoxon":
        before_label, after_label = _comparison_paired_labels(result)
        stat_text = f"W = {_apa_number(result.statistic)}"
        if language == "en":
            return (
                f"A {_comparison_test_label_en(result)} showed {_comparison_significance_en(result)} "
                f"difference between {before_label} and {after_label}, "
                f"{stat_text}, {p_text}, {effect}.{case_count_sentence}"
            )
        return (
            f"{_comparison_test_label_ko(result)} 결과, {before_label}와 {after_label}의 차이는 "
            f"{_comparison_significance_ko(result)}, {stat_text}, {p_text}, "
            f"{effect}.{case_count_sentence}"
        )

    if result.test_name == "mann_whitney":
        stat_text = f"U = {_apa_number(result.statistic)}"
        dv_label = result.dv_label or result.dv
        if language == "en":
            if first_desc and second_desc:
                return (
                    f"A {_comparison_test_label_en(result)} showed {_comparison_significance_en(result)} "
                    f"{dv_label} distribution difference between {first_desc} and {second_desc}, "
                    f"{stat_text}, {p_text}, {effect}.{case_count_sentence}"
                )
            return (
                f"A {_comparison_test_label_en(result)} showed {_comparison_significance_en(result)} "
                f"{dv_label} distribution difference, {stat_text}, {p_text}, "
                f"{effect}.{case_count_sentence}"
            )
        group_text = (
            f"{first_desc}와 {second_desc}의 " if first_desc and second_desc else ""
        )
        return (
            f"{_comparison_test_label_ko(result)} 결과, {group_text}{dv_label} 분포 차이는 "
            f"{_comparison_significance_ko(result)}, {stat_text}, {p_text}, "
            f"{effect}.{case_count_sentence}"
        )

    if result.test_name in {"student_t", "welch_t"}:
        df_text = _apa_number(result.df, 2) if result.df is not None else ""
        stat_text = (
            f"t({df_text}) = {_apa_number(result.statistic)}"
            if df_text
            else f"t = {_apa_number(result.statistic)}"
        )
        mean_diff = _comparison_mean_difference(result)
        if language == "en":
            test_label = _comparison_test_label_en(result)
            article = "A" if test_label.startswith("Welch") else "An"
            significance = (
                "a statistically significant"
                if result.p_value < 0.05
                else "no statistically significant"
            )
            if first_desc and second_desc:
                return (
                    f"{article} {test_label} showed {significance} {result.dv_label or result.dv} score difference "
                    f"between {first_desc} and {second_desc}, {stat_text}, {p_text}{ci_text}, {effect}."
                    f"{case_count_sentence}"
                )
            return (
                f"{article} {test_label} showed {significance} {result.dv_label or result.dv} score difference, "
                f"{stat_text}, {p_text}{ci_text}, {effect}.{case_count_sentence}"
            )
        significance = (
            "통계적으로 유의하였다"
            if result.p_value < 0.05
            else "통계적으로 유의하지 않았다"
        )
        group_text = (
            f"{first_desc}와 {second_desc}의 " if first_desc and second_desc else ""
        )
        return (
            f"{_comparison_test_label_ko(result)} 결과, {group_text}{result.dv_label or result.dv} 점수 차이"
            f"(Mdiff = {_apa_number(mean_diff)})는 {significance}, "
            f"{stat_text}, {p_text}{ci_text}, {effect}.{case_count_sentence}"
        )

    if result.test_name == "anova_oneway":
        stat_text = f"F = {_apa_number(result.statistic)}"
        if language == "en":
            return (
                f"A {_comparison_test_label_en(result)} was {'significant' if result.p_value < 0.05 else 'not significant'}, "
                f"{stat_text}, {p_text}, {effect}.{case_count_sentence}"
            )
        return (
            f"{_comparison_test_label_ko(result)} 결과는 "
            f"{'통계적으로 유의하였다' if result.p_value < 0.05 else '통계적으로 유의하지 않았다'}, "
            f"{stat_text}, {p_text}, {effect}.{case_count_sentence}"
        )

    raise ValueError(f"Unsupported comparison test: {result.test_name}")


def _regression_predictor_rows(result: RegressionResult) -> list[CoefficientRow]:
    return [
        row
        for row in result.coefficients
        if row.name not in {"const", "(Intercept)", "Intercept"}
    ]


def _regression_se_label_ko(result: RegressionResult) -> str:
    if result.se_type.upper() == "HC3":
        return "이분산-강건(HC3) 표준오차"
    return f"{result.se_type} 표준오차"


def _regression_se_label_en(result: RegressionResult) -> str:
    if result.se_type.upper() == "HC3":
        return "HC3 robust standard errors"
    return f"{result.se_type} standard errors"


def _regression_row_text_ko(row: CoefficientRow, df_resid: int) -> str:
    significance = "유의하게 예측하였다" if row.p_value < 0.05 else "유의하지 않았다"
    beta = (
        ""
        if row.beta is None
        else f", β = {_apa_number(row.beta, omit_leading_zero=True)}"
    )
    return (
        f"{row.name}는 {significance}"
        f"(b = {_apa_number(row.b)}, SE = {_apa_number(row.se)}, "
        f"t({df_resid}) = {_apa_number(row.t)}, {_apa_p(row.p_value)}{beta})"
    )


def _regression_row_text_en(row: CoefficientRow, df_resid: int, dv: str) -> str:
    beta = (
        ""
        if row.beta is None
        else f", beta = {_apa_number(row.beta, omit_leading_zero=True)}"
    )
    stats_text = (
        f"b = {_apa_number(row.b)}, SE = {_apa_number(row.se)}, "
        f"t({df_resid}) = {_apa_number(row.t)}, {_apa_p(row.p_value)}{beta}"
    )
    if row.p_value < 0.05:
        return f"{row.name} significantly predicted {dv} ({stats_text})"
    return f"{row.name} was not significant ({stats_text})"


def _regression_prose(result: RegressionResult, language: str = "ko") -> str:
    omnibus = f"F({result.df_model}, {result.df_resid}) = {_apa_number(result.f_statistic)}, {_apa_p(result.f_p_value)}"
    model_fit = (
        f"R² = {_apa_number(result.r_squared, omit_leading_zero=True)}, "
        f"adjusted R² = {_apa_number(result.adj_r_squared, omit_leading_zero=True)}"
    )
    predictors = _regression_predictor_rows(result)

    if language == "en":
        predictor_text = "; ".join(
            _regression_row_text_en(row, result.df_resid, result.dv)
            for row in predictors
        )
        return (
            f"Using {_regression_se_label_en(result)}, the regression model was statistically significant, "
            f"{omnibus}, with {model_fit}. {predictor_text}."
        )

    predictor_text = "; ".join(
        _regression_row_text_ko(row, result.df_resid) for row in predictors
    )
    return (
        f"{_regression_se_label_ko(result)}를 사용한 회귀모형은 통계적으로 유의하였다, "
        f"{omnibus}, {model_fit}. {predictor_text}."
    )


def prose_for(result: object, language: str = "ko") -> str:
    if isinstance(result, DescriptivesTableResult):
        return prose_for_descriptives(result, language=language)
    if isinstance(result, FrequencyCrosstabResult):
        return prose_for_frequency_crosstab(result, language=language)
    if isinstance(result, CorrelationResult):
        return prose_for_correlation(result, language=language)
    if isinstance(result, OneWayAnovaResult):
        return prose_for_anova_oneway(result, language=language)
    if isinstance(result, FriedmanResult):
        return prose_for_friedman(result, language=language)
    if isinstance(result, KruskalWallisResult):
        return prose_for_kruskal_wallis(result, language=language)
    if isinstance(result, MediationResult):
        return prose_for_mediation(result, language=language)
    if isinstance(result, ModeratedMediationResult):
        return prose_for_moderated_mediation(result, language=language)
    if isinstance(result, RepeatedMeasuresAnovaResult):
        return prose_for_repeated_measures_anova(result, language=language)
    if isinstance(result, AncovaResult):
        return prose_for_ancova(result, language=language)
    if isinstance(result, FactorPcaResult):
        return prose_for_factor_pca(result, language=language)
    if isinstance(result, FactorialAnovaResult):
        return prose_for_factorial_anova(result, language=language)
    if isinstance(result, LogisticRegressionResult):
        return prose_for_logistic(result, language=language)
    if isinstance(result, RegressionResult):
        return _regression_prose(result, language=language)
    if isinstance(result, ReliabilityResult):
        return _reliability_prose(result, language=language)
    if isinstance(result, ComparisonResult):
        return _comparison_prose(result, language=language)
    raise TypeError(f"Unsupported result for prose: {type(result).__name__}")


def table_for(result: object) -> list[dict[str, str]]:
    if isinstance(result, DescriptivesTableResult):
        return table_for_descriptives(result)

    if isinstance(result, FrequencyCrosstabResult):
        return table_for_frequency_crosstab(result)

    if isinstance(result, CorrelationResult):
        return table_for_correlation(result)

    if isinstance(result, OneWayAnovaResult):
        return table_for_anova_oneway(result)

    if isinstance(result, FriedmanResult):
        return table_for_friedman(result)

    if isinstance(result, KruskalWallisResult):
        return table_for_kruskal_wallis(result)

    if isinstance(result, MediationResult):
        return table_for_mediation(result)

    if isinstance(result, ModeratedMediationResult):
        return table_for_moderated_mediation(result)

    if isinstance(result, RepeatedMeasuresAnovaResult):
        return table_for_repeated_measures_anova(result)

    if isinstance(result, AncovaResult):
        return table_for_ancova(result)

    if isinstance(result, FactorPcaResult):
        return [
            *component_table_for_factor_pca(result),
            *loading_table_for_factor_pca(result),
        ]

    if isinstance(result, FactorialAnovaResult):
        return table_for_factorial_anova(result)

    if isinstance(result, LogisticRegressionResult):
        return table_for_logistic(result)

    if isinstance(result, RegressionResult):
        return [
            {
                "predictor": row.name,
                "b": _apa_number(row.b),
                "SE": _apa_number(row.se),
                "t": _apa_number(row.t),
                "p": _apa_p(row.p_value),
                "beta": ""
                if row.beta is None
                else _apa_number(row.beta, omit_leading_zero=True),
                "95% CI": f"[{_apa_number(row.ci[0])}, {_apa_number(row.ci[1])}]",
                "VIF": ""
                if row.vif is None
                else _apa_number(row.vif, omit_leading_zero=False),
                "vif": ""
                if row.vif is None
                else _apa_number(row.vif, omit_leading_zero=False),
            }
            for row in result.coefficients
        ]

    if isinstance(result, ReliabilityResult):
        return [
            {
                "item": item,
                "item_total_corr": _apa_number(
                    result.item_total_corr[item], omit_leading_zero=True
                ),
                "alpha_if_deleted": _apa_number(
                    result.alpha_if_deleted[item], omit_leading_zero=True
                ),
            }
            for item in result.item_total_corr
        ]

    if isinstance(result, ComparisonResult):
        groups = _comparison_groups(result)
        first_label = groups[0][0] if len(groups) > 0 else ""
        second_label = groups[1][0] if len(groups) > 1 else ""
        before_label = result.before_label or (
            str(first_label) if result.paired else ""
        )
        after_label = result.after_label or (str(second_label) if result.paired else "")
        return [
            {
                "test": result.test_name,
                "dv": result.dv,
                "group": result.group_var,
                "group_1": str(first_label),
                "group_2": str(second_label),
                "before": before_label,
                "after": after_label,
                "statistic": _apa_number(result.statistic),
                "df": _apa_number(result.df, 2) if result.df is not None else "",
                "p": _apa_p(result.p_value),
                "effect": result.effect_name,
                "effect_value": _apa_number(
                    result.effect_value, omit_leading_zero=True
                ),
                "n_obs": str(result.n_obs),
                "n_total": str(result.n_total),
                "n_dropped": str(result.n_dropped),
                "ci95": f"[{_apa_number(result.mean_diff_ci[0])}, {_apa_number(result.mean_diff_ci[1])}]"
                if result.mean_diff_ci is not None
                else "",
            }
        ]

    raise TypeError(f"Unsupported result for table: {type(result).__name__}")


def _chart_values_from_mapping(values: object) -> tuple[list[str], list[float]]:
    if isinstance(values, dict):
        return list(values.keys()), [float(value) for value in values.values()]
    if isinstance(values, Iterable) and not isinstance(values, (str, bytes)):
        numeric = [float(value) for value in values]
        return [str(index + 1) for index in range(len(numeric))], numeric
    return [], []


def _render_legacy_chart(chart: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()

    data = chart.data
    fig, ax = plt.subplots(figsize=(7, 4.5), constrained_layout=True)

    if chart.type == "horizontal_bar":
        labels = list(data.get("labels", []))
        values = list(data.get("values", []))
        if not labels:
            labels, values = _chart_values_from_mapping(data.get("values", {}))
        ax.barh(labels, values, color="#4C78A8")
        ax.set_xlabel(chart.x_label)
        ax.set_ylabel(chart.y_label)
        ax.set_title(chart.title)
        ax.invert_yaxis()
    elif chart.type == "mean_ci_jitter":
        groups = data.get("groups", [])
        for idx, group in enumerate(groups):
            values = group.get("values", [])
            if values:
                jitter_x = [
                    idx + (i - len(values) / 2) * 0.015 for i in range(len(values))
                ]
                ax.scatter(jitter_x, values, alpha=0.45, color="#72B7B2")
            mean = group.get("mean")
            ci = group.get("ci95")
            if mean is not None:
                ax.scatter([idx], [mean], color="#F58518", zorder=3)
            if ci is not None:
                ax.vlines(idx, ci[0], ci[1], color="#F58518", linewidth=3)
        ax.set_xticks(range(len(groups)))
        ax.set_xticklabels(
            [g.get("label", f"group {i + 1}") for i, g in enumerate(groups)]
        )
        ax.set_ylabel(chart.y_label)
        ax.set_title(chart.title)
    elif chart.type == "box":
        groups = data.get("groups", [])
        values = [g.get("values", []) for g in groups]
        labels = [g.get("label", f"group {i + 1}") for i, g in enumerate(groups)]
        ax.boxplot(values, labels=labels, patch_artist=True)
        ax.set_ylabel(chart.y_label)
        ax.set_title(chart.title)
    else:
        plt.close(fig)
        raise ValueError(f"Unsupported chart type: {chart.type}")

    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _render_coefficient_forest(spec: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()
    rows = list(spec.data.get("rows", []))
    if not rows:
        raise ValueError(
            "coefficient_forest chart requires at least one coefficient row"
        )

    labels = [str(row.get("name", row.get("term", ""))) for row in rows]
    estimates = [float(row.get("beta", row.get("estimate", 0.0))) for row in rows]
    ci_values = [row.get("ci", (row.get("ci_low"), row.get("ci_high"))) for row in rows]
    ci_low = [float(ci[0]) for ci in ci_values]
    ci_high = [float(ci[1]) for ci in ci_values]
    lower_errors = [estimate - low for estimate, low in zip(estimates, ci_low)]
    upper_errors = [high - estimate for estimate, high in zip(estimates, ci_high)]

    fig_height = max(3.5, 0.55 * len(rows) + 1.4)
    fig, ax = plt.subplots(figsize=(7.2, fig_height), constrained_layout=True)
    y_positions = list(range(len(rows)))
    ax.errorbar(
        estimates,
        y_positions,
        xerr=[lower_errors, upper_errors],
        fmt="o",
        color="#1f77b4",
        ecolor="#555555",
        capsize=4,
    )
    ax.axvline(0, color="#999999", linestyle="--", linewidth=1)
    ax.set_yticks(y_positions)
    ax.set_yticklabels(labels)
    ax.invert_yaxis()
    ax.set_xlabel(spec.x_label)
    ax.set_ylabel(spec.y_label)
    ax.set_title(spec.title)
    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _render_diagnostic_chart(spec: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()
    data = spec.data
    fig, ax = plt.subplots(figsize=(7, 4.5), constrained_layout=True)

    if spec.type == "residual_vs_fitted":
        fitted = [float(value) for value in data.get("fitted", [])]
        residuals = [float(value) for value in data.get("residuals", [])]
        ax.scatter(fitted, residuals, alpha=0.75, color="#4C78A8")
        ax.axhline(0, color="#999999", linestyle="--", linewidth=1)
        ax.set_xlabel(spec.x_label)
        ax.set_ylabel(spec.y_label)
        ax.set_title(spec.title)
    elif spec.type == "residual_qq":
        if "theoretical" in data and "sample" in data:
            theoretical = [float(value) for value in data.get("theoretical", [])]
            sample = [float(value) for value in data.get("sample", [])]
            ax.scatter(theoretical, sample, alpha=0.75, color="#4C78A8")
            if theoretical and sample:
                slope, intercept = (
                    pd.Series(sample).cov(pd.Series(theoretical))
                    / pd.Series(theoretical).var(),
                    pd.Series(sample).mean(),
                )
                ax.plot(
                    theoretical,
                    [slope * x + intercept for x in theoretical],
                    color="#F58518",
                    linewidth=2,
                )
        else:
            residuals = [float(value) for value in data.get("residuals", [])]
            if not residuals:
                plt.close(fig)
                raise ValueError("residual_qq chart requires residuals")
            (osm, osr), (slope, intercept, _r) = stats.probplot(residuals, dist="norm")
            ax.scatter(osm, osr, alpha=0.75, color="#4C78A8")
            ax.plot(osm, slope * osm + intercept, color="#F58518", linewidth=2)
        ax.set_xlabel(spec.x_label)
        ax.set_ylabel(spec.y_label)
        ax.set_title(spec.title)
    elif spec.type == "cooks_distance":
        cooks = [
            float(value) for value in data.get("cooks", data.get("cooks_distance", []))
        ]
        index = [int(value) for value in data.get("index", range(1, len(cooks) + 1))]
        ax.bar(index, cooks, color="#4C78A8")
        threshold = data.get("threshold")
        if threshold is not None:
            ax.axhline(float(threshold), color="#F58518", linestyle="--", linewidth=1)
        ax.set_xlabel(spec.x_label)
        ax.set_ylabel(spec.y_label)
        ax.set_title(spec.title)
    else:
        plt.close(fig)
        raise ValueError(f"Unsupported diagnostic chart type: {spec.type}")

    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _render_paired_line_chart(spec: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()

    before_values = [float(value) for value in spec.data.get("before", [])]
    after_values = [float(value) for value in spec.data.get("after", [])]
    if not before_values or len(before_values) != len(after_values):
        raise ValueError(
            "paired_line chart requires equal-length before and after values"
        )

    before_label = str(spec.data.get("before_label", "Before"))
    after_label = str(spec.data.get("after_label", "After"))
    mean_before = sum(before_values) / len(before_values)
    mean_after = sum(after_values) / len(after_values)

    fig, ax = plt.subplots(figsize=(7, 4.5), constrained_layout=True)
    for before, after in zip(before_values, after_values):
        ax.plot([0, 1], [before, after], color="#4C78A8", alpha=0.25, linewidth=1)

    ax.scatter(
        [0] * len(before_values),
        before_values,
        color="#4C78A8",
        alpha=0.45,
        s=22,
    )
    ax.scatter(
        [1] * len(after_values),
        after_values,
        color="#4C78A8",
        alpha=0.45,
        s=22,
    )
    ax.plot([0, 1], [mean_before, mean_after], color="#F58518", linewidth=2.5)
    ax.scatter([0, 1], [mean_before, mean_after], color="#F58518", s=55, zorder=3)
    ax.set_xticks([0, 1])
    ax.set_xticklabels([before_label, after_label])
    ax.set_xlim(-0.2, 1.2)
    ax.set_xlabel(spec.x_label)
    ax.set_ylabel(spec.y_label)
    ax.set_title(spec.title)

    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _factorial_chart_sequence(value: object, label: str) -> list[object]:
    if not isinstance(value, Iterable) or isinstance(value, (str, bytes, Mapping)):
        raise ValueError(f"factorial_interaction {label} must be a sequence")
    return list(value)


def _factorial_chart_values(value: object, label: str) -> list[float]:
    raw = _factorial_chart_sequence(value, label)
    values: list[float] = []
    for item in raw:
        if isinstance(item, bool):
            raise ValueError(
                f"factorial_interaction {label} must be finite numeric data"
            )
        try:
            numeric = float(item)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"factorial_interaction {label} must be finite numeric data"
            ) from exc
        if not math.isfinite(numeric):
            raise ValueError(
                f"factorial_interaction {label} must be finite numeric data"
            )
        values.append(numeric)
    return values


def _render_factorial_interaction(spec: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()
    if not isinstance(spec.data, Mapping):
        raise ValueError("factorial_interaction data must be a mapping")
    factor_a = _factorial_chart_sequence(spec.data.get("factor_a"), "factor_a")
    series = _factorial_chart_sequence(spec.data.get("series"), "series")
    factor_b_label = str(spec.data.get("factor_b_label", "")).strip()
    if not factor_b_label:
        raise ValueError("factorial_interaction requires a factor B label")
    if not 2 <= len(factor_a) <= 6 or not 2 <= len(series) <= 6:
        raise ValueError(
            "factorial_interaction requires 2 through 6 levels for both factors"
        )

    x_labels: list[str] = []
    x_tokens: list[str] = []
    for level in factor_a:
        if not isinstance(level, Mapping):
            raise ValueError("factorial_interaction factor_a levels must be mappings")
        token = str(level.get("token", "")).strip()
        label = str(level.get("label", "")).strip()
        if not token or not label:
            raise ValueError(
                "factorial_interaction factor_a levels require token and label"
            )
        x_tokens.append(token)
        x_labels.append(label)
    if len(set(x_tokens)) != len(x_tokens) or len(set(x_labels)) != len(x_labels):
        raise ValueError("factorial_interaction factor_a levels must be unique")

    validated: list[tuple[str, list[float], list[float], list[float]]] = []
    series_tokens: list[str] = []
    series_labels: list[str] = []
    expected_length = len(factor_a)
    for item in series:
        if not isinstance(item, Mapping):
            raise ValueError("factorial_interaction series entries must be mappings")
        token = str(item.get("factor_b_token", "")).strip()
        label = str(item.get("factor_b_label", "")).strip()
        if not token or not label:
            raise ValueError(
                "factorial_interaction series require factor B token and label"
            )
        means = _factorial_chart_values(item.get("means"), "means")
        lower = _factorial_chart_values(item.get("ci_low"), "ci_low")
        upper = _factorial_chart_values(item.get("ci_high"), "ci_high")
        if not (len(means) == len(lower) == len(upper) == expected_length):
            raise ValueError(
                "factorial_interaction means and intervals must be aligned with factor A"
            )
        for index in range(expected_length):
            if not lower[index] <= means[index] <= upper[index]:
                raise ValueError(
                    "factorial_interaction interval must contain its cell mean"
                )
        series_tokens.append(token)
        series_labels.append(label)
        validated.append((label, means, lower, upper))
    if len(set(series_tokens)) != len(series_tokens) or len(set(series_labels)) != len(
        series_labels
    ):
        raise ValueError("factorial_interaction factor B series must be unique")

    fig, ax = plt.subplots(figsize=(7.2, 4.8), constrained_layout=True)
    x_positions = list(range(expected_length))
    for label, means, lower, upper in validated:
        lower_error = [means[index] - lower[index] for index in x_positions]
        upper_error = [upper[index] - means[index] for index in x_positions]
        ax.errorbar(
            x_positions,
            means,
            yerr=[lower_error, upper_error],
            marker="o",
            linewidth=2,
            capsize=4,
            label=label,
        )
    ax.set_xticks(x_positions)
    ax.set_xticklabels(x_labels)
    ax.set_xlim(-0.2, expected_length - 0.8)
    ax.set_xlabel(spec.x_label)
    ax.set_ylabel(spec.y_label)
    ax.set_title(spec.title)
    ax.legend(title=factor_b_label, frameon=False)
    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _odds_ratio_ticks(lower: float, upper: float) -> tuple[list[float], list[str]]:
    if not 0.0 < lower < upper:
        raise ValueError("Odds-ratio tick bounds must be positive and ordered")
    minimum_exponent = math.floor(math.log10(lower))
    maximum_exponent = math.ceil(math.log10(upper))
    ticks = [
        multiplier * (10.0**exponent)
        for exponent in range(minimum_exponent, maximum_exponent + 1)
        for multiplier in (1.0, 2.0, 5.0)
        if lower <= multiplier * (10.0**exponent) <= upper
    ]
    if len(ticks) < 2:
        ticks = [lower, upper]

    def label(value: float) -> str:
        if 0.001 <= value < 10000.0:
            return f"{value:.6g}"
        return f"{value:.1e}"

    return ticks, [label(value) for value in ticks]


def _render_odds_ratio_forest(spec: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()
    rows = list(spec.data.get("rows", []))
    if not rows:
        raise ValueError("odds_ratio_forest chart requires coefficient rows")
    labels = [str(row.get("name", "")) for row in rows]
    estimates = [float(row["odds_ratio"]) for row in rows]
    intervals = [row["ci"] for row in rows]
    lower = [float(interval[0]) for interval in intervals]
    upper = [float(interval[1]) for interval in intervals]
    if any(value <= 0.0 for value in [*estimates, *lower, *upper]):
        raise ValueError("odds_ratio_forest values must be positive")

    fig_height = max(3.5, 0.55 * len(rows) + 1.4)
    fig, ax = plt.subplots(figsize=(7.2, fig_height), constrained_layout=True)
    positions = list(range(len(rows)))
    ax.errorbar(
        estimates,
        positions,
        xerr=[
            [estimate - bound for estimate, bound in zip(estimates, lower)],
            [bound - estimate for estimate, bound in zip(estimates, upper)],
        ],
        fmt="o",
        color="#1f77b4",
        ecolor="#555555",
        capsize=4,
    )
    ax.axvline(1.0, color="#999999", linestyle="--", linewidth=1)
    ax.set_xscale("log")
    lower_limit = min(lower) * 0.9
    upper_limit = max(upper) * 1.1
    ax.set_xlim(lower_limit, upper_limit)
    ticks, tick_labels = _odds_ratio_ticks(lower_limit, upper_limit)
    ax.set_xticks(ticks, labels=tick_labels)
    ax.tick_params(axis="x", which="minor", labelbottom=False)
    ax.set_yticks(positions)
    ax.set_yticklabels(labels)
    ax.invert_yaxis()
    ax.set_xlabel(spec.x_label)
    ax.set_ylabel(spec.y_label)
    ax.set_title(spec.title)
    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _render_roc_curve(spec: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()
    rows = list(spec.data.get("rows", []))
    if len(rows) < 2:
        raise ValueError("roc_curve chart requires at least two points")
    false_positive = [float(row["false_positive_rate"]) for row in rows]
    true_positive = [float(row["true_positive_rate"]) for row in rows]

    fig, ax = plt.subplots(figsize=(6.2, 5.4), constrained_layout=True)
    ax.plot(false_positive, true_positive, color="#1f77b4", linewidth=2)
    ax.plot([0, 1], [0, 1], color="#999999", linestyle="--", linewidth=1)
    ax.set_xlim(0.0, 1.0)
    ax.set_ylim(0.0, 1.0)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel(spec.x_label)
    ax.set_ylabel(spec.y_label)
    ax.set_title(spec.title)
    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _render_calibration_plot(spec: ChartSpec, output_path: str | Path) -> None:
    _configure_fonts()
    rows = list(spec.data.get("rows", []))
    if len(rows) < 3:
        raise ValueError("calibration_plot chart requires at least three bins")
    predicted = [float(row["mean_predicted"]) for row in rows]
    observed = [float(row["observed_rate"]) for row in rows]
    sizes = [max(28.0, 8.0 * float(row["count"])) for row in rows]

    fig, ax = plt.subplots(figsize=(6.2, 5.4), constrained_layout=True)
    ax.plot([0, 1], [0, 1], color="#999999", linestyle="--", linewidth=1)
    ax.plot(predicted, observed, color="#1f77b4", linewidth=1.5)
    ax.scatter(predicted, observed, s=sizes, color="#1f77b4", alpha=0.75)
    ax.set_xlim(0.0, 1.0)
    ax.set_ylim(0.0, 1.0)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel(spec.x_label)
    ax.set_ylabel(spec.y_label)
    ax.set_title(spec.title)
    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def _render_single_or_bundle(
    spec: ChartSpec,
    output_path: str | Path,
    key: str | None,
    renderer: Callable[[ChartSpec, str | Path], None],
) -> str | list[str]:
    if key is None:
        single_path = Path(output_path)
        single_path.parent.mkdir(parents=True, exist_ok=True)
        renderer(spec, single_path)
        return str(single_path)

    output_dir = Path(output_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{_safe_chart_key(key)}-{uuid.uuid4().hex[:8]}"
    paths: list[str] = []
    for ext in ("png", "svg", "eps"):
        path = output_dir / f"{stem}.{ext}"
        renderer(spec, path)
        paths.append(str(path))
    return paths


def render_chart(
    spec: ChartSpec, output_path: str | Path, key: str | None = None
) -> str | list[str]:
    if spec.type == "coefficient_forest":
        return _render_single_or_bundle(
            spec, output_path, key, _render_coefficient_forest
        )
    if spec.type in {"residual_vs_fitted", "residual_qq", "cooks_distance"}:
        return _render_single_or_bundle(
            spec, output_path, key, _render_diagnostic_chart
        )
    if spec.type in {"horizontal_bar", "mean_ci_jitter", "box"}:
        return _render_single_or_bundle(spec, output_path, key, _render_legacy_chart)
    if spec.type == "paired_line":
        return _render_single_or_bundle(
            spec, output_path, key, _render_paired_line_chart
        )
    if spec.type == "odds_ratio_forest":
        return _render_single_or_bundle(
            spec, output_path, key, _render_odds_ratio_forest
        )
    if spec.type == "roc_curve":
        return _render_single_or_bundle(spec, output_path, key, _render_roc_curve)
    if spec.type == "calibration_plot":
        return _render_single_or_bundle(
            spec, output_path, key, _render_calibration_plot
        )
    if spec.type == "factorial_interaction":
        return _render_single_or_bundle(
            spec, output_path, key, _render_factorial_interaction
        )
    raise ValueError(f"Unsupported chart type: {spec.type}")


def write_docx(
    prose: list[str],
    tables: dict[str, list[dict[str, str]]],
    figure_paths: dict[str, list[str]],
    path: str | Path,
) -> str:
    docx_path = Path(path)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("APA Report", level=1)

    for sentence in prose:
        document.add_paragraph(sentence)

    for key, rows in tables.items():
        document.add_heading(key, level=2)
        if not rows:
            continue
        columns = list(rows[0].keys())
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell, column in zip(table.rows[0].cells, columns):
            cell.text = str(column)
        for row in rows:
            cells = table.add_row().cells
            for idx, column in enumerate(columns):
                cells[idx].text = str(row.get(column, ""))

    for key, paths in figure_paths.items():
        if not paths:
            continue
        document.add_heading(key, level=2)
        for figure_path in paths:
            if Path(figure_path).suffix.lower() == ".png":
                document.add_picture(str(figure_path))
            else:
                document.add_paragraph(str(figure_path))

    document.save(docx_path)
    return str(docx_path)


def _normalise_include(params: dict[str, object]) -> list[str] | None:
    include = params.get("include")
    analysis_key = params.get("analysis_key")
    if include is None and analysis_key is not None:
        include = [analysis_key]
    if include is None:
        return None
    if isinstance(include, str):
        return [include]
    if not isinstance(include, list) or not all(
        isinstance(item, str) for item in include
    ):
        raise ValueError("Report include must be a list of analysis result keys")
    return list(include)


def _resolve_included(
    ctx: PipelineContext, include: list[str] | None
) -> list[tuple[str, object]]:
    if include is None:
        analyses = [
            (key, value)
            for key, value in ctx.analyses.items()
            if not key.startswith("analysis:")
        ]
        if len(analyses) != 1:
            raise ValueError(
                "ReportStep requires include when the context does not contain exactly one analysis result"
            )
        return analyses

    resolved: list[tuple[str, object]] = []
    for public_key in include:
        if public_key in ctx.analyses:
            resolved.append((public_key, ctx.analyses[public_key]))
            continue
        alias_key = f"analysis:{public_key}"
        if alias_key in ctx.analyses:
            resolved.append((public_key, ctx.analyses[alias_key]))
            continue
        raise ValueError(f"Missing analysis result: {public_key}")
    return resolved


def _cleanup_generated(
    paths: list[Path], output_dir: Path, *, remove_output_dir: bool
) -> None:
    for path in paths:
        try:
            resolved = path.resolve(strict=False)
            if not _is_relative_to_path(resolved, output_dir):
                continue
            if path.is_file() and path.suffix.lower() in {
                ".docx",
                ".png",
                ".svg",
                ".eps",
            }:
                path.unlink(missing_ok=True)
        except OSError:
            pass
    if remove_output_dir:
        try:
            output_dir.rmdir()
        except OSError:
            pass


def _is_relative_to_path(path: Path, base: Path) -> bool:
    try:
        path.relative_to(base)
    except ValueError:
        return False
    return True


def _safe_report_paths(params: dict[str, object]) -> tuple[Path, Path, Path, bool]:
    forbidden_aliases = {"output_docx", "docx_path", "output_path"} & set(params)
    if forbidden_aliases:
        raise ValueError(
            "Report direct output paths are not supported; use output_dir and filename"
        )

    filename = str(params.get("filename", "report.docx"))
    if (
        any(separator in filename for separator in ("/", "\\"))
        or Path(filename).is_absolute()
        or ":" in filename
    ):
        raise ValueError("Report filename must not contain path separators")
    if Path(filename).suffix.lower() != ".docx":
        raise ValueError("Report filename must use .docx extension")

    raw_output_dir = Path(str(params.get("output_dir", ".")))
    output_dir_existed = raw_output_dir.exists()
    if output_dir_existed:
        if raw_output_dir.is_symlink():
            raise ValueError("Report output_dir must not be a symbolic link")
        if not raw_output_dir.is_dir():
            raise ValueError("Report output_dir exists and is not a directory")
    output_dir = raw_output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    raw_docx_path = output_dir / filename
    if raw_docx_path.exists() and raw_docx_path.is_symlink():
        raise ValueError("Report output path must not be a symbolic link")
    docx_path = raw_docx_path.resolve()
    if not _is_relative_to_path(docx_path, output_dir):
        raise ValueError("Report filename resolves outside output_dir")

    chart_dir_param = params.get("chart_dir")
    if chart_dir_param is None:
        chart_dir = output_dir
    else:
        raw_chart_dir = Path(str(chart_dir_param))
        if not raw_chart_dir.is_absolute():
            raw_chart_dir = output_dir / raw_chart_dir
        if raw_chart_dir.exists() and raw_chart_dir.is_symlink():
            raise ValueError("Report chart_dir must not be a symbolic link")
        chart_dir = raw_chart_dir.resolve()
    if not _is_relative_to_path(chart_dir, output_dir):
        raise ValueError("Report chart_dir must stay within output_dir")
    if chart_dir.exists():
        if chart_dir.is_symlink():
            raise ValueError("Report chart_dir must not be a symbolic link")
        if not chart_dir.is_dir():
            raise ValueError("Report chart_dir exists and is not a directory")

    return output_dir, docx_path, chart_dir, not output_dir_existed


def _report_chart_specs(result: object) -> tuple[ChartSpec, ...]:
    specs: list[ChartSpec] = []
    single = getattr(result, "chart_spec", None)
    if single is not None:
        if not isinstance(single, ChartSpec):
            raise ValueError("Report chart_spec must be a ChartSpec")
        specs.append(single)
    multiple = getattr(result, "chart_specs", ())
    if multiple:
        if not isinstance(multiple, (tuple, list)) or not all(
            isinstance(spec, ChartSpec) for spec in multiple
        ):
            raise ValueError("Report chart_specs must contain ChartSpec values")
        specs.extend(multiple)
    return tuple(specs)


SELECTION_DISCLOSURE = {
    "ko": (
        "분석 방법 선택에 실험적 후보 안내가 사용되었습니다. "
        "계산 모듈의 수치 검증 범위와 추천 타당성은 별개입니다."
    ),
    "en": (
        "An experimental analysis-candidate aid was used to select this method. "
        "Numerical validation of the calculation module and validity of the "
        "recommendation are separate."
    ),
}

RESEARCH_OS_SELECTION_DISCLOSURE = {
    "ko": (
        "분석 방법 선택에 로컬 Research OS의 실험적 후보가 사용되었습니다. "
        "이 기록은 설정의 출처를 표시할 뿐 추천 타당성을 보증하지 않습니다."
    ),
    "en": (
        "A local Research OS experimental candidate was used to select this method. "
        "This records the setting's origin; it does not guarantee recommendation "
        "validity."
    ),
}


@dataclass
class ReportStep(Step):
    step_type = "report.apa"
    produces_analysis = True
    safe_for_untrusted_project_json = False

    def compute(self, ctx: PipelineContext) -> StepResult:
        language_param = str(self.params.get("language", "ko"))
        language_base = language_param.lower().split("-")[0]
        if language_base not in {"ko", "en"}:
            raise ValueError("Unsupported report language")
        language = "en" if language_base == "en" else "ko"
        selection_origin = str(self.params.get("selection_origin", "manual"))
        if selection_origin not in {
            "manual",
            "experimental_candidate_assisted",
            "research_os_assisted",
        }:
            raise ValueError("Unsupported selection origin")

        include = _normalise_include(self.params)
        include_figures = bool(self.params.get("include_figures", True))
        output_dir, docx_path, chart_dir, remove_output_dir_on_failure = (
            _safe_report_paths(self.params)
        )
        included_results = _resolve_included(ctx, include)

        created_paths: list[Path] = []
        prose: list[str] = []
        if selection_origin == "experimental_candidate_assisted":
            prose.append(SELECTION_DISCLOSURE[language])
        elif selection_origin == "research_os_assisted":
            prose.append(RESEARCH_OS_SELECTION_DISCLOSURE[language])
        tables: dict[str, list[dict[str, str]]] = {}
        figure_paths: dict[str, list[str]] = {}
        try:
            for public_key, result in included_results:
                prose.append(prose_for(result, language=language))
                tables[public_key] = table_for(result)
                chart_specs = _report_chart_specs(result)
                paths: list[str] = []
                if include_figures:
                    for index, chart_spec in enumerate(chart_specs, start=1):
                        chart_key = (
                            public_key
                            if len(chart_specs) == 1
                            else f"{public_key}-{index}-{chart_spec.type}"
                        )
                        rendered = render_chart(chart_spec, chart_dir, chart_key)
                        rendered_paths = (
                            [rendered] if isinstance(rendered, str) else list(rendered)
                        )
                        paths.extend(rendered_paths)
                        created_paths.extend(Path(path) for path in rendered_paths)
                figure_paths[public_key] = paths

            write_docx(prose, tables, figure_paths, docx_path)
            created_paths.append(docx_path)
        except Exception:
            if docx_path.exists():
                created_paths.append(docx_path)
            _cleanup_generated(
                created_paths,
                output_dir,
                remove_output_dir=remove_output_dir_on_failure,
            )
            raise

        report = ReportResult(
            prose=prose,
            tables=tables,
            docx_path=str(docx_path),
            figure_paths=figure_paths,
            apa_template_id="report.apa.v1",
        )
        return StepResult(analysis=report)

    def reads(self) -> set[str]:
        include = _normalise_include(self.params)
        if include is None:
            return set()
        reads: set[str] = set()
        for key in include:
            reads.add(key)
            reads.add(f"analysis:{key}")
        return reads

    def writes(self) -> set[str]:
        return {f"report:{self.id}"}


Step.register_type(ReportStep.step_type, ReportStep)
