from __future__ import annotations

import hashlib
import json
import os
import tempfile
import zipfile
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import TypeVar

from defusedxml import ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from modori.recommendation_benchmark import (
    AdjudicationRecord,
    BenchmarkContractError,
    GoldRecord,
    RecommendationIdentity,
    ReviewerAnnotation,
    ReviewerSubmission,
    canonical_json,
)


T = TypeVar("T")

_REVIEWER_SHEETS = (
    "Instructions",
    "Study Cards",
    "Case Reviews",
    "Recommendations",
    "Clarifications",
    "Abstentions",
)
_ADJUDICATION_SHEETS = (*_REVIEWER_SHEETS, "Adjudication", "Resolution Minutes")
_HEADERS = {
    "Study Cards": (
        "case_id",
        "evidence_stage",
        "title_ko",
        "research_question_ko",
        "unit_of_observation",
        "sampling",
        "grouping",
        "time_structure",
        "weights_clusters",
        "variable_meanings",
        "known_missing_codes",
        "facts_visible",
        "facts_clarification_only",
        "data_file",
    ),
    "Case Reviews": (
        "case_id",
        "evidence_stage",
        "action_class",
        "failure_severity",
        "active_minutes",
        "notes",
        "title_ko",
        "research_question_ko",
        "data_file",
    ),
    "Recommendations": (
        "case_id",
        "evidence_stage",
        "rank",
        "family",
        "design_mode",
        "outcome",
        "group",
        "factor_a",
        "factor_b",
        "predictors",
        "items",
        "variables",
        "covariates",
        "measures",
    ),
    "Clarifications": ("case_id", "evidence_stage", "fact_id"),
    "Abstentions": ("case_id", "evidence_stage", "reason_code"),
    "Adjudication": (
        "case_id",
        "evidence_stage",
        "disagreement_summary",
        "decision_notes",
    ),
    "Resolution Minutes": ("case_id", "evidence_stage", "resolution_minutes"),
}
_ROLE_COLUMNS = (
    "outcome",
    "group",
    "factor_a",
    "factor_b",
    "predictors",
    "items",
    "variables",
    "covariates",
    "measures",
)
_INSTRUCTION_CONTRACT = {
    "A1": (
        "각 사례를 다른 검토자와 상의하지 말고 독립적으로 판정하십시오. "
        "통계 용어가 아니라 연구 질문과 설계 사실을 기준으로 기록합니다."
    ),
    "A2": "검토자/판정자 ID",
    "A4": "판정 행동(action_class)",
    "B4": "recommendation_eligible / clarification_required / abstention_required",
    "A5": "목록 입력",
    "B5": "여러 변수는 세미콜론(;)으로 구분합니다.",
    "A6": "시간",
    "B6": "검토자는 사례별 실제 작업 분을 숫자로 입력합니다.",
    "A7": "보조 시트 슬롯",
    "B7": (
        "case_id, evidence_stage, 추천 rank는 미리 채워져 있습니다. "
        "답안 열만 입력하고 식별 열은 수정하지 않습니다."
    ),
    "A8": "빈 슬롯",
    "B8": "해당하지 않는 슬롯은 비워 두고 새 행을 추가하지 않습니다.",
}
_IMMUTABLE_FINGERPRINT_COLUMNS = {
    "Study Cards": tuple(range(1, 15)),
    "Case Reviews": (1, 2, 7, 8, 9),
    "Recommendations": (1, 2, 3),
    "Clarifications": (1, 2),
    "Abstentions": (1, 2),
    "Adjudication": (1, 2),
    "Resolution Minutes": (1, 2),
}
_FIXED_WORKBOOK_DATETIME = datetime(2026, 7, 10, 0, 0, 0)
_FIXED_ZIP_DATETIME = (2026, 7, 10, 0, 0, 0)
_CORE_NAMESPACES = {
    "dcterms": "http://purl.org/dc/terms/",
}


def _nonempty_text(value: object, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise BenchmarkContractError(f"{field_name} must be a non-empty string")
    return value.strip()


def _display_text(value: object) -> str:
    if value is None or not str(value).strip():
        return "없음"
    return str(value).strip()


def _display_mapping(value: object) -> str:
    if not isinstance(value, Mapping) or not value:
        return "없음"
    lines: list[str] = []
    for key in sorted(value, key=str):
        raw = value[key]
        if isinstance(raw, Sequence) and not isinstance(raw, str):
            displayed = ", ".join(str(item) for item in raw) or "없음"
        else:
            displayed = _display_text(raw)
        lines.append(f"{key}: {displayed}")
    return "\n".join(lines)


def _display_sequence(value: object) -> str:
    if isinstance(value, str) or not isinstance(value, Sequence) or not value:
        return "없음"
    return (
        "\n".join(f"• {str(item).strip()}" for item in value if str(item).strip())
        or "없음"
    )


@dataclass(frozen=True)
class PilotCaseSummary:
    case_id: str
    evidence_stage: str
    title_ko: str
    research_question_ko: str
    data_file: str
    unit_of_observation: str = ""
    sampling: str = ""
    grouping: str = ""
    time_structure: str = ""
    weights_clusters: str = ""
    variable_meanings: str = ""
    known_missing_codes: str = ""
    facts_visible: str = ""
    facts_clarification_only: str = ""

    def __post_init__(self) -> None:
        for field_name in (
            "case_id",
            "evidence_stage",
            "title_ko",
            "research_question_ko",
            "data_file",
        ):
            object.__setattr__(
                self,
                field_name,
                _nonempty_text(getattr(self, field_name), field_name),
            )
        for field_name in (
            "unit_of_observation",
            "sampling",
            "grouping",
            "time_structure",
            "weights_clusters",
            "variable_meanings",
            "known_missing_codes",
            "facts_visible",
            "facts_clarification_only",
        ):
            value = getattr(self, field_name)
            if not isinstance(value, str):
                raise BenchmarkContractError(f"{field_name} must be a string")

    @classmethod
    def from_case_mapping(cls, mapping: Mapping[str, object]) -> PilotCaseSummary:
        study_card = mapping.get("study_card")
        if not isinstance(study_card, Mapping):
            raise BenchmarkContractError("pilot case study_card must be an object")
        return cls(
            case_id=_nonempty_text(mapping.get("case_id"), "case_id"),
            evidence_stage=_nonempty_text(
                mapping.get("evidence_stage"),
                "evidence_stage",
            ),
            title_ko=_nonempty_text(mapping.get("title_ko"), "title_ko"),
            research_question_ko=_nonempty_text(
                mapping.get("research_question_ko"),
                "research_question_ko",
            ),
            data_file=_nonempty_text(mapping.get("data_file"), "data_file"),
            unit_of_observation=_display_text(study_card.get("unit_of_observation")),
            sampling=_display_text(study_card.get("sampling")),
            grouping=_display_text(study_card.get("grouping")),
            time_structure=_display_text(study_card.get("time_structure")),
            weights_clusters=_display_text(study_card.get("weights_clusters")),
            variable_meanings=_display_mapping(study_card.get("variable_meanings")),
            known_missing_codes=_display_mapping(study_card.get("known_missing_codes")),
            facts_visible=_display_sequence(study_card.get("facts_visible")),
            facts_clarification_only=_display_sequence(
                study_card.get("facts_clarification_only")
            ),
        )


@dataclass(frozen=True)
class AdjudicatedGold:
    adjudicator_id: str
    gold_records: tuple[GoldRecord, ...]
    adjudications: tuple[AdjudicationRecord, ...]


def _strict_object(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise BenchmarkContractError(f"duplicate JSON key: {key}")
        result[key] = value
    return result


def _reject_json_constant(value: str) -> object:
    raise BenchmarkContractError(f"non-standard JSON constant: {value}")


def read_jsonl(
    path: Path, record_loader: Callable[[Mapping[str, object]], T]
) -> tuple[T, ...]:
    if not path.is_file():
        raise BenchmarkContractError(f"JSONL file does not exist: {path}")
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BenchmarkContractError(f"JSONL file must be UTF-8: {path}") from exc
    records: list[T] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            raise BenchmarkContractError(f"blank line in JSONL at line {line_number}")
        try:
            value = json.loads(
                line,
                object_pairs_hook=_strict_object,
                parse_constant=_reject_json_constant,
            )
        except BenchmarkContractError:
            raise
        except json.JSONDecodeError as exc:
            raise BenchmarkContractError(
                f"invalid JSONL at line {line_number}: {exc.msg}"
            ) from exc
        if not isinstance(value, Mapping):
            raise BenchmarkContractError(
                f"JSONL record at line {line_number} must be an object"
            )
        try:
            records.append(record_loader(value))
        except BenchmarkContractError as exc:
            raise BenchmarkContractError(
                f"invalid JSONL record at line {line_number}: {exc}"
            ) from exc
    if not records:
        raise BenchmarkContractError("JSONL file must contain at least one record")
    return tuple(records)


def _record_mapping(record: object) -> Mapping[str, object]:
    if isinstance(record, Mapping):
        return record
    to_mapping = getattr(record, "to_mapping", None)
    if callable(to_mapping):
        value = to_mapping()
        if isinstance(value, Mapping):
            return value
    raise BenchmarkContractError("JSONL records must be mappings or expose to_mapping")


def write_jsonl(
    path: Path,
    records: Iterable[object],
    *,
    overwrite: bool = False,
) -> None:
    if path.exists() and not overwrite:
        raise BenchmarkContractError(f"output already exists: {path}")
    serialized = [canonical_json(_record_mapping(record)) for record in records]
    if not serialized:
        raise BenchmarkContractError("JSONL output requires at least one record")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="\n",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)
            handle.write("\n".join(serialized))
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        temporary_path.replace(path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def _expected_case_map(
    cases: Iterable[PilotCaseSummary],
) -> dict[tuple[str, str], PilotCaseSummary]:
    result: dict[tuple[str, str], PilotCaseSummary] = {}
    for case in cases:
        key = (case.case_id, case.evidence_stage)
        if key in result:
            raise BenchmarkContractError(f"duplicate pilot case-stage: {key}")
        result[key] = case
    if not result:
        raise BenchmarkContractError("pilot cases must not be empty")
    return result


def _style_header(worksheet: object, column_count: int) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
    for cell in worksheet[1][:column_count]:
        cell.font = Font(bold=True, color="1F2937")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{worksheet.cell(row=1, column=column_count).coordinate}"
    )


def _set_widths(worksheet: object, widths: Mapping[str, float]) -> None:
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def _set_literal_text(cell: object, value: str) -> None:
    cell.value = value
    cell.data_type = "s"


def _build_workbook(
    cases: tuple[PilotCaseSummary, ...],
    *,
    adjudication: bool,
) -> Workbook:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    instructions = workbook.create_sheet("Instructions")
    instructions.merge_cells("A1:B1")
    for coordinate, value in _INSTRUCTION_CONTRACT.items():
        instructions[coordinate] = value
    instructions["B2"] = None
    instructions["A1"].font = Font(bold=True, size=12)
    instructions["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    instructions["B7"].alignment = Alignment(wrap_text=True, vertical="top")
    instructions["B8"].alignment = Alignment(wrap_text=True, vertical="top")
    instructions.row_dimensions[1].height = 64
    instructions.row_dimensions[7].height = 34
    _set_widths(instructions, {"A": 28, "B": 78})

    study_cards = workbook.create_sheet("Study Cards")
    study_cards.append(_HEADERS["Study Cards"])
    for case in cases:
        row = study_cards.max_row + 1
        values = (
            case.case_id,
            case.evidence_stage,
            case.title_ko,
            case.research_question_ko,
            case.unit_of_observation or "없음",
            case.sampling or "없음",
            case.grouping or "없음",
            case.time_structure or "없음",
            case.weights_clusters or "없음",
            case.variable_meanings or "없음",
            case.known_missing_codes or "없음",
            case.facts_visible or "없음",
            case.facts_clarification_only or "없음",
            case.data_file,
        )
        for column, value in enumerate(values, start=1):
            cell = study_cards.cell(row=row, column=column)
            _set_literal_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        study_cards.row_dimensions[row].height = 72
    _style_header(study_cards, len(_HEADERS["Study Cards"]))
    study_cards.freeze_panes = "C2"
    _set_widths(
        study_cards,
        {
            "A": 24,
            "B": 18,
            "C": 30,
            "D": 48,
            "E": 34,
            "F": 32,
            "G": 32,
            "H": 34,
            "I": 38,
            "J": 42,
            "K": 28,
            "L": 42,
            "M": 42,
            "N": 32,
        },
    )

    reviews = workbook.create_sheet("Case Reviews")
    reviews.append(_HEADERS["Case Reviews"])
    for case in cases:
        row = reviews.max_row + 1
        for column, value in (
            (1, case.case_id),
            (2, case.evidence_stage),
            (7, case.title_ko),
            (8, case.research_question_ko),
            (9, case.data_file),
        ):
            _set_literal_text(reviews.cell(row=row, column=column), value)
    action_validation = DataValidation(
        type="list",
        formula1='"recommendation_eligible,clarification_required,abstention_required"',
        allow_blank=True,
    )
    severity_validation = DataValidation(
        type="list",
        formula1='"E1,E2,E3,E4,E5"',
        allow_blank=True,
    )
    reviews.add_data_validation(action_validation)
    reviews.add_data_validation(severity_validation)
    action_validation.add(f"C2:C{len(cases) + 1}")
    severity_validation.add(f"D2:D{len(cases) + 1}")
    _style_header(reviews, len(_HEADERS["Case Reviews"]))
    _set_widths(
        reviews,
        {
            "A": 24,
            "B": 18,
            "C": 25,
            "D": 18,
            "E": 16,
            "F": 35,
            "G": 30,
            "H": 52,
            "I": 32,
        },
    )

    for sheet_name in ("Recommendations", "Clarifications", "Abstentions"):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(_HEADERS[sheet_name])
        slots_per_case = 3 if sheet_name != "Abstentions" else 1
        for case in cases:
            for slot in range(1, slots_per_case + 1):
                row = worksheet.max_row + 1
                _set_literal_text(worksheet.cell(row=row, column=1), case.case_id)
                _set_literal_text(
                    worksheet.cell(row=row, column=2),
                    case.evidence_stage,
                )
                if sheet_name == "Recommendations":
                    worksheet.cell(row=row, column=3).value = slot
        _style_header(worksheet, len(_HEADERS[sheet_name]))
    _set_widths(
        workbook["Recommendations"],
        {
            "A": 24,
            "B": 18,
            "C": 8,
            "D": 28,
            "E": 22,
            "F": 20,
            "G": 20,
            "H": 30,
            "I": 30,
            "J": 30,
            "K": 30,
            "L": 30,
            "M": 30,
            "N": 30,
        },
    )
    _set_widths(workbook["Clarifications"], {"A": 24, "B": 18, "C": 36})
    _set_widths(workbook["Abstentions"], {"A": 24, "B": 18, "C": 36})

    if adjudication:
        adjudication_sheet = workbook.create_sheet("Adjudication")
        adjudication_sheet.append(_HEADERS["Adjudication"])
        resolution_sheet = workbook.create_sheet("Resolution Minutes")
        resolution_sheet.append(_HEADERS["Resolution Minutes"])
        for case in cases:
            adjudication_sheet.append([case.case_id, case.evidence_stage, None, None])
            resolution_sheet.append([case.case_id, case.evidence_stage, None])
        _style_header(adjudication_sheet, len(_HEADERS["Adjudication"]))
        _style_header(resolution_sheet, len(_HEADERS["Resolution Minutes"]))
        _set_widths(
            adjudication_sheet,
            {"A": 24, "B": 18, "C": 48, "D": 48},
        )
        _set_widths(resolution_sheet, {"A": 24, "B": 18, "C": 22})

    return workbook


def _save_reproducible_workbook(workbook: Workbook, path: Path) -> None:
    workbook.properties.created = _FIXED_WORKBOOK_DATETIME
    workbook.properties.modified = _FIXED_WORKBOOK_DATETIME
    raw_descriptor, raw_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.stem}.",
        suffix=".raw.xlsx",
    )
    normalized_descriptor, normalized_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.stem}.",
        suffix=".normalized.xlsx",
    )
    os.close(raw_descriptor)
    os.close(normalized_descriptor)
    raw_path = Path(raw_name)
    normalized_path = Path(normalized_name)
    try:
        workbook.save(raw_path)
        with (
            zipfile.ZipFile(raw_path, "r") as source,
            zipfile.ZipFile(
                normalized_path,
                "w",
            ) as target,
        ):
            for source_info in sorted(
                source.infolist(), key=lambda item: item.filename
            ):
                target_info = zipfile.ZipInfo(
                    filename=source_info.filename,
                    date_time=_FIXED_ZIP_DATETIME,
                )
                target_info.compress_type = source_info.compress_type
                target_info.comment = source_info.comment
                target_info.create_system = source_info.create_system
                target_info.external_attr = source_info.external_attr
                target_info.internal_attr = source_info.internal_attr
                target_info.extract_version = source_info.extract_version
                target_info.create_version = source_info.create_version
                data = source.read(source_info.filename)
                if source_info.filename == "docProps/core.xml":
                    data = _normalize_core_properties(data)
                target.writestr(target_info, data)
        normalized_path.replace(path)
    finally:
        raw_path.unlink(missing_ok=True)
        normalized_path.unlink(missing_ok=True)


def _normalize_core_properties(data: bytes) -> bytes:
    root = ET.fromstring(data)
    fixed_timestamp = _FIXED_WORKBOOK_DATETIME.isoformat() + "Z"
    for tag_name in ("created", "modified"):
        element = root.find(f"{{{_CORE_NAMESPACES['dcterms']}}}{tag_name}")
        if element is not None:
            element.text = fixed_timestamp
    return ET.tostring(root, encoding="utf-8")


def build_blank_pilot_workbooks(
    cases: Iterable[PilotCaseSummary],
    output_dir: Path,
    *,
    overwrite: bool = False,
) -> tuple[Path, Path, Path]:
    case_tuple = tuple(cases)
    _expected_case_map(case_tuple)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = (
        output_dir / "reviewer-a.xlsx",
        output_dir / "reviewer-b.xlsx",
        output_dir / "adjudication.xlsx",
    )
    existing = [path for path in paths if path.exists()]
    if existing and not overwrite:
        raise BenchmarkContractError(f"workbook output already exists: {existing[0]}")
    for path, adjudication in zip(paths, (False, False, True), strict=True):
        workbook = _build_workbook(case_tuple, adjudication=adjudication)
        _save_reproducible_workbook(workbook, path)
        workbook.close()
    return paths


def _worksheet_headers(worksheet: object) -> tuple[str, ...]:
    return tuple(
        str(cell.value) if cell.value is not None else "" for cell in worksheet[1]
    )


def _validate_workbook_structure(workbook: object, *, adjudication: bool) -> None:
    expected_sheets = _ADJUDICATION_SHEETS if adjudication else _REVIEWER_SHEETS
    if tuple(workbook.sheetnames) != expected_sheets:
        raise BenchmarkContractError(
            f"workbook sheets do not match expected schema: {expected_sheets}"
        )
    for sheet_name, headers in _HEADERS.items():
        if sheet_name not in expected_sheets:
            continue
        actual = _worksheet_headers(workbook[sheet_name])
        if actual != headers:
            raise BenchmarkContractError(f"workbook header drift in {sheet_name}")
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise BenchmarkContractError(
                        f"formula cells are forbidden: {worksheet.title}!{cell.coordinate}"
                    )


_PREFILLED_SLOT_COLUMNS = {
    sheet_name: columns
    for sheet_name, columns in _IMMUTABLE_FINGERPRINT_COLUMNS.items()
    if sheet_name
    in {
        "Recommendations",
        "Clarifications",
        "Abstentions",
        "Adjudication",
        "Resolution Minutes",
    }
}


def _prefilled_slot_rows(worksheet: object) -> tuple[tuple[object, ...], ...]:
    columns = _PREFILLED_SLOT_COLUMNS.get(worksheet.title)
    if columns is None:
        return ()
    return tuple(
        tuple(worksheet.cell(row=row, column=column).value for column in columns)
        for row in range(2, worksheet.max_row + 1)
    )


def _validate_prefilled_slots(
    workbook: object,
    expected: Mapping[tuple[str, str], PilotCaseSummary],
) -> None:
    case_keys = tuple(expected)
    expected_rows = {
        "Recommendations": tuple(
            (case_id, evidence_stage, rank)
            for case_id, evidence_stage in case_keys
            for rank in range(1, 4)
        ),
        "Clarifications": tuple(
            (case_id, evidence_stage)
            for case_id, evidence_stage in case_keys
            for _slot in range(3)
        ),
        "Abstentions": case_keys,
        "Adjudication": case_keys,
        "Resolution Minutes": case_keys,
    }
    for sheet_name, rows in expected_rows.items():
        if sheet_name not in workbook.sheetnames:
            continue
        if _prefilled_slot_rows(workbook[sheet_name]) != rows:
            raise BenchmarkContractError(f"{sheet_name} prefilled slot contract drift")


def _validate_instruction_contract(workbook: object) -> None:
    worksheet = workbook["Instructions"]
    if worksheet.max_row != 8 or worksheet.max_column != 2:
        raise BenchmarkContractError("instruction contract drift")
    for row in range(1, 9):
        for column in range(1, 3):
            cell = worksheet.cell(row=row, column=column)
            if cell.coordinate == "B2":
                continue
            if cell.value != _INSTRUCTION_CONTRACT.get(cell.coordinate):
                raise BenchmarkContractError("instruction contract drift")


def _immutable_contract_rows(worksheet: object) -> tuple[tuple[object, ...], ...]:
    if worksheet.title == "Instructions":
        return tuple(
            (
                worksheet.cell(row=row, column=column).coordinate,
                worksheet.cell(row=row, column=column).value,
            )
            for row in range(1, worksheet.max_row + 1)
            for column in range(1, worksheet.max_column + 1)
            if worksheet.cell(row=row, column=column).coordinate != "B2"
        )
    columns = _IMMUTABLE_FINGERPRINT_COLUMNS.get(worksheet.title)
    if columns is None:
        return ()
    return tuple(
        tuple(worksheet.cell(row=row, column=column).value for column in columns)
        for row in range(2, worksheet.max_row + 1)
    )


def workbook_schema_fingerprint(path: Path) -> str:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        sheet_contracts: list[dict[str, object]] = []
        for worksheet in workbook.worksheets:
            validations = sorted(
                (
                    str(validation.sqref),
                    str(validation.formula1),
                    str(validation.type),
                )
                for validation in worksheet.data_validations.dataValidation
            )
            sheet_contracts.append(
                {
                    "name": worksheet.title,
                    "headers": list(_worksheet_headers(worksheet)),
                    "max_row": worksheet.max_row,
                    "max_column": worksheet.max_column,
                    "freeze_panes": (
                        str(worksheet.freeze_panes)
                        if worksheet.freeze_panes is not None
                        else None
                    ),
                    "merged_ranges": sorted(
                        str(cell_range) for cell_range in worksheet.merged_cells.ranges
                    ),
                    "immutable_contract": [
                        list(row) for row in _immutable_contract_rows(worksheet)
                    ],
                    "validations": validations,
                }
            )
        payload = canonical_json({"sheets": sheet_contracts}).encode("utf-8")
        return f"sha256:{hashlib.sha256(payload).hexdigest()}"
    finally:
        workbook.close()


def _case_key_from_row(row: tuple[object, ...], sheet_name: str) -> tuple[str, str]:
    if len(row) < 2:
        raise BenchmarkContractError(f"incomplete row in {sheet_name}")
    return (
        _nonempty_text(row[0], f"{sheet_name} case_id"),
        _nonempty_text(row[1], f"{sheet_name} evidence_stage"),
    )


def _nonempty_rows(worksheet: object) -> Iterable[tuple[object, ...]]:
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if any(value is not None and str(value).strip() for value in row):
            yield row


def _review_rows(
    workbook: object,
    expected: Mapping[tuple[str, str], PilotCaseSummary],
    *,
    require_minutes: bool,
) -> dict[tuple[str, str], tuple[str, str, float | None]]:
    rows: dict[tuple[str, str], tuple[str, str, float | None]] = {}
    for row in _nonempty_rows(workbook["Case Reviews"]):
        key = _case_key_from_row(row, "Case Reviews")
        if key in rows:
            raise BenchmarkContractError(f"duplicate Case Reviews row: {key}")
        if key not in expected:
            raise BenchmarkContractError(f"unexpected Case Reviews row: {key}")
        action_class = _nonempty_text(row[2], "action_class")
        severity = _nonempty_text(row[3], "failure_severity")
        if severity not in {"E1", "E2", "E3", "E4", "E5"}:
            raise BenchmarkContractError("failure_severity must be E1 through E5")
        raw_minutes = row[4]
        minutes: float | None = None
        if require_minutes:
            if isinstance(raw_minutes, bool) or not isinstance(
                raw_minutes, (int, float)
            ):
                raise BenchmarkContractError("active_minutes must be numeric")
            minutes = float(raw_minutes)
        case = expected[key]
        if tuple(row[6:9]) != (
            case.title_ko,
            case.research_question_ko,
            case.data_file,
        ):
            raise BenchmarkContractError(f"case reference drift in Case Reviews: {key}")
        rows[key] = (action_class, severity, minutes)
    if set(rows) != set(expected):
        raise BenchmarkContractError(
            "Case Reviews must cover every expected case-stage"
        )
    return rows


def _validate_study_cards(
    workbook: object,
    expected: Mapping[tuple[str, str], PilotCaseSummary],
) -> None:
    rows: dict[tuple[str, str], tuple[object, ...]] = {}
    for row in _nonempty_rows(workbook["Study Cards"]):
        key = _case_key_from_row(row, "Study Cards")
        if key in rows:
            raise BenchmarkContractError(f"duplicate Study Cards row: {key}")
        if key not in expected:
            raise BenchmarkContractError(f"unexpected Study Cards row: {key}")
        rows[key] = row
    if set(rows) != set(expected):
        raise BenchmarkContractError("Study Cards must cover every expected case-stage")
    for key, case in expected.items():
        expected_values = (
            case.case_id,
            case.evidence_stage,
            case.title_ko,
            case.research_question_ko,
            case.unit_of_observation or "없음",
            case.sampling or "없음",
            case.grouping or "없음",
            case.time_structure or "없음",
            case.weights_clusters or "없음",
            case.variable_meanings or "없음",
            case.known_missing_codes or "없음",
            case.facts_visible or "없음",
            case.facts_clarification_only or "없음",
            case.data_file,
        )
        if tuple(rows[key][:14]) != expected_values:
            raise BenchmarkContractError(f"study card drift for {key}")


def _split_values(value: object) -> tuple[str, ...]:
    if value is None or not str(value).strip():
        return ()
    values = tuple(part.strip() for part in str(value).split(";") if part.strip())
    if len(set(values)) != len(values):
        raise BenchmarkContractError("role list contains duplicate values")
    return values


def _recommendation_rows(
    workbook: object,
    expected: Mapping[tuple[str, str], PilotCaseSummary],
) -> dict[tuple[str, str], tuple[RecommendationIdentity, ...]]:
    ranked: dict[tuple[str, str], list[tuple[int, RecommendationIdentity]]] = {}
    role_start = 5
    role_end = role_start + len(_ROLE_COLUMNS)
    for row in _nonempty_rows(workbook["Recommendations"]):
        if not any(
            value is not None and str(value).strip()
            for value in row[3:role_end]
        ):
            continue
        key = _case_key_from_row(row, "Recommendations")
        if key not in expected:
            raise BenchmarkContractError(f"unexpected Recommendations row: {key}")
        rank = row[2]
        if isinstance(rank, bool) or not isinstance(rank, int) or not 1 <= rank <= 3:
            raise BenchmarkContractError(
                "recommendation rank must be an integer from 1 to 3"
            )
        roles = tuple(
            (name, values)
            for name, raw in zip(
                _ROLE_COLUMNS,
                row[role_start:role_end],
                strict=True,
            )
            if (values := _split_values(raw))
        )
        recommendation = RecommendationIdentity(
            family=_nonempty_text(row[3], "recommendation family"),
            roles=roles,
            design_mode=_nonempty_text(row[4], "recommendation design_mode"),
        )
        ranked.setdefault(key, []).append((rank, recommendation))
    result: dict[tuple[str, str], tuple[RecommendationIdentity, ...]] = {}
    for key, values in ranked.items():
        ranks = [rank for rank, _ in values]
        if len(set(ranks)) != len(ranks):
            raise BenchmarkContractError(f"duplicate recommendation rank for {key}")
        ordered = tuple(recommendation for _, recommendation in sorted(values))
        result[key] = ordered
    return result


def _single_value_rows(
    workbook: object,
    sheet_name: str,
    expected: Mapping[tuple[str, str], PilotCaseSummary],
) -> dict[tuple[str, str], tuple[str, ...]]:
    result: dict[tuple[str, str], list[str]] = {}
    for row in _nonempty_rows(workbook[sheet_name]):
        if len(row) < 3 or row[2] is None or not str(row[2]).strip():
            continue
        key = _case_key_from_row(row, sheet_name)
        if key not in expected:
            raise BenchmarkContractError(f"unexpected {sheet_name} row: {key}")
        value = _nonempty_text(row[2], f"{sheet_name} value")
        result.setdefault(key, []).append(value)
    normalized: dict[tuple[str, str], tuple[str, ...]] = {}
    for key, values in result.items():
        if len(set(values)) != len(values):
            raise BenchmarkContractError(f"duplicate {sheet_name} value for {key}")
        normalized[key] = tuple(values)
    return normalized


def load_reviewer_workbook(
    path: Path,
    expected_cases: Iterable[PilotCaseSummary],
) -> ReviewerSubmission:
    expected = _expected_case_map(expected_cases)
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        _validate_workbook_structure(workbook, adjudication=False)
        _validate_instruction_contract(workbook)
        _validate_study_cards(workbook, expected)
        _validate_prefilled_slots(workbook, expected)
        reviewer_id = _nonempty_text(
            workbook["Instructions"]["B2"].value, "reviewer ID"
        )
        review_rows = _review_rows(workbook, expected, require_minutes=True)
        recommendations = _recommendation_rows(workbook, expected)
        clarifications = _single_value_rows(workbook, "Clarifications", expected)
        abstentions = _single_value_rows(workbook, "Abstentions", expected)
        annotations = tuple(
            ReviewerAnnotation(
                case_id=key[0],
                evidence_stage=key[1],
                action_class=review_rows[key][0],
                acceptable_recommendations=recommendations.get(key, ()),
                required_clarification_facts=clarifications.get(key, ()),
                acceptable_abstention_reasons=abstentions.get(key, ()),
                active_minutes=float(review_rows[key][2]),
                failure_severity=review_rows[key][1],
            )
            for key in expected
        )
        return ReviewerSubmission(reviewer_id=reviewer_id, annotations=annotations)
    finally:
        workbook.close()


def load_adjudication_workbook(
    path: Path,
    expected_cases: Iterable[PilotCaseSummary],
) -> AdjudicatedGold:
    expected = _expected_case_map(expected_cases)
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        _validate_workbook_structure(workbook, adjudication=True)
        _validate_instruction_contract(workbook)
        _validate_study_cards(workbook, expected)
        _validate_prefilled_slots(workbook, expected)
        adjudicator_id = _nonempty_text(
            workbook["Instructions"]["B2"].value,
            "adjudicator ID",
        )
        review_rows = _review_rows(workbook, expected, require_minutes=False)
        recommendations = _recommendation_rows(workbook, expected)
        clarifications = _single_value_rows(workbook, "Clarifications", expected)
        abstentions = _single_value_rows(workbook, "Abstentions", expected)
        gold_records = tuple(
            GoldRecord(
                case_id=key[0],
                evidence_stage=key[1],
                action_class=review_rows[key][0],
                acceptable_recommendations=recommendations.get(key, ()),
                required_clarification_facts=clarifications.get(key, ()),
                acceptable_abstention_reasons=abstentions.get(key, ()),
                failure_severity=review_rows[key][1],
            )
            for key in expected
        )
        resolution_rows: dict[tuple[str, str], AdjudicationRecord] = {}
        for row in _nonempty_rows(workbook["Resolution Minutes"]):
            key = _case_key_from_row(row, "Resolution Minutes")
            if key not in expected:
                raise BenchmarkContractError(
                    f"unexpected Resolution Minutes row: {key}"
                )
            if key in resolution_rows:
                raise BenchmarkContractError(f"duplicate Resolution Minutes row: {key}")
            raw_minutes = row[2]
            if isinstance(raw_minutes, bool) or not isinstance(
                raw_minutes, (int, float)
            ):
                raise BenchmarkContractError("resolution_minutes must be numeric")
            resolution_rows[key] = AdjudicationRecord(
                case_id=key[0],
                evidence_stage=key[1],
                resolution_minutes=float(raw_minutes),
            )
        if set(resolution_rows) != set(expected):
            raise BenchmarkContractError(
                "Resolution Minutes must cover every expected case-stage"
            )
        return AdjudicatedGold(
            adjudicator_id=adjudicator_id,
            gold_records=gold_records,
            adjudications=tuple(resolution_rows[key] for key in expected),
        )
    finally:
        workbook.close()
