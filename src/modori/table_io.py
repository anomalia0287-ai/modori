from __future__ import annotations

import csv
from dataclasses import dataclass
from io import StringIO
from pathlib import Path
from typing import Any

import pandas as pd


class TableReadError(ValueError):
    def __init__(self, message_ko: str) -> None:
        super().__init__(message_ko)
        self.message_ko = message_ko


@dataclass(frozen=True)
class TableReadSource:
    path: Path
    file_type: str
    sheet_name: str | None = None
    sheet_names: tuple[str, ...] = ()


@dataclass(frozen=True)
class TableLayoutOverride:
    sheet_name: str | None = None
    header_row_index: int | None = None
    header_row_count: int | None = None
    data_start_row_index: int | None = None


@dataclass(frozen=True)
class TableInferenceReport:
    file_type: str
    header_row_index: int
    header_row_count: int
    data_start_row_index: int
    column_count: int
    confidence: str
    reasons: tuple[str, ...] = ()
    sheet_name: str | None = None
    sheet_names: tuple[str, ...] = ()
    leading_rows: tuple[tuple[str, ...], ...] = ()

    @property
    def requires_user_confirmation(self) -> bool:
        return self.confidence == "low"


@dataclass(frozen=True)
class TableReadResult:
    frame: pd.DataFrame
    metadata: Any | None
    source: TableReadSource
    warnings: tuple[str, ...] = ()
    inference_report: TableInferenceReport | None = None

    @property
    def columns(self) -> tuple[str, ...]:
        return tuple(str(column) for column in self.frame.columns)

    @property
    def row_count(self) -> int:
        return int(self.frame.shape[0])

    def __iter__(self):
        yield self.frame
        yield self.metadata


@dataclass(frozen=True)
class TableHeaderResult:
    columns: tuple[str, ...]
    source: TableReadSource

    def __iter__(self):
        return iter(self.columns)

    def __len__(self) -> int:
        return len(self.columns)

    def __getitem__(self, index: int) -> str:
        return self.columns[index]


@dataclass(frozen=True)
class TablePreviewResult(TableReadResult):
    preview_limit: int = 0
    sample_rows: tuple[dict[str, Any], ...] = ()

    @property
    def previewed_rows(self) -> int:
        return self.row_count


@dataclass(frozen=True)
class PreviewReadLimits:
    max_file_bytes: int | None = None
    max_rows: int = 30
    max_columns: int | None = 50
    max_cells: int | None = 1_500


@dataclass(frozen=True)
class FullReadLimits:
    max_file_bytes: int | None = None
    max_rows: int | None = None
    max_columns: int | None = None
    max_cells: int | None = None


@dataclass(frozen=True)
class _DelimitedReadContext:
    encoding: str | None
    delimiter: str
    header_row_index: int
    header_row_count: int
    data_start_row_index: int
    header_cells: tuple[Any, ...]
    layout_overridden: bool = False
    warnings: tuple[str, ...] = ()
    leading_rows: tuple[tuple[str, ...], ...] = ()


@dataclass(frozen=True)
class _XlsxReadContext:
    source: TableReadSource
    header_row_index: int
    header_row_count: int
    data_start_row_index: int
    header_cells: tuple[Any, ...]
    header_rows: tuple[tuple[Any, ...], ...] = ()
    layout_overridden: bool = False
    warnings: tuple[str, ...] = ()
    leading_rows: tuple[tuple[str, ...], ...] = ()


DEFAULT_PREVIEW_READ_LIMITS = PreviewReadLimits()
DEFAULT_FULL_READ_LIMITS = FullReadLimits()
_LEGACY_XLS_MESSAGE = "구형 Excel(.xls) 파일은 현재 지원하지 않습니다. Excel에서 .xlsx 또는 .csv로 저장한 뒤 다시 열어 주세요."
_EMPTY_TABLE_MESSAGE = "표 데이터가 없습니다. 원본 포털에서 CSV 파일을 다시 받거나 표가 있는 시트를 선택해 주세요."
_LAYOUT_OVERRIDE_MESSAGE = "사용자 지정 표 레이아웃을 적용했습니다."
_INVALID_LAYOUT_MESSAGE = "지정한 표 레이아웃을 적용할 수 없습니다. 헤더 행과 데이터 시작 행을 확인해 주세요."
_MISSING_SHEET_MESSAGE = "지정한 시트를 찾지 못했습니다. 시트 이름을 확인해 주세요."
_HEADER_INFERENCE_MESSAGE = "표 헤더를 자동으로 찾지 못했습니다. 가져오기 창에서 헤더 행을 지정해 주세요."
_AGGREGATE_ROW_LABELS = frozenset({"합계", "총계", "소계", "총합계", "합계액"})
_AGGREGATE_CONTEXT_LABELS = frozenset({"전체", "계"})
_ROW_NUMBER_MARKERS = frozenset({"-", "－", "–", "—"})
_AGGREGATE_ROW_DETECTED_MESSAGE = "집계/합계 행 {count}개를 감지했습니다. 필요한 경우 가져오기 창에서 제외할 수 있습니다."
_AGGREGATE_ROW_DROPPED_MESSAGE = "집계/합계 행 {count}개를 제외했습니다."
_REVIEW_MAX_ROWS = 25
_REVIEW_MAX_CELLS = 8
_REVIEW_MAX_CELL_CHARS = 60


def normalize_file_type(path: Path, file_type: str | None = None) -> str:
    value = file_type if file_type else path.suffix
    normalized = str(value).lower().lstrip(".")
    if normalized not in {"csv", "xlsx", "xls", "sav"}:
        raise ValueError(f"Unsupported table file type: {value}")
    return normalized


def read_full(
    path: Path,
    file_type: str | None = None,
    *,
    limits: FullReadLimits = DEFAULT_FULL_READ_LIMITS,
    layout: TableLayoutOverride | None = None,
    drop_aggregate_rows: bool = False,
) -> TableReadResult:
    normalized = normalize_file_type(path, file_type)
    _enforce_file_size(path, limits.max_file_bytes)
    if normalized == "csv":
        context = _csv_read_context(path, layout)
        read_kwargs = _csv_read_kwargs(context)
        if limits.max_rows is not None:
            read_kwargs["nrows"] = limits.max_rows + 1
        frame = pd.read_csv(path, **read_kwargs)
        frame, cleanup_warnings = _sanitize_frame(frame, context.header_cells)
        metadata = None
        source = _table_source(path, normalized)
        warnings = _merge_warnings(context.warnings, cleanup_warnings)
        inference_report = _delimited_inference_report(
            context,
            file_type=normalized,
            column_count=len(frame.columns),
        )
    elif normalized == "xlsx":
        if _xlsx_needs_limited_read(limits):
            frame, source, warnings, inference_report = _read_xlsx_limited(
                path,
                limits,
                layout=layout,
            )
        else:
            context = _xlsx_read_context(path, layout)
            source = context.source
            frame = pd.read_excel(path, **_xlsx_read_kwargs(context))
            frame, cleanup_warnings = _sanitize_frame(frame, context.header_cells)
            _reject_notice_only_table(frame, context.header_rows)
            warnings = _merge_warnings(context.warnings, cleanup_warnings)
            inference_report = _xlsx_inference_report(context, column_count=len(frame.columns))
        metadata = None
    elif normalized == "xls":
        if _is_text_table_file(path):
            result = _read_delimited_full(
                path,
                limits,
                file_type=normalized,
                source_warning=("XLS 확장자이지만 텍스트 표로 읽었습니다.",),
            )
            frame = result.frame
            metadata = result.metadata
            source = _table_source(path, normalized)
            warnings = result.warnings
            inference_report = result.inference_report
        else:
            context = _excel_read_context(path, normalized)
            read_kwargs = _excel_read_kwargs(context)
            if limits.max_rows is not None:
                read_kwargs["nrows"] = limits.max_rows + 1
            try:
                frame = pd.read_excel(path, **read_kwargs)
            except ImportError as exc:
                raise TableReadError(_LEGACY_XLS_MESSAGE) from exc
            metadata = None
            source = _table_source(path, normalized)
            frame, cleanup_warnings = _sanitize_frame(frame, context.header_cells)
            warnings = _merge_warnings(context.warnings, cleanup_warnings)
            inference_report = _xlsx_inference_report(context, column_count=len(frame.columns))
    elif normalized == "sav":
        import pyreadstat

        read_kwargs = {"user_missing": True}
        if limits.max_rows is not None:
            read_kwargs["row_limit"] = limits.max_rows + 1
        frame, metadata = pyreadstat.read_sav(path, **read_kwargs)
        source = _table_source(path, normalized)
        warnings = ()
        inference_report = None
    else:
        raise ValueError(f"Unsupported table file type: {normalized}")
    frame, aggregate_warnings = _apply_aggregate_row_policy(
        frame,
        drop=drop_aggregate_rows,
    )
    warnings = _merge_warnings(warnings, aggregate_warnings)
    _enforce_shape_limits(frame, limits)
    return TableReadResult(
        frame=frame,
        metadata=metadata,
        source=source,
        warnings=warnings,
        inference_report=inference_report,
    )


def read_header(
    path: Path,
    file_type: str | None = None,
    *,
    layout: TableLayoutOverride | None = None,
) -> list[str]:
    return list(read_header_result(path, file_type, layout=layout).columns)


def read_header_result(
    path: Path,
    file_type: str | None = None,
    *,
    layout: TableLayoutOverride | None = None,
) -> TableHeaderResult:
    normalized = normalize_file_type(path, file_type)
    if normalized == "csv":
        context = _csv_read_context(path, layout)
        frame = pd.read_csv(path, nrows=0, **_csv_read_kwargs(context))
        frame, _ = _sanitize_frame(frame, context.header_cells)
        columns = frame.columns
        source = _table_source(path, normalized)
    elif normalized == "xlsx":
        context = _safe_xlsx_read_context(path, layout)
        source = context.source
        frame = pd.read_excel(
            path,
            nrows=0,
            **_xlsx_read_kwargs(context),
        )
        frame, _ = _sanitize_frame(frame, context.header_cells)
        columns = frame.columns
    elif normalized == "xls":
        if _is_text_table_file(path):
            context = _csv_read_context(path, layout)
            frame = pd.read_csv(path, nrows=0, **_csv_read_kwargs(context))
            frame, _ = _sanitize_frame(frame, context.header_cells)
            columns = frame.columns
        else:
            try:
                context = _excel_read_context(path, normalized)
                frame = pd.read_excel(path, nrows=0, **_excel_read_kwargs(context))
                frame, _ = _sanitize_frame(frame, context.header_cells)
                columns = frame.columns
            except ImportError as exc:
                raise TableReadError(_LEGACY_XLS_MESSAGE) from exc
        source = _table_source(path, normalized)
    elif normalized == "sav":
        import pyreadstat

        frame, metadata = pyreadstat.read_sav(
            path,
            metadataonly=True,
            user_missing=True,
        )
        columns = getattr(metadata, "column_names", None) or frame.columns
        source = _table_source(path, normalized)
    else:
        raise ValueError(f"Unsupported table file type: {normalized}")
    return TableHeaderResult(
        columns=tuple(str(column) for column in columns),
        source=source,
    )


def read_preview(
    path: Path,
    file_type: str | None = None,
    *,
    limits: PreviewReadLimits = DEFAULT_PREVIEW_READ_LIMITS,
    layout: TableLayoutOverride | None = None,
    drop_aggregate_rows: bool = False,
) -> TablePreviewResult:
    _validate_preview_limits(limits)
    normalized = normalize_file_type(path, file_type)
    if normalized == "csv":
        _enforce_file_size(path, limits.max_file_bytes)
        frame, warnings, inference_report = _read_csv_preview(path, limits, layout)
        return _preview_result(
            frame,
            None,
            _table_source(path, normalized),
            preview_limit=_preview_row_limit(limits),
            warnings=warnings,
            inference_report=inference_report,
            drop_aggregate_rows=drop_aggregate_rows,
        )
    if normalized == "xlsx":
        _enforce_file_size(path, limits.max_file_bytes)
        frame, source, warnings, inference_report = _read_xlsx_preview(path, limits, layout)
        return _preview_result(
            frame,
            None,
            source,
            preview_limit=_preview_row_limit(limits),
            warnings=warnings,
            inference_report=inference_report,
            drop_aggregate_rows=drop_aggregate_rows,
        )
    if normalized == "xls":
        _enforce_file_size(path, limits.max_file_bytes)
        if _is_text_table_file(path):
            frame, warnings, inference_report = _read_delimited_preview(
                path,
                limits,
                file_type=normalized,
                source_warning=("XLS 확장자이지만 텍스트 표로 읽었습니다.",),
            )
        else:
            frame, warnings, inference_report = _read_excel_preview(path, limits)
        return _preview_result(
            frame,
            None,
            _table_source(path, normalized),
            preview_limit=_preview_row_limit(limits),
            warnings=warnings,
            inference_report=inference_report,
            drop_aggregate_rows=drop_aggregate_rows,
        )
    if normalized == "sav":
        import pyreadstat

        _enforce_file_size(path, limits.max_file_bytes)
        _, preview_metadata = pyreadstat.read_sav(
            path,
            metadataonly=True,
            user_missing=True,
        )
        columns = tuple(str(column) for column in getattr(preview_metadata, "column_names", ()))
        selected_columns, warnings = _limited_preview_columns(columns, limits)
        read_kwargs: dict[str, Any] = {
            "row_limit": _preview_row_limit(limits),
            "user_missing": True,
        }
        if selected_columns is not None:
            read_kwargs["usecols"] = list(selected_columns)
        frame, metadata = pyreadstat.read_sav(
            path,
            **read_kwargs,
        )
        return _preview_result(
            frame,
            metadata,
            _table_source(path, normalized),
            preview_limit=_preview_row_limit(limits),
            warnings=warnings,
            drop_aggregate_rows=drop_aggregate_rows,
        )
    raise ValueError(f"Unsupported table file type: {normalized}")


def _read_csv_preview(
    path: Path,
    limits: PreviewReadLimits,
    layout: TableLayoutOverride | None = None,
) -> tuple[pd.DataFrame, tuple[str, ...], TableInferenceReport]:
    return _read_delimited_preview(path, limits, layout=layout)


def _read_delimited_full(
    path: Path,
    limits: FullReadLimits,
    *,
    file_type: str = "csv",
    layout: TableLayoutOverride | None = None,
    source_warning: tuple[str, ...] = (),
) -> TableReadResult:
    context = _csv_read_context(path, layout)
    read_kwargs = _csv_read_kwargs(context)
    if limits.max_rows is not None:
        read_kwargs["nrows"] = limits.max_rows + 1
    frame = pd.read_csv(path, **read_kwargs)
    frame, cleanup_warnings = _sanitize_frame(frame, context.header_cells)
    _enforce_shape_limits(frame, limits)
    return TableReadResult(
        frame=frame,
        metadata=None,
        source=_table_source(path, "csv"),
        warnings=_merge_warnings(source_warning, context.warnings, cleanup_warnings),
        inference_report=_delimited_inference_report(
            context,
            file_type=file_type,
            column_count=len(frame.columns),
            extra_reasons=source_warning,
        ),
    )


def _read_delimited_preview(
    path: Path,
    limits: PreviewReadLimits,
    *,
    file_type: str = "csv",
    layout: TableLayoutOverride | None = None,
    source_warning: tuple[str, ...] = (),
) -> tuple[pd.DataFrame, tuple[str, ...], TableInferenceReport]:
    context = _csv_read_context(path, layout)
    header = pd.read_csv(path, nrows=0, **_csv_read_kwargs(context))
    selected_columns, warnings = _limited_preview_columns(
        tuple(str(column) for column in header.columns),
        limits,
    )
    read_kwargs: dict[str, Any] = {
        **_csv_read_kwargs(context),
        "nrows": _preview_row_limit(limits),
    }
    header_cells = context.header_cells
    if selected_columns is not None:
        read_kwargs["usecols"] = list(selected_columns)
        header_cells = _selected_header_cells(header.columns, context.header_cells, selected_columns)
    frame = pd.read_csv(path, **read_kwargs)
    frame, cleanup_warnings = _sanitize_frame(frame, header_cells)
    return (
        frame,
        _merge_warnings(source_warning, context.warnings, warnings, cleanup_warnings),
        _delimited_inference_report(
            context,
            file_type=file_type,
            column_count=len(header.columns),
            extra_reasons=source_warning,
        ),
    )


def _read_excel_preview(
    path: Path,
    limits: PreviewReadLimits,
) -> tuple[pd.DataFrame, tuple[str, ...], TableInferenceReport]:
    context = _excel_read_context(path, "xls")
    try:
        header = pd.read_excel(path, nrows=0, **_excel_read_kwargs(context))
    except (ImportError, ValueError) as exc:
        raise TableReadError(_LEGACY_XLS_MESSAGE) from exc
    selected_columns, warnings = _limited_preview_columns(
        tuple(str(column) for column in header.columns),
        limits,
    )
    read_kwargs: dict[str, Any] = {
        **_excel_read_kwargs(context),
        "nrows": _preview_row_limit(limits),
    }
    header_cells = context.header_cells
    if selected_columns is not None:
        read_kwargs["usecols"] = list(selected_columns)
        header_cells = _selected_header_cells(header.columns, context.header_cells, selected_columns)
    try:
        frame = pd.read_excel(path, **read_kwargs)
    except (ImportError, ValueError) as exc:
        raise TableReadError(_LEGACY_XLS_MESSAGE) from exc
    frame, cleanup_warnings = _sanitize_frame(frame, header_cells)
    return (
        frame,
        _merge_warnings(context.warnings, warnings, cleanup_warnings),
        _xlsx_inference_report(context, column_count=len(header.columns)),
    )


def _read_xlsx_preview(
    path: Path,
    limits: PreviewReadLimits,
    layout: TableLayoutOverride | None = None,
) -> tuple[pd.DataFrame, TableReadSource, tuple[str, ...], TableInferenceReport]:
    return _read_xlsx_rows(
        path,
        _preview_row_limit(limits),
        _preview_column_limit(limits),
        layout=layout,
    )


def _read_xlsx_limited(
    path: Path,
    limits: FullReadLimits,
    context: _XlsxReadContext | None = None,
    layout: TableLayoutOverride | None = None,
) -> tuple[pd.DataFrame, TableReadSource, tuple[str, ...], TableInferenceReport]:
    workbook = _load_xlsx_workbook(path)
    try:
        worksheet = _xlsx_layout_worksheet(workbook, layout)
        context = context or _xlsx_read_context_from_workbook(path, workbook, worksheet, layout)
        iterator = worksheet.iter_rows(values_only=True)
        _skip_rows(iterator, context.data_start_row_index)
        columns = _xlsx_columns(context.header_cells)
        row_limit = _xlsx_limited_row_read_count(limits, len(columns))
        rows = []
        for row in iterator:
            rows.append(list(row))
            if row_limit is not None and len(rows) >= row_limit:
                break
        frame = pd.DataFrame(rows, columns=columns)
        frame, cleanup_warnings = _sanitize_frame(frame, context.header_cells)
        _reject_notice_only_table(frame, context.header_rows)
        return (
            frame,
            context.source,
            _merge_warnings(context.warnings, cleanup_warnings),
            _xlsx_inference_report(context, column_count=len(columns)),
        )
    finally:
        workbook.close()


def _read_xlsx_rows(
    path: Path,
    max_rows: int,
    max_columns: int | None = None,
    layout: TableLayoutOverride | None = None,
) -> tuple[pd.DataFrame, TableReadSource, tuple[str, ...], TableInferenceReport]:
    context = _xlsx_read_context(path, layout)
    workbook = _load_xlsx_workbook(path)
    try:
        worksheet = _xlsx_layout_worksheet(workbook, layout)
        iterator = worksheet.iter_rows(values_only=True)
        _skip_rows(iterator, context.data_start_row_index)
        if not context.header_cells:
            return (
                pd.DataFrame(),
                context.source,
                context.warnings,
                _xlsx_inference_report(context, column_count=0),
            )
        columns = _xlsx_columns(context.header_cells)
        inferred_column_count = len(columns)
        selected_columns, warnings = _limited_preview_columns(tuple(columns), max_columns)
        header_cells = context.header_cells
        if selected_columns is not None:
            columns = list(selected_columns)
            header_cells = tuple(header_cells[: len(columns)])
        rows = []
        for _, row in zip(range(max_rows), iterator, strict=False):
            rows.append(list(row[: len(columns)]))
        frame = pd.DataFrame(rows, columns=columns)
        frame, cleanup_warnings = _sanitize_frame(frame, header_cells)
        _reject_notice_only_table(frame, context.header_rows)
        return (
            frame,
            context.source,
            _merge_warnings(context.warnings, warnings, cleanup_warnings),
            _xlsx_inference_report(context, column_count=inferred_column_count),
        )
    finally:
        workbook.close()


def _load_xlsx_workbook(path: Path):
    from openpyxl import load_workbook

    return load_workbook(path, read_only=True, data_only=True)


def _csv_read_context(path: Path, layout: TableLayoutOverride | None = None) -> _DelimitedReadContext:
    sample = path.read_bytes()[:128 * 1024]
    if not sample:
        return _DelimitedReadContext(
            encoding=None,
            delimiter=",",
            header_row_index=0,
            header_row_count=1,
            data_start_row_index=1,
            header_cells=(),
        )
    encoding, encoding_warnings = _detect_text_encoding(sample)
    text = sample.decode(encoding, errors="strict")
    delimiter = _detect_delimiter(text)
    rows = _parse_delimited_sample(text, delimiter)
    layout_overridden = layout is not None and layout.header_row_index is not None
    if layout_overridden:
        header_row_index = int(layout.header_row_index)
        header_row_count = int(layout.header_row_count or 1)
        data_start_row_index = int(
            layout.data_start_row_index
            if layout.data_start_row_index is not None
            else header_row_index + header_row_count
        )
        _validate_layout_bounds(
            row_count=len(rows),
            header_row_index=header_row_index,
            header_row_count=header_row_count,
            data_start_row_index=data_start_row_index,
        )
    else:
        header_row_index, header_row_count = _detect_header_layout(rows)
        data_start_row_index = header_row_index + header_row_count
    header_rows = tuple(rows[header_row_index : header_row_index + header_row_count])
    header_cells = _flatten_header_rows(header_rows)
    warnings = list(encoding_warnings)
    if layout_overridden:
        warnings.append(_LAYOUT_OVERRIDE_MESSAGE)
    if delimiter != ",":
        warnings.append(f"CSV 구분자: {_delimiter_label(delimiter)}")
    warnings.extend(_header_offset_warnings(header_row_index))
    warnings.extend(_multi_header_warnings(header_row_count))
    return _DelimitedReadContext(
        encoding=None if encoding in {"utf-8", "utf-8-sig"} else encoding,
        delimiter=delimiter,
        header_row_index=header_row_index,
        header_row_count=header_row_count,
        data_start_row_index=data_start_row_index,
        header_cells=header_cells,
        layout_overridden=layout_overridden,
        warnings=tuple(warnings),
        leading_rows=_leading_review_rows(rows, data_start_row_index),
    )


def _detect_text_encoding(sample: bytes) -> tuple[str, tuple[str, ...]]:
    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig", ()
    for encoding in ("utf-8", "cp949", "euc-kr"):
        try:
            sample.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            continue
        if encoding == "utf-8":
            return encoding, ()
        return encoding, (f"CSV 인코딩: {encoding}",)
    return "utf-8", ()


def _detect_delimiter(text: str) -> str:
    sample = "\n".join(line for line in text.splitlines()[:20] if line.strip())
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        counts = {
            delimiter: sum(line.count(delimiter) for line in sample.splitlines())
            for delimiter in (",", "\t", ";", "|")
        }
        delimiter, count = max(counts.items(), key=lambda item: item[1])
        return delimiter if count > 0 else ","


def _delimiter_label(delimiter: str) -> str:
    return "탭" if delimiter == "\t" else delimiter


def _parse_delimited_sample(text: str, delimiter: str) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    for row in csv.reader(StringIO(text), delimiter=delimiter):
        rows.append(tuple(row))
        if len(rows) >= 30:
            break
    return rows


def _validate_layout_bounds(
    *,
    row_count: int,
    header_row_index: int,
    header_row_count: int,
    data_start_row_index: int,
) -> None:
    if (
        header_row_index < 0
        or header_row_count < 1
        or header_row_index + header_row_count > row_count
        or data_start_row_index < header_row_index + header_row_count
        or data_start_row_index > row_count
    ):
        raise TableReadError(_INVALID_LAYOUT_MESSAGE)


def _csv_read_kwargs(context: _DelimitedReadContext) -> dict[str, Any]:
    kwargs: dict[str, Any] = {}
    if context.encoding:
        kwargs["encoding"] = context.encoding
    if context.delimiter != ",":
        kwargs["sep"] = context.delimiter
        kwargs["engine"] = "python"
    if (
        context.header_row_count > 1
        or context.data_start_row_index != context.header_row_index + context.header_row_count
    ):
        kwargs["skiprows"] = context.data_start_row_index
        kwargs["header"] = None
        kwargs["names"] = list(context.header_cells)
    elif context.header_row_index > 0:
        kwargs["skiprows"] = context.header_row_index
    return kwargs


def _leading_review_rows(
    rows: list[tuple[Any, ...]],
    data_start_row_index: int,
) -> tuple[tuple[str, ...], ...]:
    limit = min(len(rows), data_start_row_index + 3, _REVIEW_MAX_ROWS)
    review: list[tuple[str, ...]] = []
    for row in rows[:limit]:
        cells = tuple(
            ("" if cell is None else str(cell).strip())[:_REVIEW_MAX_CELL_CHARS]
            for cell in row[:_REVIEW_MAX_CELLS]
        )
        while cells and not cells[-1]:
            cells = cells[:-1]
        review.append(cells)
    return tuple(review)


def _delimited_inference_report(
    context: _DelimitedReadContext,
    *,
    file_type: str,
    column_count: int,
    extra_reasons: tuple[str, ...] = (),
) -> TableInferenceReport:
    return TableInferenceReport(
        file_type=file_type,
        header_row_index=context.header_row_index,
        header_row_count=context.header_row_count,
        data_start_row_index=context.data_start_row_index,
        column_count=column_count,
        confidence=_table_inference_confidence(
            column_count,
            header_row_index=context.header_row_index,
            layout_overridden=context.layout_overridden,
        ),
        reasons=_table_inference_reasons(
            header_row_index=context.header_row_index,
            header_row_count=context.header_row_count,
            extra_reasons=(
                *extra_reasons,
                *((_LAYOUT_OVERRIDE_MESSAGE,) if context.layout_overridden else ()),
            ),
        ),
        leading_rows=context.leading_rows,
    )


def _table_inference_confidence(
    column_count: int,
    *,
    header_row_index: int,
    layout_overridden: bool = False,
) -> str:
    if column_count < 2:
        return "low"
    if layout_overridden:
        return "high"
    if header_row_index >= 10:
        return "medium"
    return "high"


def _table_inference_reasons(
    *,
    header_row_index: int,
    header_row_count: int,
    extra_reasons: tuple[str, ...] = (),
) -> tuple[str, ...]:
    reasons: list[str] = []
    reasons.extend(extra_reasons)
    reasons.extend(_header_offset_warnings(header_row_index))
    reasons.extend(_multi_header_warnings(header_row_count))
    if not reasons and header_row_index == 0 and header_row_count == 1:
        reasons.append("첫 번째 행을 헤더로 인식했습니다.")
    return tuple(reasons)


def _safe_xlsx_read_context(
    path: Path,
    layout: TableLayoutOverride | None = None,
) -> _XlsxReadContext:
    try:
        return _xlsx_read_context(path, layout)
    except Exception:
        return _XlsxReadContext(
            source=_table_source(path, "xlsx"),
            header_row_index=0,
            header_row_count=1,
            data_start_row_index=1,
            header_cells=(),
        )


def _xlsx_read_context(path: Path, layout: TableLayoutOverride | None = None) -> _XlsxReadContext:
    workbook = _load_xlsx_workbook(path)
    try:
        worksheet = _xlsx_layout_worksheet(workbook, layout)
        return _xlsx_read_context_from_workbook(path, workbook, worksheet, layout)
    finally:
        workbook.close()


def _xlsx_layout_worksheet(workbook: Any, layout: TableLayoutOverride | None) -> Any:
    if layout is not None and layout.sheet_name:
        try:
            return workbook[layout.sheet_name]
        except KeyError as exc:
            raise TableReadError(_MISSING_SHEET_MESSAGE) from exc
    return workbook.active


def _xlsx_read_context_from_workbook(
    path: Path,
    workbook: Any,
    worksheet: Any,
    layout: TableLayoutOverride | None = None,
) -> _XlsxReadContext:
    source = _xlsx_source_from_workbook(path, workbook, worksheet)
    rows: list[tuple[Any, ...]] = []
    for _, row in zip(range(30), worksheet.iter_rows(values_only=True), strict=False):
        rows.append(tuple(row))
    layout_overridden = layout is not None and (
        bool(layout.sheet_name)
        or layout.header_row_index is not None
        or layout.data_start_row_index is not None
    )
    if layout is not None and layout.header_row_index is not None:
        header_row_index = int(layout.header_row_index)
        header_row_count = int(layout.header_row_count or 1)
        data_start_row_index = int(
            layout.data_start_row_index
            if layout.data_start_row_index is not None
            else header_row_index + header_row_count
        )
        _validate_layout_bounds(
            row_count=len(rows),
            header_row_index=header_row_index,
            header_row_count=header_row_count,
            data_start_row_index=data_start_row_index,
        )
    else:
        header_row_index, header_row_count = _detect_header_layout(rows)
        data_start_row_index = header_row_index + header_row_count
    header_rows = tuple(rows[header_row_index : header_row_index + header_row_count])
    header_cells = _flatten_header_rows(header_rows)
    return _XlsxReadContext(
        source=source,
        header_row_index=header_row_index,
        header_row_count=header_row_count,
        data_start_row_index=data_start_row_index,
        header_cells=header_cells,
        header_rows=header_rows,
        layout_overridden=layout_overridden,
        warnings=_merge_warnings(
            ((_LAYOUT_OVERRIDE_MESSAGE,) if layout_overridden else ()),
            tuple(_header_offset_warnings(header_row_index)),
            tuple(_multi_header_warnings(header_row_count)),
        ),
        leading_rows=_leading_review_rows(rows, data_start_row_index),
    )


def _xlsx_read_kwargs(context: _XlsxReadContext) -> dict[str, Any]:
    kwargs: dict[str, Any] = dict(_xlsx_read_excel_kwargs(context.source))
    if context.header_row_count > 1:
        kwargs["skiprows"] = context.data_start_row_index
        kwargs["header"] = None
        kwargs["names"] = list(context.header_cells)
    else:
        kwargs["header"] = context.header_row_index
    return kwargs


def _xlsx_inference_report(context: _XlsxReadContext, *, column_count: int) -> TableInferenceReport:
    return TableInferenceReport(
        file_type=context.source.file_type,
        sheet_name=context.source.sheet_name,
        sheet_names=context.source.sheet_names,
        header_row_index=context.header_row_index,
        header_row_count=context.header_row_count,
        data_start_row_index=context.data_start_row_index,
        column_count=column_count,
        confidence=_table_inference_confidence(
            column_count,
            header_row_index=context.header_row_index,
            layout_overridden=context.layout_overridden,
        ),
        reasons=_table_inference_reasons(
            header_row_index=context.header_row_index,
            header_row_count=context.header_row_count,
            extra_reasons=((_LAYOUT_OVERRIDE_MESSAGE,) if context.layout_overridden else ()),
        ),
        leading_rows=context.leading_rows,
    )


def _excel_read_context(path: Path, file_type: str) -> _XlsxReadContext:
    try:
        sample = pd.read_excel(path, header=None, nrows=30)
    except ImportError as exc:
        raise TableReadError(_LEGACY_XLS_MESSAGE) from exc
    except ValueError as exc:
        raise TableReadError(_LEGACY_XLS_MESSAGE) from exc
    rows = _frame_to_rows(sample)
    header_row_index, header_row_count = _detect_header_layout(rows)
    header_rows = tuple(rows[header_row_index : header_row_index + header_row_count])
    header_cells = _flatten_header_rows(header_rows)
    return _XlsxReadContext(
        source=_table_source(path, file_type),
        header_row_index=header_row_index,
        header_row_count=header_row_count,
        data_start_row_index=header_row_index + header_row_count,
        header_cells=header_cells,
        header_rows=header_rows,
        warnings=_merge_warnings(
            tuple(_header_offset_warnings(header_row_index)),
            tuple(_multi_header_warnings(header_row_count)),
        ),
        leading_rows=_leading_review_rows(rows, header_row_index + header_row_count),
    )


def _excel_read_kwargs(context: _XlsxReadContext) -> dict[str, Any]:
    if context.header_row_count > 1:
        return {
            "skiprows": context.data_start_row_index,
            "header": None,
            "names": list(context.header_cells),
        }
    return {"header": context.header_row_index}


def _frame_to_rows(frame: pd.DataFrame) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    for row in frame.itertuples(index=False, name=None):
        rows.append(tuple(None if pd.isna(value) else value for value in row))
    return rows


def _detect_header_layout(rows: list[tuple[Any, ...]]) -> tuple[int, int]:
    best_index = 0
    best_count = 1
    best_score = float("-inf")
    for index, row in enumerate(rows[:20]):
        width = _row_width(row)
        if width < 2:
            continue
        for count in range(1, min(3, len(rows) - index) + 1):
            header_rows = rows[index : index + count]
            if count > 1 and not _first_header_row_has_multi_level_markers(header_rows):
                continue
            if not _header_block_is_plausible(header_rows):
                continue
            score = _header_layout_score(rows, index, count)
            if score > best_score:
                best_index = index
                best_count = count
                best_score = score
    if best_score == float("-inf") and any(_row_width(row) >= 2 for row in rows):
        raise TableReadError(_HEADER_INFERENCE_MESSAGE)
    return best_index, best_count


def _header_block_is_plausible(rows: list[tuple[Any, ...]]) -> bool:
    width = max((_row_width(row) for row in rows), default=0)
    if width < 2:
        return False
    for row in rows:
        cells = [cell for cell in _trim_cells(row) if cell]
        if not cells:
            return False
        if _looks_like_metadata_row(cells):
            return False
        if _looks_like_data_row(cells):
            return False
    return True


def _first_header_row_has_multi_level_markers(rows: list[tuple[Any, ...]]) -> bool:
    if len(rows) < 2:
        return False
    width = max((_row_width(row) for row in rows), default=0)
    if width < 2:
        return False
    first_row = _header_row_cells(rows[0], width)
    if any(
        not first_cell
        and any(lower_row[index] for lower_row in (_header_row_cells(row, width) for row in rows[1:]))
        for index, first_cell in enumerate(first_row)
    ):
        return True
    non_empty = [cell for cell in first_row if cell]
    if len(set(non_empty)) < len(non_empty):
        return True
    return sum(1 for cell in non_empty if _is_year_like(cell)) >= 2


def _header_row_cells(row: tuple[Any, ...], width: int) -> tuple[str, ...]:
    cells = ["" if cell is None else str(cell).strip() for cell in row[:width]]
    if len(cells) < width:
        cells.extend([""] * (width - len(cells)))
    return tuple(cells)


def _header_layout_score(
    rows: list[tuple[Any, ...]],
    index: int,
    count: int,
) -> float:
    header_rows = rows[index : index + count]
    width = max(_row_width(row) for row in header_rows)
    names = _flatten_header_rows(tuple(header_rows))
    score = float(width)
    score += _following_data_row_score(rows, index + count - 1, width)
    score += min(len(set(names)), len(names))
    if count > 1:
        score += 10 + (count * 3)
    if index == 0:
        score += 1
    else:
        previous_widths = [_row_width(previous) for previous in rows[max(0, index - 3) : index]]
        if any(width > previous_width for previous_width in previous_widths):
            score += 3
        if any(width == previous_width for previous_width in previous_widths):
            score -= 12
    return score


def _looks_like_data_row(cells: list[str]) -> bool:
    numeric_cells = [cell for cell in cells if _is_number_like(cell)]
    if not numeric_cells:
        return False
    if len(numeric_cells) / len(cells) <= 0.35:
        return False
    return not all(_is_year_like(cell) for cell in numeric_cells)


def _is_year_like(value: str) -> bool:
    cleaned = value.strip()
    return len(cleaned) == 4 and cleaned.isdigit() and 1800 <= int(cleaned) <= 2200


def _flatten_header_rows(header_rows: tuple[tuple[Any, ...], ...]) -> tuple[str, ...]:
    if not header_rows:
        return ()
    width = max((_row_width(row) for row in header_rows), default=0)
    if len(header_rows) == 1:
        row = header_rows[0]
        cells = ["" if cell is None else str(cell).strip() for cell in row[:width]]
        if len(cells) < width:
            cells.extend([""] * (width - len(cells)))
        return tuple(cells)
    filled_rows = [_fill_header_row(row, width) for row in header_rows]
    names: list[str] = []
    for column_index in range(width):
        parts: list[str] = []
        for row in filled_rows:
            cell = row[column_index] if column_index < len(row) else ""
            if not cell or cell.lower().startswith("unnamed:"):
                continue
            if parts and parts[-1] == cell:
                continue
            parts.append(cell)
        names.append(" ".join(parts).strip() or f"column_{column_index + 1}")
    return _unique_names(tuple(names))


def _fill_header_row(row: tuple[Any, ...], width: int) -> tuple[str, ...]:
    cells = ["" if cell is None else str(cell).strip() for cell in row[:width]]
    if len(cells) < width:
        cells.extend([""] * (width - len(cells)))
    filled: list[str] = []
    current = ""
    for cell in cells:
        if cell:
            current = cell
        filled.append(current)
    return tuple(filled)


def _detect_header_row_index(rows: list[tuple[Any, ...]]) -> int:
    best_index = 0
    best_score = float("-inf")
    for index, row in enumerate(rows[:20]):
        cells = _trim_cells(row)
        non_empty = [cell for cell in cells if cell]
        if len(non_empty) < 2:
            continue
        width = len(cells)
        score = float(len(non_empty))
        if any(not _is_number_like(cell) for cell in non_empty):
            score += 4
        if all(_is_number_like(cell) for cell in non_empty):
            score -= 8
        if _looks_like_metadata_row(non_empty):
            score -= 10
        score += _following_data_row_score(rows, index, width)
        if len(set(non_empty)) == len(non_empty):
            score += 2
        else:
            score -= 1
        if index == 0:
            score += 1
        elif any(_row_width(previous) < width for previous in rows[:index]):
            score += 2
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def _trim_cells(row: tuple[Any, ...]) -> tuple[str, ...]:
    cells = ["" if cell is None else str(cell).strip() for cell in row]
    while cells and not cells[-1]:
        cells.pop()
    return tuple(cells)


def _row_width(row: tuple[Any, ...]) -> int:
    return len(_trim_cells(row))


def _looks_like_metadata_row(non_empty: list[str]) -> bool:
    if len(non_empty) > 2:
        marker_text = " ".join(non_empty)
        return any(marker in marker_text for marker in ("검색조건", "검색 조건", "조회조건"))
    marker_text = " ".join(non_empty)
    return any(
        marker in marker_text
        for marker in (
            "자료기준",
            "기준일",
            "단위",
            "출처",
            "제공기관",
            "저작권",
            "검색조건",
            "검색 조건",
            "조회조건",
        )
    )


def _is_number_like(value: str) -> bool:
    cleaned = "".join(character for character in value if character not in ",%").strip()
    if not cleaned:
        return False
    try:
        float(cleaned)
    except ValueError:
        return False
    return True


def _following_data_row_score(
    rows: list[tuple[Any, ...]],
    header_index: int,
    header_width: int,
) -> float:
    score = 0.0
    for row in rows[header_index + 1 : header_index + 6]:
        cells = _trim_cells(row)
        non_empty = [cell for cell in cells if cell]
        if len(non_empty) < 2:
            continue
        if abs(len(cells) - header_width) <= 1:
            score += 3
        if any(_is_number_like(cell) for cell in non_empty):
            score += 1
    return score


def _header_offset_warnings(header_row_index: int) -> tuple[str, ...]:
    if header_row_index <= 0:
        return ()
    return (f"표 헤더 앞의 안내 행 {header_row_index}개를 건너뛰었습니다.",)


def _multi_header_warnings(header_row_count: int) -> tuple[str, ...]:
    if header_row_count <= 1:
        return ()
    return (f"다중 헤더 {header_row_count}행을 하나의 열 이름으로 합쳤습니다.",)


def _selected_header_cells(
    columns: Any,
    header_cells: tuple[Any, ...],
    selected_columns: tuple[str, ...],
) -> tuple[Any, ...]:
    column_names = [str(column) for column in columns]
    selected_cells: list[Any] = []
    for selected in selected_columns:
        try:
            index = column_names.index(selected)
        except ValueError:
            selected_cells.append(selected)
            continue
        selected_cells.append(header_cells[index] if index < len(header_cells) else selected)
    return tuple(selected_cells)


def _skip_rows(iterator: Any, count: int) -> None:
    for _ in range(count):
        try:
            next(iterator)
        except StopIteration:
            return


def _sanitize_frame(
    frame: pd.DataFrame,
    header_cells: tuple[Any, ...],
) -> tuple[pd.DataFrame, tuple[str, ...]]:
    if frame.empty and len(frame.columns) == 0:
        return frame, ()
    keep_indexes: list[int] = []
    names: list[str] = []
    dropped_blank = 0
    renamed_blank = 0
    duplicate_count = 0
    seen: dict[str, int] = {}
    for index, current in enumerate(frame.columns):
        raw = header_cells[index] if index < len(header_cells) else current
        name = _clean_column_name(raw)
        series = frame.iloc[:, index]
        if not name:
            if _is_empty_column(series):
                dropped_blank += 1
                continue
            renamed_blank += 1
            name = f"column_{index + 1}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        if count:
            duplicate_count += 1
            name = f"{name}_{count + 1}"
        keep_indexes.append(index)
        names.append(name)
    cleaned = frame.iloc[:, keep_indexes].copy()
    cleaned.columns = names
    warnings: list[str] = []
    if dropped_blank:
        warnings.append(f"빈 열 {dropped_blank}개를 제외했습니다.")
    if renamed_blank:
        warnings.append(f"빈 헤더 {renamed_blank}개를 column_N 형식으로 바꿨습니다.")
    if duplicate_count:
        warnings.append(f"중복 열 이름 {duplicate_count}개를 고유한 이름으로 바꿨습니다.")
    return cleaned, tuple(warnings)


def _apply_aggregate_row_policy(
    frame: pd.DataFrame,
    *,
    drop: bool,
) -> tuple[pd.DataFrame, tuple[str, ...]]:
    row_mask = _aggregate_row_mask(frame)
    count = int(row_mask.sum())
    if count == 0:
        return frame, ()
    if drop:
        return (
            frame.loc[~row_mask].reset_index(drop=True),
            (_AGGREGATE_ROW_DROPPED_MESSAGE.format(count=count),),
        )
    return frame, (_AGGREGATE_ROW_DETECTED_MESSAGE.format(count=count),)


def _aggregate_row_mask(frame: pd.DataFrame) -> pd.Series:
    if frame.empty:
        return pd.Series(False, index=frame.index)
    return frame.apply(_looks_like_aggregate_row, axis=1)


def _looks_like_aggregate_row(row: pd.Series) -> bool:
    labels = tuple(
        _normalize_aggregate_label(value)
        for value in row
        if not _is_blank_value(value)
    )
    if not labels:
        return False
    first = labels[0]
    if first in _AGGREGATE_ROW_LABELS:
        return True
    if _is_row_number_marker(first) and len(labels) > 1:
        return labels[1] in _AGGREGATE_ROW_LABELS
    if first == "전국":
        return any(label in _AGGREGATE_CONTEXT_LABELS for label in labels[1:3])
    return False


def _normalize_aggregate_label(value: object) -> str:
    return "".join(character for character in str(value).strip() if not character.isspace())


def _is_row_number_marker(label: str) -> bool:
    return label in _ROW_NUMBER_MARKERS or label.isdecimal()


def _is_blank_value(value: object) -> bool:
    if pd.isna(value):
        return True
    return str(value).strip() == ""


def _clean_column_name(value: Any) -> str:
    if value is None:
        return ""
    name = str(value).strip()
    if not name:
        return ""
    if name.lower().startswith("unnamed:"):
        return ""
    return name


def _is_empty_column(series: pd.Series) -> bool:
    for value in series:
        if pd.isna(value):
            continue
        if str(value).strip() == "":
            continue
        return False
    return True


def _unique_names(names: tuple[str, ...]) -> tuple[str, ...]:
    seen: dict[str, int] = {}
    unique: list[str] = []
    for name in names:
        count = seen.get(name, 0)
        seen[name] = count + 1
        unique.append(f"{name}_{count + 1}" if count else name)
    return tuple(unique)


def _merge_warnings(*groups: tuple[str, ...]) -> tuple[str, ...]:
    merged: list[str] = []
    for group in groups:
        for warning in group:
            if warning and warning not in merged:
                merged.append(warning)
    return tuple(merged)


def _is_text_table_file(path: Path) -> bool:
    sample = path.read_bytes()[:4096]
    if not sample:
        return False
    if sample.startswith((b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", b"PK\x03\x04")):
        return False
    try:
        encoding, _ = _detect_text_encoding(sample)
        text = sample.decode(encoding, errors="strict")
    except UnicodeDecodeError:
        return False
    delimiter = _detect_delimiter(text)
    rows = _parse_delimited_sample(text, delimiter)
    return any(_row_width(row) >= 2 for row in rows[:5])


def _reject_notice_only_table(frame: pd.DataFrame, header_rows: tuple[tuple[Any, ...], ...]) -> None:
    if not frame.empty or len(frame.columns) > 1:
        return
    header_text = " ".join(
        str(cell).strip()
        for row in header_rows
        for cell in row
        if cell is not None and str(cell).strip()
    )
    if len(frame.columns) <= 1 or not header_text or _looks_like_notice_text(header_text):
        raise TableReadError(_EMPTY_TABLE_MESSAGE)


def _looks_like_notice_text(text: str) -> bool:
    return any(
        marker in text
        for marker in ("□", "참고용", "서비스", "검색조건", "제공하는 정보", "활용하시기")
    )


def _table_source(path: Path, file_type: str) -> TableReadSource:
    return TableReadSource(path=path, file_type=file_type)


def _xlsx_source(path: Path) -> TableReadSource:
    workbook = _load_xlsx_workbook(path)
    try:
        return _xlsx_source_from_workbook(path, workbook, workbook.active)
    finally:
        workbook.close()


def _safe_xlsx_source(path: Path) -> TableReadSource:
    try:
        return _xlsx_source(path)
    except Exception:
        return _table_source(path, "xlsx")


def _xlsx_read_excel_kwargs(source: TableReadSource) -> dict[str, str]:
    if source.sheet_name is None:
        return {}
    return {"sheet_name": source.sheet_name}


def _xlsx_source_from_workbook(
    path: Path,
    workbook: Any,
    worksheet: Any,
) -> TableReadSource:
    return TableReadSource(
        path=path,
        file_type="xlsx",
        sheet_name=getattr(worksheet, "title", None),
        sheet_names=tuple(str(name) for name in getattr(workbook, "sheetnames", ())),
    )


def _preview_result(
    frame: pd.DataFrame,
    metadata: Any | None,
    source: TableReadSource,
    *,
    preview_limit: int,
    warnings: tuple[str, ...] = (),
    inference_report: TableInferenceReport | None = None,
    drop_aggregate_rows: bool = False,
) -> TablePreviewResult:
    frame, aggregate_warnings = _apply_aggregate_row_policy(
        frame,
        drop=drop_aggregate_rows,
    )
    return TablePreviewResult(
        frame=frame,
        metadata=metadata,
        source=source,
        warnings=_merge_warnings(warnings, aggregate_warnings),
        inference_report=inference_report,
        preview_limit=preview_limit,
        sample_rows=_sample_rows(frame),
    )


def _sample_rows(frame: pd.DataFrame) -> tuple[dict[str, Any], ...]:
    rows: list[dict[str, Any]] = []
    for row in frame.to_dict(orient="records"):
        rows.append({str(key): _display_cell(value) for key, value in row.items()})
    return tuple(rows)


def _display_cell(value: Any) -> Any:
    if pd.isna(value):
        return None
    return value


def _xlsx_columns(header: tuple[Any, ...]) -> list[str]:
    return [
        str(value) if value is not None else f"Unnamed: {index}"
        for index, value in enumerate(header)
    ]


def _validate_preview_limits(limits: PreviewReadLimits) -> None:
    if limits.max_rows < 1:
        raise ValueError("Preview max_rows must be at least 1")
    if limits.max_file_bytes is not None and limits.max_file_bytes < 1:
        raise ValueError("Preview max_file_bytes must be at least 1")
    if limits.max_columns is not None and limits.max_columns < 1:
        raise ValueError("Preview max_columns must be at least 1")
    if limits.max_cells is not None and limits.max_cells < 1:
        raise ValueError("Preview max_cells must be at least 1")


def _preview_row_limit(limits: PreviewReadLimits) -> int:
    column_limit = _preview_column_limit(limits)
    if limits.max_cells is None or column_limit is None:
        return limits.max_rows
    return min(limits.max_rows, max(1, limits.max_cells // column_limit))


def _preview_column_limit(limits: PreviewReadLimits) -> int | None:
    column_limit = limits.max_columns
    if limits.max_cells is None:
        return column_limit
    cell_column_limit = max(1, limits.max_cells // limits.max_rows)
    if column_limit is None:
        return cell_column_limit
    return min(column_limit, cell_column_limit)


def _limited_preview_columns(
    columns: tuple[str, ...],
    limits_or_max_columns: PreviewReadLimits | int | None,
) -> tuple[tuple[str, ...] | None, tuple[str, ...]]:
    if isinstance(limits_or_max_columns, PreviewReadLimits):
        max_columns = _preview_column_limit(limits_or_max_columns)
    else:
        max_columns = limits_or_max_columns
    if max_columns is None or len(columns) <= max_columns:
        return None, ()
    return (
        tuple(columns[:max_columns]),
        (f"미리보기 열 제한: {len(columns)}개 중 {max_columns}개 열만 표시합니다.",),
    )


def _xlsx_needs_limited_read(limits: FullReadLimits) -> bool:
    return (
        limits.max_rows is not None
        or limits.max_columns is not None
        or limits.max_cells is not None
    )


def _xlsx_limited_row_read_count(
    limits: FullReadLimits,
    column_count: int,
) -> int | None:
    row_limit = limits.max_rows + 1 if limits.max_rows is not None else None
    if limits.max_cells is not None and column_count > 0:
        cell_row_limit = limits.max_cells // column_count + 1
        if row_limit is None:
            return cell_row_limit
        return min(row_limit, cell_row_limit)
    return row_limit


def _enforce_file_size(path: Path, max_file_bytes: int | None) -> None:
    if max_file_bytes is None:
        return
    if path.stat().st_size > max_file_bytes:
        raise ValueError(f"Table file exceeds the configured size limit of {max_file_bytes} bytes.")


def _enforce_shape_limits(frame: pd.DataFrame, limits: FullReadLimits) -> None:
    row_count = int(frame.shape[0])
    column_count = int(frame.shape[1])
    if limits.max_rows is not None and row_count > limits.max_rows:
        raise ValueError(
            f"Table row count {row_count} exceeds the configured row limit "
            f"of {limits.max_rows}."
        )
    if limits.max_columns is not None and column_count > limits.max_columns:
        raise ValueError(
            f"Table column count {column_count} exceeds the configured column limit "
            f"of {limits.max_columns}."
        )
    if limits.max_cells is None:
        return
    cell_count = int(frame.shape[0] * frame.shape[1])
    if cell_count > limits.max_cells:
        raise ValueError(
            f"Table contains {cell_count} cells, exceeding the cell limit "
            f"of {limits.max_cells}."
        )
