from __future__ import annotations

import os
import zipfile
from pathlib import Path
from shutil import copyfile
from tempfile import NamedTemporaryFile
from time import sleep
from typing import Callable

from docx import Document

from modori.steps.reporting import (
    RESEARCH_OS_SELECTION_DISCLOSURE,
    SELECTION_DISCLOSURE,
)
from modori.ui.contracts import CommandResult, ReportExportOptions


EXPERIMENTAL_DISCLOSURE_KO = (
    "선택 경로 안내: 분석 방법 선택에는 실험적 후보 안내를 사용했습니다. "
    "수치 결과는 직접 실행과 동일한 계산 모듈에서 생성되었습니다."
)
EXPERIMENTAL_DISCLOSURE_EN = (
    "Selection path: Experimental candidate guidance informed the analysis-method "
    "selection. Numerical results were produced by the same calculation module used "
    "for direct execution."
)
RESEARCH_OS_DISCLOSURE_KO = (
    "선택 경로 안내: 로컬 Research OS의 실험적 후보를 검토하고 분석 설정을 "
    "확정했습니다. 이 안내는 추천 타당성을 보증하지 않습니다."
)
RESEARCH_OS_DISCLOSURE_EN = (
    "Selection path: A local Research OS experimental candidate was reviewed and "
    "its analysis settings were confirmed. This notice does not guarantee "
    "recommendation validity."
)
_KNOWN_DISCLOSURES = {
    EXPERIMENTAL_DISCLOSURE_KO,
    EXPERIMENTAL_DISCLOSURE_EN,
    RESEARCH_OS_DISCLOSURE_KO,
    RESEARCH_OS_DISCLOSURE_EN,
    *SELECTION_DISCLOSURE.values(),
    *RESEARCH_OS_SELECTION_DISCLOSURE.values(),
}
_RESTORE_RETRY_DELAYS = (0.02, 0.05, 0.1, 0.2)


def _remove_paragraph(paragraph: object) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def _apply_selection_disclosure(path: Path, options: ReportExportOptions) -> None:
    provenance = options.selection_provenance
    if provenance not in {
        "manual",
        "experimental_candidate_assisted",
        "research_os_assisted",
    }:
        raise RuntimeError("Unsupported selection provenance")

    if not zipfile.is_zipfile(path):
        if provenance != "manual":
            raise RuntimeError("Report cannot represent selection provenance")
        return

    document = Document(path)
    changed = False
    for paragraph in list(document.paragraphs):
        if paragraph.text in _KNOWN_DISCLOSURES:
            _remove_paragraph(paragraph)
            changed = True

    if provenance == "experimental_candidate_assisted":
        disclosure = (
            EXPERIMENTAL_DISCLOSURE_EN
            if options.language == "en"
            else EXPERIMENTAL_DISCLOSURE_KO
        )
        document.add_paragraph(disclosure)
        changed = True
    elif provenance == "research_os_assisted":
        disclosure = (
            RESEARCH_OS_DISCLOSURE_EN
            if options.language == "en"
            else RESEARCH_OS_DISCLOSURE_KO
        )
        document.add_paragraph(disclosure)
        changed = True

    if changed:
        document.save(path)


def _publish_with_selection_disclosure(
    output_path: Path,
    options: ReportExportOptions,
) -> None:
    with NamedTemporaryFile(
        dir=output_path.parent,
        prefix=f".{output_path.name}.",
        suffix=".tmp",
        delete=False,
    ) as staging_file:
        staging_path = Path(staging_file.name)

    try:
        if not _replace_with_retry(output_path, staging_path):
            raise PermissionError("Report staging path remained locked")
        _apply_selection_disclosure(staging_path, options)
        if not _replace_with_retry(staging_path, output_path):
            raise PermissionError("Report output path remained locked")
    except Exception:
        staging_path.unlink(missing_ok=True)
        output_path.unlink(missing_ok=True)
        raise


def _backup_existing(path: Path) -> Path:
    with NamedTemporaryFile(
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".backup",
        delete=False,
    ) as backup_file:
        backup_path = Path(backup_file.name)
    try:
        copyfile(path, backup_path)
    except Exception:
        backup_path.unlink(missing_ok=True)
        raise
    return backup_path


def _restore_backup(backup_path: Path | None, output_path: Path) -> bool:
    if backup_path is None:
        return True
    if not backup_path.exists():
        return False
    return _replace_with_retry(backup_path, output_path)


def _replace_with_retry(source: Path, destination: Path) -> bool:
    for delay in (*_RESTORE_RETRY_DELAYS, None):
        try:
            os.replace(source, destination)
        except PermissionError:
            if delay is None:
                return False
            sleep(delay)
        else:
            return True
    return False


def _restore_failure_result(
    backup_path: Path | None,
    pipeline_version: int,
) -> CommandResult:
    backup_ids = (
        [str(backup_path.resolve())]
        if backup_path is not None and backup_path.exists()
        else []
    )
    return CommandResult(
        ok=False,
        message_ko=(
            "기존 보고서를 자동으로 복구하지 못했습니다. "
            "원래 보고서 폴더의 보관 파일을 확인해 주세요."
        ),
        error_code="report_restore_failed",
        pipeline_version=pipeline_version,
        result_ids=backup_ids,
    )


class ReportExportService:
    def export(
        self,
        *,
        pipeline: object | None,
        options: ReportExportOptions,
        exporter: Callable[[object, ReportExportOptions], str | Path],
        pipeline_version: int,
        expected_output_path: Path | None = None,
    ) -> CommandResult:
        if pipeline is None:
            return CommandResult(
                ok=False,
                message_ko="내보낼 분석 결과가 없습니다.",
                error_code="no_pipeline",
                pipeline_version=pipeline_version,
            )
        expected_path = (
            expected_output_path.resolve() if expected_output_path is not None else None
        )
        backup_path: Path | None = None
        if expected_path is not None and expected_path.exists():
            if not options.replace_existing:
                return CommandResult(
                    ok=False,
                    message_ko=(
                        "같은 이름의 보고서가 이미 있습니다. "
                        "기존 파일을 바꿀지 확인해 주세요."
                    ),
                    error_code="report_destination_exists",
                    pipeline_version=pipeline_version,
                    result_ids=[str(expected_path)],
                )
            try:
                backup_path = _backup_existing(expected_path)
            except Exception:
                return CommandResult(
                    ok=False,
                    message_ko="기존 보고서를 안전하게 보관하지 못했습니다.",
                    error_code="engine_error",
                    pipeline_version=pipeline_version,
                )
        try:
            output_path = Path(exporter(pipeline, options))
        except Exception:
            if expected_path is not None and not _restore_backup(
                backup_path, expected_path
            ):
                return _restore_failure_result(backup_path, pipeline_version)
            return CommandResult(
                ok=False,
                message_ko="보고서를 내보내지 못했습니다.",
                error_code="engine_error",
                pipeline_version=pipeline_version,
            )
        if (
            backup_path is not None
            and expected_path is not None
            and output_path.resolve() != expected_path
        ):
            if not _restore_backup(backup_path, expected_path):
                return _restore_failure_result(backup_path, pipeline_version)
            return CommandResult(
                ok=False,
                message_ko="보고서 저장 경로를 확인하지 못했습니다.",
                error_code="engine_error",
                pipeline_version=pipeline_version,
            )
        if not output_path.exists():
            if expected_path is not None and not _restore_backup(
                backup_path, expected_path
            ):
                return _restore_failure_result(backup_path, pipeline_version)
            return CommandResult(
                ok=False,
                message_ko="보고서 파일이 생성되지 않았습니다.",
                error_code="engine_error",
                pipeline_version=pipeline_version,
            )
        try:
            _publish_with_selection_disclosure(output_path, options)
        except Exception:
            if expected_path is not None and not _restore_backup(
                backup_path, expected_path
            ):
                return _restore_failure_result(backup_path, pipeline_version)
            return CommandResult(
                ok=False,
                message_ko="보고서에 선택 경로 안내를 기록하지 못했습니다.",
                error_code="engine_error",
                pipeline_version=pipeline_version,
            )
        if backup_path is not None:
            backup_path.unlink(missing_ok=True)
        return CommandResult(
            ok=True,
            message_ko="보고서를 내보냈습니다.",
            pipeline_version=pipeline_version,
            result_ids=[str(output_path)],
        )
